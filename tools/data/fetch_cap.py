#!/usr/bin/env python3
"""Fetches the published data the kinds of plant are derived from, into data/sources/raw/cap/.

- The BEA's detailed fixed assets of private nonresidential industries by industry and asset: investment at historical
  cost (the weights of each asset's service life and declining-balance rate), and net stocks and depreciation at
  current cost (each kind's geometric rate, depreciation over the stock).
- The Census Bureau's manufacturers' shipments and unfilled orders of nondefense capital goods (M3 survey, through
  FRED): the months of shipments the orders on the books stand for, an order's wait for its equipment.
- The Census Bureau's months from start to completion of private nonresidential construction projects.

The service lives and declining-balance rates themselves are the BEA's table "BEA Rates of Depreciation, Service
Lives, Declining-Balance Rates, and Hulten-Wykoff Categories", a document of a few pages transcribed in derive_cap.py.

    python3 tools/data/fetch_cap.py [--cache DIR]
"""
import argparse
import csv
import datetime
import io
import json
import re
import tempfile
from pathlib import Path

import openpyxl

from fetch import RAW, cached, log
from fetch_pop import table

BEA = "https://apps.bea.gov/national/FA2004/Details/xls/"
# Workbook, the prefix of its series codes, and the years kept.
BEA_BOOKS = {
    "investment": ("detailnonres_inv1.xlsx", "I3N", [str(y) for y in range(2015, 2025)],
                   "Investment in private nonresidential fixed assets by industry and asset, historical cost, "
                   "millions of dollars"),
    "net_stock": ("detailnonres_stk1.xlsx", "K1N", ["2023", "2024"],
                  "Current-cost net stock of private nonresidential fixed assets by industry and asset, millions of "
                  "dollars, year end"),
    "depreciation": ("detailnonres_dep1.xlsx", "M1N", ["2023", "2024"],
                     "Current-cost depreciation of private nonresidential fixed assets by industry and asset, "
                     "millions of dollars"),
}
FRED = "https://fred.stlouisfed.org/graph/fredgraph.csv?id={series}"
FRED_SERIES = {
    "ANXAUO": "Unfilled orders, nondefense capital goods excluding aircraft, millions of dollars, seasonally adjusted",
    "ANXAVS": "Value of shipments, nondefense capital goods excluding aircraft, millions of dollars, seasonally "
              "adjusted",
}
CENSUS_T1 = "https://www.census.gov/construction/c30/xlsx/t125.xlsx"


def bea(cache: Path, manifest: dict) -> None:
    """Each workbook's series by industry and asset, the years kept."""
    for name, (book, prefix, years, title) in BEA_BOOKS.items():
        ws = openpyxl.load_workbook(cached(cache, book, BEA + book), read_only=True)["Datasets"]
        rows = ws.iter_rows(values_only=True)
        head = [str(h) if h is not None else None for h in next(rows)]
        cols = {y: head.index(y) for y in years if y in head}
        pattern = re.compile(rf"^{prefix}(\w{{4}})1(\w{{4}})\.A$")
        out = []
        for r in rows:
            m = pattern.match(str(r[0] or ""))
            if m:
                for y, c in cols.items():
                    if r[c] is not None:
                        out.append((m.group(1), m.group(2), y, f"{float(r[c]):.1f}"))
        manifest["series"][f"cap/bea_{name}"] = {
            "title": f"{title} (BEA Fixed Assets, detailed estimates, {book})",
            "rows": table(RAW / "cap" / f"bea_{name}.csv", ["industry", "asset", "year", "value"], out),
        }
    manifest["sources"]["bea_fixed_assets"] = {
        "title": "U.S. Bureau of Economic Analysis, Fixed Assets Accounts, detailed estimates by industry and asset",
        "url": BEA + "<workbook>",
        "fetched": datetime.date.today().isoformat(),
    }


def fred(cache: Path, manifest: dict) -> None:
    """The M3 survey's monthly unfilled orders and shipments of nondefense capital goods excluding aircraft."""
    for series, title in FRED_SERIES.items():
        text = cached(cache, f"{series}.csv", FRED.format(series=series)).read_text()
        rows = [(r["observation_date"], r[series]) for r in csv.DictReader(io.StringIO(text))
                if r[series] not in ("", ".")]
        manifest["series"][f"cap/{series}"] = {
            "title": f"{title} (Census Bureau M3 survey, through FRED)",
            "rows": table(RAW / "cap" / f"{series}.csv", ["month", "value"], rows),
        }
    manifest["sources"]["fred_m3"] = {"title": "Federal Reserve Bank of St. Louis, FRED", "url": FRED,
                                      "fetched": datetime.date.today().isoformat()}


def census(cache: Path, manifest: dict) -> None:
    """Average months from start to completion of private nonresidential projects completed in 2024 and 2025, by
    value of project."""
    ws = openpyxl.load_workbook(cached(cache, "t125.xlsx", CENSUS_T1)).active
    rows, types = [], []
    for r in ws.iter_rows(values_only=True):
        head = str(r[0] or "").strip()
        if head.startswith("Value of Project"):
            types = [str(c).strip() if c is not None else None for c in r[1:]]
            continue
        for kind, months in zip(types, r[1:]):
            if kind is not None and isinstance(months, (int, float)) and head.startswith(("All", "$", "Less")):
                rows.append((kind.rstrip("*"), head, f"{float(months):.1f}"))
    manifest["series"]["cap/construction_months"] = {
        "title": "Average number of months from start to completion for private nonresidential construction projects "
                 "completed in 2024-2025, by value of project, all types (Census Bureau, Construction Length of Time, "
                 "Table 1)",
        "rows": table(RAW / "cap" / "construction_months.csv", ["type", "projects", "months"], rows),
    }
    manifest["sources"]["census_length"] = {"title": "U.S. Census Bureau, Construction Length of Time",
                                            "url": CENSUS_T1, "fetched": datetime.date.today().isoformat()}


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--cache", type=Path, default=None)
    args = parser.parse_args()
    cache = args.cache or Path(tempfile.mkdtemp())
    cache.mkdir(parents=True, exist_ok=True)
    manifest = json.loads((RAW / "manifest.json").read_text())
    for source in (bea, fred, census):
        source(cache, manifest)
        log(f"{source.__name__} done")
    (RAW / "manifest.json").write_text(json.dumps(manifest, indent=2, sort_keys=True) + "\n")


if __name__ == "__main__":
    main()
