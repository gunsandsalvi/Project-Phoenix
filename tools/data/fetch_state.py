#!/usr/bin/env python3
"""Fetches the published data on firms' institutions and the state, into data/sources/raw/ (state/ unless noted):

- World Bank, Doing Business 2020 historical data (data as of May 2019): every measured indicator (not the scores or
  ranks) of starting a business, construction permits, registering property, getting credit (bureau and registry
  coverage, legal rights), paying taxes, trading across borders, enforcing contracts and resolving insolvency
  (recovery rate, time, cost); the national figure where a large economy is measured in two cities.
- World Bank, Enterprise Surveys indicators (portal API), by economy and survey year, all firms and by size class:
  finance (investment and working capital by source, loans, credit constraints, electronic payments), firm profile
  (age, ownership, sales, employment and labour productivity growth, capacity use), job flows and jobs share by size,
  innovation.
- OECD tax databases: corporate income tax rates (statutory and small business), central government personal income
  tax rates and thresholds, top rates and their thresholds, employee, employer and self-employed social security
  contribution schedules; Global Revenue Statistics, general government tax revenue by tax category, % of GDP.
- OECD TaxBEN net replacement rates in unemployment, 2019, singles with and without children at 67% and 100% of the
  average wage, by month of unemployment, with and without social assistance.
- IMF Government Finance Statistics by function: general government expense on health (GF07) and education (GF09),
  % of GDP (imf_sdmx/).
- World Bank WDI series through the API: hospital beds, physicians, nurses, health and education spending,
  pupil-teacher ratios, patent applications, R&D spending and researchers (wdi/).
- OECD Main Science and Technology Indicators and ANBERD business R&D by ISIC Rev. 4 industry.
- WIPO patent grants by technology field (35 fields), by office and by origin.
- V-Dem V-Party (version 2): parties' seats, votes, government support and positions, elections 2000-2019.
- ParlGov (development version): parliamentary elections and cabinets since 2000, party left-right positions.
- International IDEA: voter turnout (parliamentary and presidential elections since 1990), electoral system design,
  political finance (public funding and spending limits); the parliamentary term derived from the elections' dates.

Typed tables (small CSVs under state/, read from documents downloaded when this was written; each row names its
table and page): productivity dispersion (Hsieh and Klenow), capital budgeting methods (Graham and Harvey), hurdle
rates and planning horizons (Graham), the BEA's modified Winfrey S-3 retirement pattern, technology adoption lags
(Comin and Hobijn), learning curves (Nagy et al.), the EU late payment directive's terms, and patents in force by
office (WIPO). This script registers them in the manifest.

    python3 tools/data/fetch_state.py [--cache DIR] [--only NAME ...]
"""
import argparse
import csv
import datetime
import io
import json
import statistics
import tempfile
import zipfile
from pathlib import Path

import openpyxl

from fetch import RAW, cached, get, log, merge_manifest
from fetch_pop import countries, table

TODAY = datetime.date.today().isoformat()
OUT = RAW / "state"
FIRST, LAST = 2015, 2025
YEAR = 2019

DB_URL = ("https://archive.doingbusiness.org/content/dam/doingBusiness/excel/db2020/"
          "Historical-data---COMPLETE-dataset-with-scores.xlsx")
# The report whose data are of May 2019, the opening's year.
DB_YEAR = 2020
DB_TOPICS = {"Starting a business", "Dealing with construction permits", "Registering property", "Getting credit",
             "Paying taxes", "Trading across borders", "Enforcing contracts", "Resolving insolvency"}
# The economies measured in two cities: the code of the national figure, which the file weights from the two.
DB_NATIONAL = {"BANG": "BGD", "BRAZ": "BRA", "CHIN": "CHN", "INDI": "IND", "INDO": "IDN", "MEXI": "MEX",
               "NIGE": "NGA", "PAKI": "PAK", "RUSS": "RUS", "US": "USA"}

ES_API = "https://extdataportal.worldbank.org/api/esapi"
ES_DATA = ES_API + "/GetTopicIndicatorData/topicid/{topic}/subtopicid/0/regionid/0/cutsId/6/colstartindex/0?lang=en"
ES_CUTS = ES_API + "/GetCutsAndSubCutsData/topicid/{topic}/cutsId/6/?lang=en"
# Finance, firm profile, jobs and innovation.
# Surveys from this year on: each economy's last one or two before the opening's year and after.
ES_FIRST = 2013
ES_TOPICS = {7: "finance", 13: "firm profile", 89: "jobs", 9: "innovation"}
ES_SIZES = {"subcut_1280": "small_5_19", "subcut_1279": "medium_20_99", "subcut_1278": "large_100_plus"}
# The indicators the portal publishes without a size split, titled from the Enterprise Surveys' indicator
# descriptions.
ES_TITLES = {"jobs2": "Firms that added jobs (%)", "jobs3": "Job expansion (%)", "jobs4": "Share in job expansion (%)",
             "jobs5": "Firms that reduced jobs (%)", "jobs6": "Job contraction (%)",
             "jobs7": "Share in job contraction (%)", "jobs8": "Net job creation (%)",
             "jobs9": "Share in net job creation (%)"}
ES_NAMES = {"Bahamas": "BHS", "Brunei": "BRN", "Cote d'Ivoire": "CIV", "Egypt": "EGY", "Gambia": "GMB",
            "Korea": "KOR", "Micronesia": "FSM", "Russia": "RUS", "Slovakia": "SVK", "Somalia": "SOM",
            "Turkiye": "TUR", "Venezuela": "VEN", "Yemen": "YEM"}

OECD = "https://sdmx.oecd.org/public/rest/data/{agency},{flow},/{key}?startPeriod={first}&endPeriod={last}" \
       "&format=csvfile"
# Name: agency, flow, key, first and last year, title.
OECD_FLOWS = {
    "tax_cit": ("OECD.CTP.TPS", "DSD_TAX_CIT@DF_CIT", "all", FIRST, LAST,
                "Corporate income tax rates, statutory and targeted at small business: central, sub-central and "
                "combined, % of taxable income, OECD Corporate Tax Statistics"),
    "tax_pit_central": ("OECD.CTP.TPS", "DSD_TAX_PIT@DF_PIT_CENT", "all", FIRST, LAST,
                        "Personal income tax, central government: marginal rates (% of taxable income) and income "
                        "thresholds, allowances and credits (national currency) by band, single person, OECD Tax "
                        "Database"),
    "tax_pit_top": ("OECD.CTP.TPS", "DSD_TAX_PIT@DF_PIT_TOP_EARN_THRESH", "all", FIRST, LAST,
                    "Top statutory personal income tax rate and its earnings threshold (national currency, multiple "
                    "of the average wage), and the average wage, OECD Tax Database"),
    "ssc_employee": ("OECD.CTP.TPS", "DSD_TAX_SSC@DF_SSC_EMPLOYEE", "all", FIRST, LAST,
                     "Employee social security contribution schedules: rates by band (% of earnings), lower and "
                     "upper thresholds and maxima (national currency), OECD Tax Database"),
    "ssc_employer": ("OECD.CTP.TPS", "DSD_TAX_SSC@DF_SSC_EMPLOYER", "all", FIRST, LAST,
                     "Employer social security contribution schedules: rates by band (% of earnings), thresholds "
                     "and maxima (national currency), OECD Tax Database"),
    "ssc_self": ("OECD.CTP.TPS", "DSD_TAX_SSC@DF_SSC_SELF", "all", FIRST, LAST,
                 "Self-employed social security contribution schedules: rates by band (% of earnings), thresholds "
                 "and maxima (national currency), OECD Tax Database"),
    "revenue_by_tax": ("OECD.CTP.TPS", "DSD_REV_COMP_GLOBAL@DF_RSGLOBAL", ".TAX_REV.S13.._T.PT_B1GQ.A", FIRST, LAST,
                       "Tax revenue of general government by tax category (OECD classification: income and "
                       "profits, social security, payroll, property incl. recurrent immovable property and estate, "
                       "inheritance and gift, goods and services), % of GDP, OECD Global Revenue Statistics"),
    "net_replacement_rates": ("OECD.ELS.JAI", "DSD_TAXBEN_NRR@DF_NRR",
                              "..PT_INC_DISP_HH_BJL.S_C0+S_C2.Y4_6.AW67+AW100._Z..NO+YES.NO.A", YEAR, YEAR,
                              f"Net replacement rates in unemployment, % of disposable household income before job "
                              f"loss, {YEAR}, singles without children and with two, previous earnings 67% and 100% "
                              "of the average wage, by month of unemployment (1-60), with and without social "
                              "assistance, no housing benefit, OECD TaxBEN"),
    "msti": ("OECD.STI.STP", "DSD_MSTI@DF_MSTI", "all", FIRST, LAST,
             "Main Science and Technology Indicators: R&D expenditure by performing sector and source of funds (% of "
             "GDP, % of GERD or BERD), researchers and R&D personnel (full-time equivalent), triadic and PCT patent "
             "counts, OECD MSTI"),
    "anberd": ("OECD.STI.STP", "DSD_ANBERD@DF_ANBERDi4", "..MA..USD_PPP.V.B", FIRST, LAST,
               "Business enterprise R&D expenditure by main activity, ISIC Rev. 4 industry, current prices, US dollars "
               "PPP, OECD ANBERD"),
}
# The columns of an OECD file that are not dimensions of the observation.
OECD_META = {"STRUCTURE", "STRUCTURE_ID", "STRUCTURE_NAME", "ACTION", "REVENUE_CODE", "REF_AREA", "TIME_PERIOD",
             "OBS_VALUE", "OBS_STATUS", "OBS_STATUS_2", "OBS_STATUS_3", "AUX_OBS_STATUS", "CONF_STATUS", "DECIMALS",
             "BASE_PER",
             "FREQ"}
# MSTI's growth rates and indices, and its national-currency and exchange-rate series, are left out: the levels in
# shares and persons carry what the mechanisms read.
MSTI_UNITS = {"PT_B1GQ", "PT_GERD", "PT_BERD", "PT_GOVERD", "PT_HERD", "PT_GBARD", "PT_GBARD_C", "FTE", "PS",
              "PATN", "PT_RSH", "PT_EMP"}

IMF_SDMX = "https://api.imf.org/external/sdmx/2.1/data/IMF.STA,{flow}/{key}?startPeriod={first}"
IMF_COFOG = {
    "GF07_S13": ("*.S13.G2MF.GF07_T.POGDP_PT.A",
                 "General government expense on health (% of GDP), IMF Government Finance Statistics (COFOG 07)"),
    "GF09_S13": ("*.S13.G2MF.GF09_T.POGDP_PT.A",
                 "General government expense on education (% of GDP), IMF Government Finance Statistics (COFOG 09)"),
}

WB_API = ("https://api.worldbank.org/v2/country/all/indicator/{code}?format=json&date={first}:{last}&per_page=20000"
          "&source=2")
WDI = {
    "SH.MED.BEDS.ZS": "S5.02", "SH.MED.PHYS.ZS": "S5.02", "SH.MED.NUMW.P3": "S5.02", "SH.XPD.CHEX.GD.ZS": "S5.02",
    "SH.XPD.GHED.GD.ZS": "S5.02", "SE.XPD.TOTL.GD.ZS": "S5.02", "SE.PRM.ENRL.TC.ZS": "S5.02",
    "SE.SEC.ENRL.TC.ZS": "S5.02", "IP.PAT.RESD": "S6.01", "IP.PAT.NRES": "S6.01", "GB.XPD.RSDV.GD.ZS": "S6.01",
    "SP.POP.SCIE.RD.P6": "S6.01",
}

WIPO_URL = "https://www.wipo.int/documents/d/ip-statistics/wipo-data-patent-indicators.zip"
WIPO_GRANTS = "dc_indicator_patent_5_grant_by_technology.csv"
WIPO_FIELDS = "patent_technology_field_names.xlsx"
# The office code of a count with no one office, left out of the counts by office.
WIPO_NO_OFFICE = "**"

VPARTY_URL = "https://www.v-dem.net/media/datasets/CPD_V-Party_CSV_v2.zip"
VPARTY_FIRST = 2000
VPARTY_COLUMNS = ["v2paid", "v2paenname", "v2pashname", "v2paseatshare", "v2panumbseat", "v2patotalseat", "v2pavote",
                  "v2pagovsup", "v2pariglef", "v2pawelf", "v2paimmig", "v2paculsup", "v2parelig", "v2xpa_popul",
                  "v2paanteli", "v2papeople"]

PARLGOV = "https://www.parlgov.org/data/parlgov-development_csv-utf-8/{view}.csv"
PARLGOV_FIRST = "2000"

IDEA = "https://www.idea.int/data-tools/export?type=region_only&themeId={theme}&world=all&loc=home"
IDEA_TURNOUT, IDEA_SYSTEMS, IDEA_FINANCE = 293, 307, 302
IDEA_FIRST = 1990
# The political finance questions kept: public funding (28 to 37) and limits on spending (39 to 42).
IDEA_FINANCE_SHEETS = ("Public funding", "Regulations of spending")
IDEA_FINANCE_QUESTIONS = {"28", "29", "30", "31", "32", "33", "34", "35", "36", "37", "39", "40", "41", "42"}
DAYS_A_YEAR = 365.25

TYPED = {
    "typed_tfp_dispersion": (["S1.22", "S1.24"],
                             "Dispersion of plants' log TFPQ and log TFPR around their industry means (standard "
                             "deviation, 75-25 and 90-10 differences; industries weighted by value added), "
                             "manufacturing, United States 1977-97, China 1998-2005, India 1987-94, with the World "
                             "Bank level each economy stands at in 2019",
                             "Typed from Hsieh and Klenow, NBER Working Paper 13290 (2007), Tables 1 and 2. The one "
                             "per-level dispersion the draw can read: the United States for developed, China "
                             "(upper-middle income) for emerging, India (lower-middle income) for developing; the "
                             "method is the paper's (deviations of log TFP from the industry mean, value-added "
                             "weighted). Cusolito and Maloney (2018, Productivity Revisited) and the Enterprise "
                             "Surveys' TFP note (Francis et al. 2020) were read: they give these moments only as "
                             "figures. OECD MultiProd and CompNet publish no downloadable tables."),
    "typed_capital_budgeting_methods": (["S2.131"],
                                        "Share of CFOs always or almost always using each capital budgeting "
                                        "technique (%) and mean score (0 never to 4 always), all, small and large "
                                        "firms, 392 US and Canadian CFOs, 1999",
                                        "Typed from Graham and Harvey (2001), Journal of Financial Economics 60, "
                                        "Table 2."),
    "typed_hurdle_rates": (["S2.131", "S8.138"],
                           "Mean hurdle rate, weighted average cost of capital and their difference (% a year) by "
                           "company characteristic, 220 US firms, Duke CFO Survey March 2019",
                           "Typed from Graham (2022), NBER Working Paper 29841, Table II."),
    "typed_planning_horizon": (["S2.131"],
                               "Years over which US CFOs judge their plans reliable, 2018 and 2013",
                               "Typed from the text of Graham (2022), NBER Working Paper 29841, p. 20; the paper's "
                               "Figure 7 has the distribution only as bars."),
    "typed_winfrey_s3": (["S2.129"],
                         "Cumulative percent of a vintage's expenditure discarded by age, as percent of the mean "
                         "service life: the BEA's modified Winfrey S-3 retirement pattern (retirements from 45% to "
                         "155% of the mean life)",
                         "Typed from BEA, Fixed Assets and Consumer Durable Goods in the United States, 1925-97, "
                         "Table D. With each kind's mean service life (cap/, derive_cap.py) it gives the annual "
                         "retirement hazard by age; the BEA applies it explicitly only to missiles and nuclear fuel, "
                         "other assets' retirements being implied by their geometric rates, and no source measures "
                         "failures of plant as such."),
    "typed_adoption_lags": (["S6.100", "S6.101"],
                            "Estimated adoption lags (years from invention to a country's adoption): mean, standard "
                            "deviation and percentiles over countries, 15 technologies, with invention years",
                            "Typed from Comin and Hobijn, An Exploration of Technology Diffusion (HBS Working Paper "
                            "08-093; AER 2010), Table 2. The CHAT data the lags are estimated from "
                            "(data.nber.org/data-appendix/w15319) were not kept: 6 MB of raw usage series."),
    "typed_learning_curves": (["S6.101"],
                              "Learning curves of 62 technologies: Wright's exponent w (cost against cumulative "
                              "production), the progress ratio 2^-w, production growth and cost decline exponents, "
                              "periods",
                              "Typed from Nagy, Farmer, Bui and Trancik, PLoS ONE 8(2) e52669 (2013), File S1, Table "
                              "1 (the Santa Fe Performance Curve Database)."),
    "typed_late_payment_terms": (["S2.111"],
                                 "EU late payment rules: default and maximum payment periods between undertakings and "
                                 "for public authorities (days), the statutory interest margin (percentage points "
                                 "over the reference rate) and the fixed recovery compensation (EUR)",
                                 "Typed from Directive 2011/7/EU, Articles 2, 3, 4 and 6 (OJ L 48, 23.2.2011)."),
    "typed_patents_in_force": (["S6.101", "S6.113"],
                               "Patent grants by office (total, resident, non-resident), equivalent grants by origin, "
                               "and patents in force by office, 2019, national and regional offices",
                               "Typed from WIPO, World Intellectual Property Indicators 2020, Table A59, read from "
                               "the report's text; office names matched to iso3, regional offices kept under their "
                               "own codes (EPO, OAPI, ARIPO, EAPO, GCCPO, WORLD for the total). The WIPO data "
                               "center has no open download of patents in force."),
}


def number(text):
    """A published figure as a float, or None where the source prints none."""
    try:
        return float(str(text).replace(",", "").replace("%", "").strip())
    except ValueError:
        return None


def doing_business(cache: Path, iso3: set) -> tuple:
    wb = openpyxl.load_workbook(cached(cache, "db_historical.xlsx", DB_URL), read_only=True)
    rows = wb["All Data"].iter_rows(values_only=True)
    next(rows), next(rows)
    topics, header = next(rows), next(rows)
    topic, column_topic = None, []
    for t in topics:
        topic = t.strip() if t else topic
        column_topic.append(topic)
    national = {code for code in DB_NATIONAL.values()}
    out = []
    for r in rows:
        code, year = r[0], r[4]
        if year != DB_YEAR or not code:
            continue
        iso = DB_NATIONAL.get(code, code if code not in national else None)
        if iso not in iso3:
            continue
        for i, name in enumerate(header):
            if column_topic[i] not in DB_TOPICS or not name or name.startswith(("Score", "Rank")):
                continue
            if "DB06-15" in name or "DB05-14" in name or "DB04-15" in name or "DB06-14" in name:
                continue
            value = r[i]
            if value in (None, "", "..", "N/A", "No Practice", "No VAT"):
                continue
            out.append((iso, column_topic[i], name.strip(), value))
    series = {"state/doing_business": {
        "title": f"Doing Business {DB_YEAR} (data as of May 2019): every measured indicator of starting a business, "
                 "construction permits, registering property, getting credit (legal rights, credit information "
                 "depth, registry and bureau coverage % of adults), paying taxes (payments, hours, total tax and "
                 "contribution rate % of profit and its parts), trading across borders (hours and USD, export and "
                 "import, documentary and border), enforcing contracts (days, % of claim), resolving insolvency "
                 "(recovery cents on the dollar, years, % of estate), World Bank",
        "rows": table(OUT / "doing_business.csv", ["iso3", "topic", "indicator", "value"], out),
        "for": ["S2.132", "S5.102", "S5.174"],
        "note": "Economies measured in two cities take the national (population-weighted) figure the file gives.",
    }}
    return {"state/doing_business": {"title": "World Bank, Doing Business historical data (DB2020 complete dataset "
                                              "with scores, corrected 2021)", "url": DB_URL, "fetched": TODAY}}, series


def enterprise_surveys(cache: Path, iso3: set) -> tuple:
    names = {r["name"]: r["iso3"] for r in csv.DictReader((RAW / "wb" / "countries.csv").open())}
    out, titles = [], {}
    for topic in ES_TOPICS:
        data = json.loads(cached(cache, f"es_topic_{topic}.json", ES_DATA.format(topic=topic)).read_text())
        for c in json.loads(cached(cache, f"es_cuts_{topic}.json", ES_CUTS.format(topic=topic)).read_text()):
            titles[c["name"].split("|", 1)[1].split(" - ", 1)[0]] = c["indicatorName"]
        for r in data["table1"]:
            if r["economyType"] != "Economy" or int(r["year"]) < ES_FIRST:
                continue
            iso = ES_NAMES.get(r["economy"], names.get(r["economy"]))
            if iso not in iso3:
                continue
            for key, value in r.items():
                if "|" not in key or number(value) is None:
                    continue
                code, _, cut = key.split("|", 1)[1].partition(" - cut_6|")
                out.append((iso, int(r["year"]), code, ES_SIZES[cut] if cut else "all", value))
    used = {r[2] for r in out}
    titles.update({k: v for k, v in ES_TITLES.items() if k not in titles})
    series = {
        "state/enterprise_surveys": {
            "title": f"Enterprise Surveys indicators by economy and survey year ({ES_FIRST} on), all firms and by size "
                     "(5-19, 20-99, 100+ employees): finance (shares of investment and working capital by source, "
                     "loans, collateral, credit constraints, electronic payments), firm profile (age, ownership, real "
                     "sales, employment and labour productivity growth, capacity use), job flows and jobs share, "
                     "innovation and R&D, World Bank; titles in state/enterprise_surveys_indicators",
            "rows": table(OUT / "enterprise_surveys.csv", ["iso3", "year", "indicator", "size", "value"], out),
            "for": ["S1.22", "S1.481", "S1.482", "S2.111", "S6.113"],
            "note": "The shares of sales and purchases on credit (fin17, fin18) and of working capital from supplier "
                    "credit (fin8, fin22) are not published by the portal; investment financed by supplier credit "
                    "(fin3) is.",
        },
        "state/enterprise_surveys_indicators": {
            "title": "Enterprise Surveys indicator codes and titles",
            "rows": table(OUT / "enterprise_surveys_indicators.csv", ["indicator", "title"],
                          [(k, titles.get(k, "")) for k in used]),
            "for": ["S1.22", "S6.113"],
        },
    }
    return {"state/enterprise_surveys": {"title": "World Bank Enterprise Surveys, indicators portal API",
                                         "url": ES_DATA.format(topic="<topic>"), "fetched": TODAY}}, series


def oecd(cache: Path, iso3: set) -> tuple:
    sources, series = {}, {}
    for name, (agency, flow, key, first, last, title) in OECD_FLOWS.items():
        url = OECD.format(agency=agency, flow=flow, key=key, first=first, last=last)
        with cached(cache, f"oecd_{name}.csv", url).open(encoding="utf-8-sig") as f:
            rows = [r for r in csv.DictReader(f) if r["REF_AREA"] in iso3 and r["OBS_VALUE"]]
        if name == "msti":
            rows = [r for r in rows if r["UNIT_MEASURE"] in MSTI_UNITS and r["TRANSFORMATION"] == "_Z"]
        # Only the dimensions that vary are kept; the constant ones are in the title.
        dims = [c for c in rows[0] if c not in OECD_META and len({r[c] for r in rows}) > 1]
        out = [(r["REF_AREA"],) + tuple(r[c] for c in dims) + (int(r["TIME_PERIOD"]), r["OBS_VALUE"]) for r in rows]
        series[f"state/{name}"] = {
            "title": title,
            "rows": table(OUT / f"{name}.csv", ["iso3"] + [d.lower() for d in dims] + ["year", "value"], out),
            "for": ["S6.101", "S6.113"] if name in ("msti", "anberd") else ["S5.122"] if name == "net_replacement_rates"
            else ["S5.100", "S5.102", "S5.103"],
        }
        sources[f"state/oecd_{name}"] = {"title": f"OECD Data Explorer, SDMX API, {agency} {flow}", "url": url,
                                         "fetched": TODAY}
    return sources, series


def imf_cofog(cache: Path, iso3: set) -> tuple:
    series = {}
    for name, (key, title) in IMF_COFOG.items():
        url = IMF_SDMX.format(flow="GFS_COFOG", key=key, first=FIRST)
        text = get(url, timeout=600, accept="application/vnd.sdmx.data+csv;version=1.0.0").decode("utf-8-sig")
        rows = [(r["COUNTRY"], int(r["TIME_PERIOD"]), r["OBS_VALUE"]) for r in csv.DictReader(io.StringIO(text))
                if r["COUNTRY"] in iso3 and r["OBS_VALUE"] and FIRST <= int(r["TIME_PERIOD"]) <= LAST]
        series[f"imf_sdmx/{name}"] = {"title": title, "for": ["S5.122"],
                                      "rows": table(RAW / "imf_sdmx" / f"{name}.csv", ["iso3", "year", "value"], rows)}
    return {"state/imf_cofog": {"title": "IMF SDMX 2.1 API, GFS_COFOG", "url": IMF_SDMX, "fetched": TODAY}}, series


def wdi(cache: Path, iso3: set) -> tuple:
    series = {}
    for code, step in WDI.items():
        page = json.loads(get(WB_API.format(code=code, first=FIRST, last=LAST)))
        rows = [(r["countryiso3code"], int(r["date"]), r["value"]) for r in page[1] or []
                if r["value"] is not None and r["countryiso3code"] in iso3]
        title = (page[1] or [{}])[0].get("indicator", {}).get("value", code)
        series[f"wdi/{code}"] = {"title": f"{title}, World Bank WDI", "for": [step],
                                 "rows": table(RAW / "wdi" / f"{code}.csv", ["iso3", "year", "value"], rows)}
    return {"state/wdi_api": {"title": "World Bank API, WDI (source 2)", "url": WB_API, "fetched": TODAY}}, series


def wipo(cache: Path, iso3: set) -> tuple:
    iso3_of = {r["iso2"]: r["iso3"] for r in csv.DictReader((RAW / "wb" / "countries.csv").open()) if r["iso2"]}
    z = zipfile.ZipFile(cached(cache, "wipo_patent_indicators.zip", WIPO_URL))
    sheet = openpyxl.load_workbook(io.BytesIO(z.read(WIPO_FIELDS)), read_only=True).worksheets[0]
    fields = {str(r[0]): r[1] for r in sheet.iter_rows(values_only=True) if r[0] is not None and r[1]}
    by_office, by_origin = {}, {}
    for r in csv.DictReader(io.TextIOWrapper(z.open(WIPO_GRANTS), encoding="utf-8-sig")):
        year = int(r["year"])
        if not FIRST <= year <= LAST:
            continue
        office, origin, count = iso3_of.get(r["office"]), iso3_of.get(r["origin"]), int(r["count"])
        if r["office"] != WIPO_NO_OFFICE and office in iso3:
            k = (office, year, r["tec_id"])
            by_office[k] = by_office.get(k, 0) + count
        if origin in iso3:
            k = (origin, year, r["tec_id"])
            by_origin[k] = by_origin.get(k, 0) + count
    rows = [(iso, "office", year, tec, n) for (iso, year, tec), n in by_office.items()] + \
           [(iso, "origin", year, tec, n) for (iso, year, tec), n in by_origin.items()]
    series = {"state/wipo_grants_by_field": {
        "title": "Patent grants by technology field (WIPO 35-field concordance of IPC), by granting office (iso3) and "
                 "by applicant's origin (summed over the offices granting), count, WIPO IP Statistics",
        "rows": table(OUT / "wipo_grants_by_field.csv", ["iso3", "by", "year", "field", "grants"], rows),
        "for": ["S6.101", "S6.113"],
        "note": "By origin sums one origin's grants over every office, so a family granted at several offices counts "
                "at each. Field names in state/wipo_fields.",
    }, "state/wipo_fields": {
        "title": "WIPO technology fields: number and name",
        "rows": table(OUT / "wipo_fields.csv", ["field", "field_name"],
                      [(k, v) for k, v in fields.items() if k.isdigit()]),
        "for": ["S6.101"],
    }}
    return {"state/wipo": {"title": "WIPO IP Statistics, patent indicators by technology (bulk zip)", "url": WIPO_URL,
                           "fetched": TODAY}}, series


def vparty(cache: Path, iso3: set) -> tuple:
    z = zipfile.ZipFile(cached(cache, "vparty_v2.zip", VPARTY_URL))
    name = next(n for n in z.namelist() if n.endswith(".csv"))
    rows = [(r["country_text_id"], int(r["year"])) + tuple(r[c] for c in VPARTY_COLUMNS)
            for r in csv.DictReader(io.TextIOWrapper(z.open(name), encoding="utf-8"))
            if r["country_text_id"] in iso3 and int(r["year"]) >= VPARTY_FIRST]
    series = {"state/vparty": {
        "title": f"Parties at each national legislative election {VPARTY_FIRST}-2019: seat share (%), seats, total "
                 "seats, vote share (%), government support, left-right economic position, welfare, immigration, "
                 "cultural, religious positions, populism and anti-elitism (expert-coded latent scales), V-Dem "
                 "V-Party v2",
        "rows": table(OUT / "vparty.csv", ["iso3", "year"] + VPARTY_COLUMNS, rows),
        "for": ["S5.134"],
    }}
    return {"state/vparty": {"title": "V-Dem Institute, V-Party dataset version 2 (country-party-date, CSV)",
                             "url": VPARTY_URL, "release": "2022-02", "fetched": TODAY}}, series


def parlgov(cache: Path, iso3: set) -> tuple:
    def view(name):
        with cached(cache, f"parlgov_{name}.csv", PARLGOV.format(view=name)).open(encoding="utf-8") as f:
            return list(csv.DictReader(f))
    elections = [(r["country_name_short"], r["election_date"], r["party_id"], r["party_name_english"],
                  r["vote_share"], r["seats"], r["seats_total"], r["left_right"]) for r in view("view_election")
                 if r["country_name_short"] in iso3 and r["election_type"] == "parliament"
                 and r["election_date"] >= PARLGOV_FIRST]
    cabinets = [(r["country_name_short"], r["start_date"], r["cabinet_id"], r["cabinet_name"], r["caretaker"],
                 r["party_id"], r["party_name_english"], r["cabinet_party"], r["prime_minister"], r["seats"],
                 r["election_seats_total"], r["left_right"]) for r in view("view_cabinet")
                if r["country_name_short"] in iso3 and r["start_date"] >= PARLGOV_FIRST]
    used = {r[2] for r in elections} | {r[5] for r in cabinets}
    parties = [(r["country_name_short"], r["party_id"], r["party_name_english"], r["family_name"], r["left_right"],
                r["state_market"], r["liberty_authority"], r["eu_anti_pro"]) for r in view("view_party")
               if r["party_id"] in used]
    series = {
        "state/parlgov_elections": {
            "title": f"Parliamentary elections since {PARLGOV_FIRST}: each party's vote share (%), seats and the "
                     "chamber's seats, with its left-right position (0-10), ParlGov",
            "rows": table(OUT / "parlgov_elections.csv", ["iso3", "date", "party_id", "party", "vote_share", "seats",
                                                         "seats_total", "left_right"], elections),
            "for": ["S5.100", "S5.134"]},
        "state/parlgov_cabinets": {
            "title": f"Cabinets since {PARLGOV_FIRST}: every party in parliament, whether in the cabinet and holding "
                     "the prime minister, its seats, ParlGov",
            "rows": table(OUT / "parlgov_cabinets.csv", ["iso3", "start", "cabinet_id", "cabinet", "caretaker",
                                                        "party_id", "party", "in_cabinet", "prime_minister", "seats",
                                                        "seats_total", "left_right"], cabinets),
            "for": ["S5.100"]},
        "state/parlgov_parties": {
            "title": "Parties' family and positions: left-right, state-market, liberty-authority, EU (0-10 expert "
                     "means), ParlGov",
            "rows": table(OUT / "parlgov_parties.csv", ["iso3", "party_id", "party", "family", "left_right",
                                                       "state_market", "liberty_authority", "eu_anti_pro"], parties),
            "for": ["S5.134"]},
    }
    return {"state/parlgov": {"title": "ParlGov development version, CSV views",
                              "url": PARLGOV.format(view="<view>"), "fetched": TODAY}}, series


def idea(cache: Path, iso3: set) -> tuple:
    def book(theme):
        return openpyxl.load_workbook(cached(cache, f"idea_{theme}.xlsx", IDEA.format(theme=theme)), read_only=True)
    turnout = []
    for r in list(book(IDEA_TURNOUT)["All"].iter_rows(values_only=True))[1:]:
        iso, kind, date = r[2], r[3], str(r[4])
        if iso in iso3 and kind in ("Parliamentary", "Presidential") and int(date[:4]) >= IDEA_FIRST:
            turnout.append((iso, kind, date[:10], number(r[5]), number(r[6]), number(r[7]), number(r[8]),
                            number(r[11]), r[12]))
    systems = [(r[2], int(r[3]), r[4], r[5], r[6], r[7], r[8], r[9])
               for r in list(book(IDEA_SYSTEMS).worksheets[0].iter_rows(values_only=True))[1:]
               if r[2] in iso3 and r[3] and str(r[3]).isdigit()]
    finance = []
    fb = book(IDEA_FINANCE)
    for sheet in IDEA_FINANCE_SHEETS:
        rows = list(fb[sheet].iter_rows(values_only=True))
        for r in rows[1:]:
            if r[2] not in iso3:
                continue
            for question, answer in zip(rows[0][3:], r[3:]):
                if question.split(".", 1)[0] in IDEA_FINANCE_QUESTIONS and answer:
                    finance.append((r[2], question.split(".", 1)[0], question.split(".", 1)[1].strip(),
                                    str(answer).strip()))
    # The legislature's term: the median of the years between consecutive parliamentary elections since 1990.
    dates = {}
    for iso, kind, date, *_ in turnout:
        if kind == "Parliamentary":
            dates.setdefault(iso, []).append(datetime.date.fromisoformat(date))
    cycle = []
    for iso, ds in dates.items():
        ds = sorted(set(ds))
        gaps = [(b - a).days / DAYS_A_YEAR for a, b in zip(ds, ds[1:])]
        if gaps:
            cycle.append((iso, len(ds), f"{statistics.median(gaps):.2f}", f"{min(gaps):.2f}", f"{max(gaps):.2f}"))
    series = {
        "state/idea_turnout": {
            "title": f"Parliamentary and presidential elections since {IDEA_FIRST}: date, turnout of registered "
                     "voters (%), total vote, registered voters, turnout of the voting-age population (%), invalid "
                     "votes (%), compulsory voting, International IDEA Voter Turnout Database",
            "rows": table(OUT / "idea_turnout.csv", ["iso3", "election", "date", "turnout_pct", "total_vote",
                                                    "registered", "vap_turnout_pct", "invalid_pct", "compulsory"],
                          turnout),
            "for": ["S5.100"]},
        "state/idea_electoral_systems": {
            "title": "Electoral system family and system for the national legislature, tiers, legislature size "
                     "(directly elected and voting members), system for the president, by year of change, "
                     "International IDEA Electoral System Design Database",
            "rows": table(OUT / "idea_electoral_systems.csv", ["iso3", "year", "family", "legislature_system",
                                                              "tiers", "seats_elected", "seats_voting",
                                                              "president_system"], systems),
            "for": ["S5.134"]},
        "state/idea_political_finance": {
            "title": "Political finance rules: direct public funding of parties, its eligibility and allocation, "
                     "earmarking, free media access, and limits on parties' and candidates' spending (questions "
                     "28-37, 39-42), International IDEA Political Finance Database",
            "rows": table(OUT / "idea_political_finance.csv", ["iso3", "question", "text", "answer"], finance),
            "for": ["S5.134"],
            "note": "The database records the rules, not their amounts: no public funding per vote is published."},
        "state/election_cycle": {
            "title": f"Years between consecutive parliamentary elections since {IDEA_FIRST}: median, minimum and "
                     "maximum, and the number of elections",
            "rows": table(OUT / "election_cycle.csv", ["iso3", "elections", "median_years", "min_years", "max_years"],
                          cycle),
            "for": ["S5.134"],
            "note": "Derived from state/idea_turnout's parliamentary election dates."},
    }
    return {"state/idea": {"title": "International IDEA data tools, xlsx exports (voter turnout, electoral system "
                                    "design, political finance)", "url": IDEA.format(theme="<theme>"),
                           "fetched": TODAY}}, series


def typed(cache: Path, iso3: set) -> tuple:
    series = {}
    for name, (steps, title, note) in TYPED.items():
        with (OUT / f"{name}.csv").open() as f:
            n = sum(1 for _ in f) - 1
        series[f"state/{name}"] = {"title": title, "rows": n, "for": steps, "note": note}
    return {}, series


SOURCES = {"doing_business": doing_business, "enterprise_surveys": enterprise_surveys, "oecd": oecd,
           "imf_cofog": imf_cofog, "wdi": wdi, "wipo": wipo, "vparty": vparty, "parlgov": parlgov, "idea": idea,
           "typed": typed}


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--cache", type=Path, default=None)
    parser.add_argument("--only", nargs="*", choices=sorted(SOURCES))
    args = parser.parse_args()
    cache = args.cache or Path(tempfile.mkdtemp())
    cache.mkdir(parents=True, exist_ok=True)
    iso3 = countries()
    for name in args.only or sorted(SOURCES):
        sources, series = SOURCES[name](cache, iso3)
        merge_manifest(sources, series)
        log(f"{name} done: " + ", ".join(f"{k} {v['rows']}" for k, v in series.items()))


if __name__ == "__main__":
    main()
