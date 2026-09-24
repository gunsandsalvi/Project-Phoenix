#!/usr/bin/env python3
"""Fetches the published data the opening population is derived from, into data/sources/raw/.

The sources surveyed in data/sources/notes/population-sources.md, each kept as a compact CSV of the rows the
derivation reads, with its release and download date in the manifest beside fetch.py's. The files are committed, so
tools/data/derive_pop.py reruns without the network.

    python3 tools/data/fetch_pop.py [--cache DIR] [--only NAME ...]

The downloads are kept in the cache directory (default: a temporary directory) and reused when present.
"""
import argparse
import csv
import datetime
import gzip
import io
import json
import tempfile
from pathlib import Path

from fetch import RAW, cached, get, log

SNAPSHOT = 2023
WPP = "https://population.un.org/wpp/assets/Excel%20Files/1_Indicator%20(Standard)/CSV_FILES/"
WPP_FILES = {
    "life_female": "WPP2024_Life_Table_Complete_Medium_Female_1950-2023.csv.gz",
    "life_male": "WPP2024_Life_Table_Complete_Medium_Male_1950-2023.csv.gz",
    "population": "WPP2024_PopulationBySingleAgeSex_Medium_1950-2023.csv.gz",
    "fertility": "WPP2024_Fertility_by_Age1.csv.gz",
    "indicators": "WPP2024_Demographic_Indicators_Medium.csv.gz",
}
FERTILE = range(15, 50)


def countries() -> set:
    return {r["iso3"] for r in csv.DictReader((RAW / "wb" / "countries.csv").open())}


def table(path: Path, header: list, rows: list) -> int:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="") as f:
        w = csv.writer(f)
        w.writerow(header)
        for r in sorted(rows):
            w.writerow(r)
    return len(rows)


def rows_of(path: Path):
    with gzip.open(path, "rt", encoding="utf-8-sig") as f:
        yield from csv.DictReader(f)


def age(r: dict) -> int:
    return int(r["AgeGrpStart"])


def wpp(cache: Path, manifest: dict, iso3: set) -> None:
    """The snapshot year's single-age life tables by sex and population by single age and sex; age-specific
    fertility for every year, for the children each cohort of mothers bore; under-five mortality for every year and
    the sex ratio at birth."""
    files = {k: cached(cache, v, WPP + v) for k, v in WPP_FILES.items()}
    life = []
    for sex in ("female", "male"):
        for r in rows_of(files[f"life_{sex}"]):
            if r["ISO3_code"] in iso3 and int(r["Time"]) == SNAPSHOT:
                life.append((r["ISO3_code"], sex, age(r), r["qx"], r["lx"]))
    manifest["series"]["wpp/life_table"] = {
        "title": f"Complete life tables by sex, {SNAPSHOT}: probability of dying between ages x and x+1 (qx) and "
                 "survivors of 100,000 born (lx), UN World Population Prospects 2024",
        "rows": table(RAW / "wpp" / "life_table.csv", ["iso3", "sex", "age", "qx", "lx"], life),
    }
    pop = [(r["ISO3_code"], age(r), r["PopMale"], r["PopFemale"]) for r in rows_of(files["population"])
           if r["ISO3_code"] in iso3 and int(r["Time"]) == SNAPSHOT]
    manifest["series"]["wpp/population_by_age"] = {
        "title": f"Population by single age and sex, 1 July {SNAPSHOT}, thousands, UN World Population Prospects 2024",
        "rows": table(RAW / "wpp" / "population_by_age.csv", ["iso3", "age", "male", "female"], pop),
    }
    asfr = {}
    for r in rows_of(files["fertility"]):
        if r["ISO3_code"] in iso3 and int(r["Time"]) <= SNAPSHOT and age(r) in FERTILE:
            asfr.setdefault((r["ISO3_code"], int(r["Time"])), {})[age(r)] = r["ASFR"]
    fert = [(iso, year, *(v[a] for a in FERTILE)) for (iso, year), v in asfr.items()]
    manifest["series"]["wpp/fertility_by_age"] = {
        "title": "Age-specific fertility rates by single age of mother 15-49, births per 1,000 women, 1950-"
                 f"{SNAPSHOT}, UN World Population Prospects 2024",
        "rows": table(RAW / "wpp" / "fertility_by_age.csv", ["iso3", "year", *(f"f{a}" for a in FERTILE)], fert),
    }
    ind = [(r["ISO3_code"], int(r["Time"]), r["Q5"], r["SRB"]) for r in rows_of(files["indicators"])
           if r["ISO3_code"] in iso3 and int(r["Time"]) <= SNAPSHOT]
    manifest["series"]["wpp/indicators"] = {
        "title": "Under-five deaths per 1,000 live births (Q5) and males per 100 females at birth (SRB), 1950-"
                 f"{SNAPSHOT}, UN World Population Prospects 2024",
        "rows": table(RAW / "wpp" / "indicators.csv", ["iso3", "year", "q5", "srb"], ind),
    }
    manifest["sources"]["wpp"] = {
        "title": "UN DESA, World Population Prospects 2024, standard projections (medium variant), CSV files",
        "url": WPP + "<file>",
    }


ILO = "https://sdmx.ilo.org/rest/data/ILO,{flow},1.0/all?startPeriod=2010"
ILO_CSV = "application/vnd.sdmx.data+csv;version=1.0.0"


def ilo_rows(cache: Path, flow: str):
    path = cache / f"{flow}.csv"
    if not path.exists():
        log(f"downloading {flow}")
        path.write_bytes(get(ILO.format(flow=flow), timeout=900, accept=ILO_CSV))
    with path.open(encoding="utf-8-sig") as f:
        yield from csv.DictReader(f)


def ilo_disability(cache: Path, manifest: dict, iso3: set) -> None:
    """Persons of working age and over by disability status, sex and age band, from labour force and household
    surveys; a year's rows come from one survey, named with them."""
    bands = {"AGE_AGGREGATE_Y15-24", "AGE_AGGREGATE_Y25-54", "AGE_AGGREGATE_Y55-64", "AGE_AGGREGATE_YGE65"}
    keep = {}
    for r in ilo_rows(cache, "DF_POP_XWAP_SEX_AGE_DSB_NB"):
        if r["FREQ"] == "A" and r["REF_AREA"] in iso3 and r["AGE"] in bands and r["DSB"] in ("DSB_STATUS_DIS", "DSB_STATUS_TOTAL") \
                and r["SEX"] in ("SEX_M", "SEX_F") and r["OBS_VALUE"]:
            key = (r["REF_AREA"], int(r["TIME_PERIOD"]), r["SEX"][4:], r["AGE"].split("_Y")[-1], r["SOURCE"])
            keep.setdefault(key, {})[r["DSB"][11:]] = r["OBS_VALUE"]
    rows = [(*k[:4], v["DIS"], v["TOTAL"], k[4]) for k, v in keep.items() if "DIS" in v and "TOTAL" in v]
    manifest["series"]["ilo/disability"] = {
        "title": "Population aged 15 and over with a disability and in total, thousands, by sex and age band "
                 "(DF_POP_XWAP_SEX_AGE_DSB_NB), ILOSTAT",
        "rows": table(RAW / "ilo" / "disability.csv",
                      ["iso3", "year", "sex", "age", "disabled", "total", "source"], rows),
    }
    manifest["sources"]["ilo_disability"] = {"title": "ILOSTAT SDMX API", "url": ILO.format(flow="<dataflow>")}


UN_HH = "https://population.un.org/household/assets/UNDESA_PD_2026_hh-size-composition.xlsx"
UN_OLDER = "https://population.un.org/LivingArrangements/assets/UNDESA_PD_2026_living-arrangements-older-persons.xlsx"
HH_COLUMNS = {
    "Average household size (number of members)": "mean_size",
    "1 member": "size_1",
    "2-3 members": "size_2_3",
    "4-5 members": "size_4_5",
    "6 or more members": "size_6_plus",
    "One-person": "one_person",
    "Couple only": "couple_only",
    "Couple with children": "couple_children",
    "Single parent with children": "single_parent",
    "Extended family": "extended",
    "Non-relatives": "non_relatives",
    "Unknown": "unknown",
    "Multi-generation": "multi_generation",
    "Three generation": "three_generation",
}
OLDER_COLUMNS = {
    "One person": "one_person",
    "Couple only": "couple_only",
    "With spouse or partner": "with_partner",
    "With children under age 20 years": "with_children_under_20",
    "With children aged 20 years or over": "with_children_20_plus",
}


def un_sheet(path: Path, sheet: str, columns: dict, keys: list) -> list:
    """The rows of a UN DESA household workbook, its header on the fifth row; each economy's rows keep their
    source, catalogue number and reference year, and '..' stays missing."""
    import openpyxl
    rows = list(openpyxl.load_workbook(path, read_only=True)[sheet].iter_rows(values_only=True))
    header = [str(h).strip() if h is not None else "" for h in rows[4]]
    at = {name: header.index(name) for name in ["ISO3 Code", "Data source category", "Data source catalog ID",
                                                 "Reference year", *keys, *columns]}
    out = []
    for r in rows[5:]:
        if r[at["ISO3 Code"]] is None:
            continue
        values = []
        for name in columns:
            v = r[at[name]]
            values.append("" if v in (None, "..", "") else str(v))
        if not any(values):
            continue
        out.append((str(r[at["ISO3 Code"]]), int(r[at["Reference year"]]), str(r[at["Data source category"]]),
                    str(r[at["Data source catalog ID"]]), *(str(r[at[k]]) for k in keys), *values))
    return out


def un_households(cache: Path, manifest: dict, iso3: set) -> None:
    """Households by size and basic type, and older persons by living arrangement, from censuses and surveys."""
    hh = [r for r in un_sheet(cached(cache, "un_hh.xlsx", UN_HH), "HH size and composition 2026", HH_COLUMNS, [])
          if r[0] in iso3]
    manifest["series"]["un/households"] = {
        "title": "Households by size and by basic and intergenerational type, % of households, and mean size, each "
                 "source's reference year, UN DESA Database on Household Size and Composition 2026",
        "rows": table(RAW / "un" / "households.csv",
                      ["iso3", "year", "source", "catalog", *HH_COLUMNS.values()], hh),
    }
    older = [r for r in un_sheet(cached(cache, "un_older.xlsx", UN_OLDER), "HHLA of Older Persons 2026",
                                 OLDER_COLUMNS, ["Age range", "Sex"]) if r[0] in iso3]
    manifest["series"]["un/older_persons"] = {
        "title": "Older persons (60+, 65+, 80+) by sex living alone, as a couple only, with a partner, with children "
                 "under 20 and with children 20 or over, % of older persons, UN DESA Database on the Households and "
                 "Living Arrangements of Older Persons 2026",
        "rows": table(RAW / "un" / "older_persons.csv",
                      ["iso3", "year", "source", "catalog", "ages", "sex", *OLDER_COLUMNS.values()], older),
    }
    manifest["sources"]["un_households"] = {
        "title": "UN DESA Population Division, Database on Household Size and Composition 2026 and Database on the "
                 "Households and Living Arrangements of Older Persons 2026",
        "url": f"{UN_HH}; {UN_OLDER}",
    }


WID_URL = "https://wid.world/bulk_download/wid_all_data.zip"
WID_SHARES = {"sptincj992": "income", "shwealj992": "wealth"}
WID_GROUPS = ["p0p50", "p50p90", "p90p100", "p99p100", "p99.9p100"]


def wid_shares(cache: Path, manifest: dict, iso3: set) -> None:
    """Shares of pre-tax national income and of net personal wealth held by the bottom half, the middle 40%, the top
    tenth, hundredth and thousandth, equal-split adults, 2010 on."""
    import zipfile
    z = zipfile.ZipFile(cached(cache, "wid_all_data.zip", WID_URL))
    two = {r["iso2"]: r["iso3"] for r in csv.DictReader((RAW / "wb" / "countries.csv").open()) if r["iso2"]}
    rows = []
    for info in z.infolist():
        name = info.filename
        if not (name.startswith("WID_data_") and name.endswith(".csv")):
            continue
        iso = two.get(name[len("WID_data_"):-len(".csv")])
        if iso not in iso3:
            continue
        for r in csv.DictReader(io.TextIOWrapper(z.open(info), encoding="utf-8"), delimiter=";"):
            if r["variable"] in WID_SHARES and r["percentile"] in WID_GROUPS and int(r["year"]) >= 2010 \
                    and r["value"]:
                rows.append((iso, int(r["year"]), WID_SHARES[r["variable"]], r["percentile"], r["value"]))
    manifest["series"]["wid/shares"] = {
        "title": "Shares of pre-tax national income (sptincj992) and net personal wealth (shwealj992), equal-split "
                 "adults, held by the bottom 50%, the middle 40%, the top 10%, 1% and 0.1%, 2010 on, WID",
        "rows": table(RAW / "wid" / "shares.csv", ["iso3", "year", "what", "group", "share"], rows),
    }
    release = max(i.date_time for i in z.infolist())
    manifest["sources"]["wid_shares"] = {
        "title": "World Inequality Database, bulk download", "url": WID_URL,
        "release": datetime.date(*release[:3]).isoformat(),
    }


def ilo_employment(cache: Path, manifest: dict, iso3: set) -> None:
    """Employed persons by sex and occupation (ISCO-08 major groups) and by status in employment (ICSE-93), from
    labour force surveys and censuses; the ILO's modelled estimates are left out."""
    specs = [
        ("DF_EMP_TEMP_SEX_OCU_NB", "OCU", "OCU_ISCO08_", "occupation",
         "Employed persons by sex and ISCO-08 major group, thousands (DF_EMP_TEMP_SEX_OCU_NB), ILOSTAT"),
        ("DF_EMP_TEMP_SEX_STE_NB", "STE", "STE_ICSE93_", "status",
         "Employed persons by sex and ICSE-93 status in employment, thousands (DF_EMP_TEMP_SEX_STE_NB), ILOSTAT"),
    ]
    for flow, dim, prefix, name, title in specs:
        rows = []
        for r in ilo_rows(cache, flow):
            if r["FREQ"] != "A" or r["REF_AREA"] not in iso3 or r["SEX"] not in ("SEX_M", "SEX_F") \
                    or not r[dim].startswith(prefix) or not r["OBS_VALUE"] or "Modelled" in r["SOURCE"]:
                continue
            rows.append((r["REF_AREA"], int(r["TIME_PERIOD"]), r["SEX"][4:], r[dim][len(prefix):], r["OBS_VALUE"],
                         r["SOURCE"]))
        manifest["series"][f"ilo/{name}"] = {
            "title": title,
            "rows": table(RAW / "ilo" / f"{name}.csv", ["iso3", "year", "sex", "class", "employed", "source"], rows),
        }
    manifest["sources"]["ilo_employment"] = {"title": "ILOSTAT SDMX API", "url": ILO.format(flow="<dataflow>")}


WCDE_URL = "https://wicshiny2023.iiasa.ac.at/wcde-data/wcde-v3-batch/2/prop.rds"
WCDE_YEAR = 2020
WCDE_LEVELS = ["No Education", "Incomplete Primary", "Primary", "Lower Secondary", "Upper Secondary",
               "Short Post Secondary", "Bachelor", "Master and higher"]


def wcde(cache: Path, manifest: dict, iso3: set) -> None:
    """Educational attainment by five-year age group and sex, the Wittgenstein Centre's reconstruction for 2020 (the
    last year before its projections begin), countries matched by their UN M49 codes through WPP's locations."""
    import pyreadr
    d = next(iter(pyreadr.read_r(str(cached(cache, "wcde_prop.rds", WCDE_URL))).values()))
    m49 = {}
    for r in rows_of(cached(cache, WPP_FILES["indicators"], WPP + WPP_FILES["indicators"])):
        if r["ISO3_code"]:
            m49[int(r["LocID"])] = r["ISO3_code"]
    d = d[(d.year == WCDE_YEAR) & d.sex.isin(["Male", "Female"]) & d.education.isin(WCDE_LEVELS)]
    rows = []
    for r in d.itertuples():
        iso = m49.get(int(r.country_code))
        if iso in iso3:
            rows.append((iso, r.sex[0], int(r.age.split("--")[0].rstrip("+")), WCDE_LEVELS.index(r.education),
                         f"{r.prop:.2f}"))
    manifest["series"]["wcde/attainment"] = {
        "title": f"Population by highest level of education (0 none, 1 incomplete primary, 2 primary, 3 lower "
                 f"secondary, 4 upper secondary, 5 short post-secondary, 6 bachelor, 7 master and higher), % of each "
                 f"five-year age group from 15 and sex, {WCDE_YEAR}, Wittgenstein Centre Human Capital Data Explorer "
                 f"v3, SSP2",
        "rows": table(RAW / "wcde" / "attainment.csv", ["iso3", "sex", "age", "level", "percent"], rows),
    }
    manifest["sources"]["wcde"] = {
        "title": "Wittgenstein Centre for Demography and Global Human Capital, Human Capital Data Explorer v3 (2023), "
                 "batch file",
        "url": WCDE_URL,
    }


# Names other publishers give economies, where the World Bank's differ: the only hand-kept mapping, each checked.
ALIASES = {
    "Korea": "KOR", "Republic of Korea": "KOR", "Slovakia": "SVK", "Turkey": "TUR", "Czech Republic": "CZE",
    "Bolivia (Plurinational State of)": "BOL", "Venezuela (Bolivarian Republic of)": "VEN", "Venezuela": "VEN",
    "Bahamas": "BHS", "Saint Kitts and Nevis": "KNA", "Saint Lucia": "LCA", "Saint Vincent and the Grenadines": "VCT",
    "Russia": "RUS", "Egypt": "EGY", "Iran": "IRN",
}
EUROSTAT_GEO = {"EL": "GRC", "UK": "GBR", "XK": "XKX"}


def iso3_by_name() -> dict:
    by_name = {r["name"]: r["iso3"] for r in csv.DictReader((RAW / "wb" / "countries.csv").open())}
    return {**by_name, **ALIASES}


AHD = "https://webfs.oecd.org/Els-com/Affordable_Housing_Database/"
AHD_TENURE = {"Own outright": "own_outright", "Owner with mortgage": "own_mortgage", "Rent (private)": "rent_private",
              "Rent (subsidised)": "rent_subsidised", "Other, unknown": "other"}
AHD_COST = {"Owner with mortgage": "mortgage_burden", "Rent (private and subsidized)": "rent_burden",
            "Rent (private and subsidised)": "rent_burden"}


def ahd_by_year(path: Path, sheet: str, labels: dict, scale: float, unmatched: set) -> list:
    """An Affordable Housing Database annex sheet: an economy's name, then one row per measure with a value per
    year from its header row; '..' and blanks are missing."""
    import openpyxl
    names = iso3_by_name()
    rows = list(openpyxl.load_workbook(path, read_only=True, data_only=True)[sheet].iter_rows(values_only=True))
    header = next(r for r in rows if sum(isinstance(c, int) and 2000 <= c <= 2100 for c in r) > 3)
    years = {i: c for i, c in enumerate(header) if isinstance(c, int) and 2000 <= c <= 2100}
    out, country = [], None
    for r in rows[rows.index(header) + 1:]:
        first = r[0].strip() if isinstance(r[0], str) else ""
        if first and first not in labels:
            country = names.get(first)
            if country is None and len(first) < 60:
                unmatched.add(first)
        label = next((c.strip() for c in r[:2] if isinstance(c, str) and c.strip() in labels), None)
        if label is None or country is None:
            continue
        for i, year in years.items():
            v = r[i] if i < len(r) else None
            if isinstance(v, (int, float)):
                out.append((country, year, labels[label], f"{v * scale:.6f}"))
    return out


def housing(cache: Path, manifest: dict, iso3: set) -> None:
    """Tenure and housing costs: the OECD Affordable Housing Database's tenure shares of households and median
    mortgage and rent burdens by year; Eurostat's tenure shares of persons; ECLAC's owners, tenants and other forms
    of tenancy among households; and the DHS surveys' women and men owning a house alone or jointly."""
    unmatched = set()
    tenure = ahd_by_year(cached(cache, "ahd_hm13.xlsx", AHD + "HM1-3-Housing-tenures.xlsx"), "HM1.3.A1",
                         AHD_TENURE, 0.01, unmatched)
    cost = ahd_by_year(cached(cache, "ahd_hc12.xlsx", AHD + "HC1-2-Housing-costs-over-income.xlsx"), "HC12_A1",
                       AHD_COST, 1.0, unmatched)
    manifest["series"]["oecd/ahd_tenure"] = {
        "title": "Households by tenure (own outright, owner with mortgage, rent private, rent subsidised, other), share "
                 "of households, by year, OECD Affordable Housing Database HM1.3.A1",
        "rows": table(RAW / "oecd" / "ahd_tenure.csv", ["iso3", "year", "tenure", "share"],
                      [r for r in tenure if r[0] in iso3]),
    }
    manifest["series"]["oecd/ahd_cost"] = {
        "title": "Median mortgage burden (principal and interest) of owners with a mortgage and rent burden of tenants, "
                 "share of disposable income, by year, OECD Affordable Housing Database HC1.2.A1",
        "rows": table(RAW / "oecd" / "ahd_cost.csv", ["iso3", "year", "measure", "share"],
                      [r for r in cost if r[0] in iso3]),
    }
    two = {r["iso2"]: r["iso3"] for r in csv.DictReader((RAW / "wb" / "countries.csv").open()) if r["iso2"]}
    url = ("https://ec.europa.eu/eurostat/api/dissemination/sdmx/2.1/data/ilc_lvho02?format=SDMX-CSV"
           "&startPeriod=2015")
    es = []
    with cached(cache, "eurostat_lvho02.csv", url).open(encoding="utf-8-sig") as f:
        for r in csv.DictReader(f):
            iso = EUROSTAT_GEO.get(r["geo"], two.get(r["geo"]))
            if r["rskpovth"] == "TOTAL" and r["hhcomp"] == "TOTAL" and r["unit"] == "PC" and iso in iso3 \
                    and r["OBS_VALUE"]:
                es.append((iso, int(r["TIME_PERIOD"]), r["tenure"], f"{float(r['OBS_VALUE']) / 100:.4f}"))
    manifest["series"]["eurostat/tenure"] = {
        "title": "Persons by tenure status of their household (OWN, OWN_L with a mortgage or loan, OWN_NL, RENT, "
                 "RENT_MKT, RENT_FR reduced or free), share of persons, Eurostat EU-SILC ilc_lvho02",
        "rows": table(RAW / "eurostat" / "tenure.csv", ["iso3", "year", "tenure", "share"], es),
    }
    cep = json.loads(cached(cache, "cepal166.json",
                            "https://api-cepalstat.cepal.org/cepalstat/api/v1/indicator/166/data?lang=en&format=json")
                     .read_text())["body"]
    dims = {d["id"]: {m["id"]: m["name"] for m in d["members"]} for d in cep["dimensions"]}
    names = iso3_by_name()
    tenure_of = {"Owner": "owner", "Tenat": "tenant", "Other forms of tenancy": "other"}
    cp = []
    for r in cep["data"]:
        if dims[326].get(r["dim_326"]) != "National" or dims[1412].get(r["dim_1412"]) not in tenure_of:
            continue
        name = dims[208].get(r["dim_208"])
        iso = names.get(name)
        if iso is None:
            unmatched.add(name)
            continue
        year = int(dims[29117][r["dim_29117"]])
        if iso in iso3 and year >= 2010 and r["value"] not in (None, ""):
            cp.append((iso, year, tenure_of[dims[1412][r["dim_1412"]]], f"{float(r['value']) / 100:.4f}"))
    manifest["series"]["cepalstat/tenure"] = {
        "title": "Households by tenure status of the dwelling (owner, tenant, other forms), national, share of "
                 "households, ECLAC CEPALSTAT indicator 166 (household surveys)",
        "rows": table(RAW / "cepalstat" / "tenure.csv", ["iso3", "year", "tenure", "share"], cp),
    }
    dhs_iso = {c["DHS_CountryCode"]: c["ISO3_CountryCode"] for c in json.loads(cached(
        cache, "dhs_countries.json", "https://api.dhsprogram.com/rest/dhs/countries?f=json").read_text())["Data"]}
    dhs = json.loads(cached(cache, "dhs_house.json",
                            "https://api.dhsprogram.com/rest/dhs/data?indicatorIds=WE_OWNA_W_HNO,WE_OWNA_W_HDK,"
                            "WE_OWNA_M_HNO,WE_OWNA_M_HDK&surveyYearStart=2010&breakdown=national&perpage=5000&f=json")
                     .read_text())["Data"]
    parts = {}
    for r in dhs:
        if r["IsPreferred"] and r["IsTotal"]:
            key = (dhs_iso.get(r["DHS_CountryCode"]), int(r["SurveyYear"]), "women" if "_W_" in r["IndicatorId"] else "men")
            parts.setdefault(key, {})[r["IndicatorId"][-3:]] = r["Value"]
    dh = [(*k, f"{(100 - v['HNO'] - v.get('HDK', 0)) / 100:.4f}") for k, v in parts.items() if "HNO" in v]
    manifest["series"]["dhs/house_owners"] = {
        "title": "Women and men aged 15-49 who own a house, alone or jointly: 100% less those who do not own one "
                 "and those who do not know, DHS Program API (WE_OWNA_*_HNO, WE_OWNA_*_HDK)",
        "rows": table(RAW / "dhs" / "house_owners.csv", ["iso3", "year", "sex", "share"],
                      [r for r in dh if r[0] in iso3]),
    }
    if unmatched:
        log(f"names matched to no economy (aggregates or not in the World Bank's list): {sorted(unmatched)}")
    manifest["sources"]["housing"] = {
        "title": "OECD Affordable Housing Database (HM1.3, HC1.2); Eurostat SDMX API (ilc_lvho02); ECLAC CEPALSTAT API "
                 "(indicator 166); DHS Program API",
        "url": f"{AHD}<file>; {url}; https://api-cepalstat.cepal.org/; https://api.dhsprogram.com/",
    }


FINDEX = {
    "account.t.d": "Account (% age 15+)",
    "fin17a": "Saved at a bank or similar financial institution (% age 15+)",
    "fin22a": "Borrowed from a formal bank or similar financial institution (% age 15+)",
}
FINDEX_URL = "https://api.worldbank.org/v2/country/all/indicator/{code}?format=json&date=2010:2025&per_page=20000&source=28"


def findex(cache: Path, manifest: dict, iso3: set) -> None:
    """Adults with an account, who saved at a bank and who borrowed from one, by survey wave."""
    for code, title in FINDEX.items():
        page = json.loads(get(FINDEX_URL.format(code=code)))
        rows = [(r["country"]["id"], int(r["date"]), f"{r['value'] / 100:.4f}") for r in page[1] or []
                if r["value"] is not None and r["country"]["id"] in iso3]
        manifest["series"][f"findex/{code}"] = {
            "title": f"{title}, World Bank Global Findex",
            "rows": table(RAW / "findex" / f"{code}.csv", ["iso3", "year", "share"], rows),
        }
    manifest["sources"]["findex"] = {"title": "World Bank Global Findex Database (API source 28)",
                                     "url": FINDEX_URL.format(code="<series>")}


PAG_URL = "https://sdmx.oecd.org/public/rest/data/OECD.ELS.SPD,DSD_PAG@DF_PAG,/all?format=csvfile"
PAG = {"CRPLF22": "pension_age", "GPRR100": "replacement_rate", "OCOP": "occupational_income", "PTOP": "public_income"}


def pensions(cache: Path, manifest: dict, iso3: set) -> None:
    """The OECD's pension ages and gross replacement rates of mandatory schemes for a worker entering at 22 on
    average earnings, by sex, and the shares of the over-65s' income from public and occupational transfers; the
    ILO's shares of persons above pensionable age receiving an old-age pension and of the severely disabled
    receiving a disability benefit; and pension funds' assets to GDP."""
    pag = []
    with cached(cache, "oecd_pag.csv", PAG_URL).open(encoding="utf-8-sig") as f:
        for r in csv.DictReader(f):
            if r["MEASURE"] in PAG and r["REF_AREA"] in iso3 and r["OBS_VALUE"] \
                    and r["OPTIONALITY"] in ("M", "_Z") and int(r["TIME_PERIOD"]) >= 2010:
                pag.append((r["REF_AREA"], int(r["TIME_PERIOD"]), PAG[r["MEASURE"]], r["SEX"], r["OBS_VALUE"]))
    manifest["series"]["oecd/pensions"] = {
        "title": "Current normal pension age for a worker entering at 22 (years, by sex), gross replacement rate of "
                 "mandatory schemes at average earnings (% of pre-retirement earnings, by sex), and public and "
                 "occupational transfers' shares of the disposable income of people over 65 (%), OECD Pensions at a "
                 "Glance (DF_PAG)",
        "rows": table(RAW / "oecd" / "pensions.csv", ["iso3", "year", "measure", "sex", "value"], pag),
    }
    cov = []
    for r in ilo_rows(cache, "DF_SDG_0131_SEX_SOC_RT"):
        if r["FREQ"] == "A" and r["REF_AREA"] in iso3 and r["SEX"] in ("SEX_M", "SEX_F", "SEX_T") \
                and r["SOC"] in ("SOC_CONTIG_PENSION", "SOC_CONTIG_DISAB") and r["OBS_VALUE"]:
            cov.append((r["REF_AREA"], int(r["TIME_PERIOD"]), r["SOC"][len("SOC_CONTIG_"):].lower(), r["SEX"][4:],
                        f"{float(r['OBS_VALUE']) / 100:.4f}", r["SOURCE"]))
    manifest["series"]["ilo/protection"] = {
        "title": "Persons above statutory pensionable age receiving an old-age pension, and persons with severe "
                 "disabilities receiving a disability benefit, share, by sex (SDG indicator 1.3.1, "
                 "DF_SDG_0131_SEX_SOC_RT), ILOSTAT",
        "rows": table(RAW / "ilo" / "protection.csv", ["iso3", "year", "function", "sex", "share", "source"], cov),
    }
    page = json.loads(get("https://api.worldbank.org/v2/country/all/indicator/GFDD.DI.13?format=json"
                          "&date=2010:2025&per_page=20000&source=32"))
    funds = [(r["countryiso3code"], int(r["date"]), r["value"]) for r in page[1] or []
             if r["value"] is not None and r["countryiso3code"] in iso3]
    manifest["series"]["wb/GFDD.DI.13"] = {
        "title": "Pension fund assets to GDP (%), World Bank Global Financial Development Database",
        "rows": table(RAW / "wb" / "GFDD.DI.13.csv", ["iso3", "year", "value"], funds),
    }
    manifest["sources"]["pensions"] = {
        "title": "OECD SDMX API, Pensions at a Glance (DF_PAG); ILOSTAT SDMX API (SDG 1.3.1); World Bank GFDD",
        "url": f"{PAG_URL}; {ILO.format(flow='DF_SDG_0131_SEX_SOC_RT')}",
    }


SOURCES = {"pensions": pensions, "housing": housing, "findex": findex, "wpp": wpp, "ilo_disability": ilo_disability, "un_households": un_households, "wid_shares": wid_shares,
           "ilo_employment": ilo_employment, "wcde": wcde}


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--cache", type=Path, default=None)
    parser.add_argument("--only", nargs="*", choices=sorted(SOURCES))
    args = parser.parse_args()
    cache = args.cache or Path(tempfile.mkdtemp())
    cache.mkdir(parents=True, exist_ok=True)
    manifest = json.loads((RAW / "manifest.json").read_text())
    iso3 = countries()
    for name in args.only or sorted(SOURCES):
        SOURCES[name](cache, manifest, iso3)
        manifest["sources"][name]["fetched"] = datetime.date.today().isoformat()
        log(f"{name} done")
    (RAW / "manifest.json").write_text(json.dumps(manifest, indent=2, sort_keys=True) + "\n")


if __name__ == "__main__":
    main()
