#!/usr/bin/env python3
"""Fetches the published data on people — labour, demography, preferences, expectations and migration — that the
remaining build steps read, into data/sources/raw/. Years 2015 to 2025 unless a table is a snapshot, which is the
survey year nearest 2019 at or before it (else the first after), one survey per economy.

- ILOSTAT (SDMX API), from surveys and administrative records: mean monthly earnings of employees by ISCO-08 major
  group and by education, local currency; collective bargaining coverage and trade union density; employment by
  public/private sector, and by ISIC Rev. 4 section and sector (snapshot); employment by ISIC Rev. 4 division
  (snapshot); fatal and non-fatal occupational injuries per 100 000 workers by section, and days lost per case;
  unemployment by duration.
- OECD (SDMX API): employment in general government as a share of employment (Government at a Glance 2025); trade
  union density and collective bargaining coverage (OECD/AIAS); sick leave of full-time employees, weeks a year.
- World Bank API: enrolment ratios, school entry ages and durations, pupil-teacher ratios, harmonised test scores,
  personal remittances, migrant stock and net migration.
- The Global Preferences Survey's country means (Falk et al. 2018), from its site's rankings service.
- Households' expectations: the University of Michigan Surveys of Consumers by age and income group (yearly
  distributions); the New York Fed's Survey of Consumer Expectations chart data by demographic group (monthly), and
  its public microdata reduced to the expectation and group columns (2015-2019); the ECB's Consumer Expectations
  Survey by age and income quintile; and a derived per-level table, the US pattern shifted to each level's inflation.
- UN DESA International Migrant Stock 2024 by destination and origin, and the World Bank/KNOMAD bilateral remittance
  matrix 2021, both from Our World in Data's catalogue (the primary hosts are unreachable).
- The OECD Family Database's crude marriage and divorce rates (via Our World in Data's catalogue); DHS household
  roster, marital status and ideal-number-of-children aggregates; the OECD Survey of Adult Skills 2023 country means
  and proficiency shares (via the Irish CSO's PxStat, which republishes the OECD's table).
- Tables typed from papers: Montenegro and Patrinos (2014), returns to schooling by economy (Annex Table 1); Bernard,
  Bell and Cooper (2018), five-year internal migration intensities by education (Tables A1, A4.1, A4.2).
- Donovan, Lu and Schoellman (2023) replication data: wage workers' job tenure distribution and quarterly transitions.

    python3 tools/data/fetch_people.py [--cache DIR] [--only NAME ...]
"""
import argparse
import csv
import datetime
import gzip
import io
import json
import re
import statistics
import tempfile
import urllib.parse
import urllib.request
from pathlib import Path

from fetch import RAW, cached, get, log, merge_manifest
from fetch_pop import countries, table

FIRST, LAST = 2015, 2025
SNAPSHOT = 2019
TODAY = datetime.date.today().isoformat()

ILO = "https://sdmx.ilo.org/rest/data/ILO,{flow},1.0/{key}?startPeriod={first}&endPeriod={last}"
ILO_CSV = "application/vnd.sdmx.data+csv;version=1.0.0"
OECD = "https://sdmx.oecd.org/public/rest/data/{flow}/{key}?startPeriod={first}&format=csvfile"
WB_API = "https://api.worldbank.org/v2/country/all/indicator/{code}?format=json&date={first}:{last}&per_page=20000"
GPS = "https://gps.econ.uni-bonn.de/"
SCA = "https://data.sca.isr.umich.edu/subset/output.php"
SCE = "https://www.newyorkfed.org/medialibrary/interactives/sce/sce/downloads/data/"
CES = ("https://data-api.ecb.europa.eu/service/data/CES/M...T.C1120+C1220+C3220+C4031+C1150_EXP+C1150_UNCERT."
       "NUM_VAR.WM+WA+WP25+WP75?format=csvdata")
OWID = "https://catalog.ourworldindata.org/"
DHS = "https://api.dhsprogram.com/rest/dhs/"
CSO = "https://ws.cso.ie/public/api.restful/PxStat.Data.Cube_API.ReadDataset/{table}/CSV/1.0/en"
RETURNS = "https://documents1.worldbank.org/curated/en/830831468147839247/pdf/WPS7020.pdf"
INTERNAL = "https://arxiv.org/pdf/1812.08913"
DLS = "https://dataverse.harvard.edu/api/access/datafile/7102919"


def levels() -> dict:
    """Each economy's development level: the World Bank income groups, lower-middle and low together."""
    level = {"High income": "developed", "Upper middle income": "emerging", "Lower middle income": "developing",
             "Low income": "developing"}
    return {r["iso3"]: level[r["income_group"]] for r in csv.DictReader((RAW / "wb" / "countries.csv").open())}


def iso3_of_iso2() -> dict:
    return {r["iso2"]: r["iso3"] for r in csv.DictReader((RAW / "wb" / "countries.csv").open()) if r["iso2"]}


def snapshot(rows: list) -> list:
    """The rows of each economy's survey nearest the snapshot year: the latest at or before it, else the first after;
    within that year the source reporting the most rows."""
    by = {}
    for r in rows:
        by.setdefault(r["iso3"], []).append(r)
    out = []
    for iso, rs in by.items():
        years = {r["year"] for r in rs}
        before = [y for y in years if y <= SNAPSHOT]
        year = max(before) if before else min(years)
        count = {}
        for r in rs:
            if r["year"] == year:
                count[r["source"]] = count.get(r["source"], 0) + 1
        source = min(count, key=lambda s: (-count[s], s))
        out += [r for r in rs if r["year"] == year and r["source"] == source]
    return out


def one_source(rows: list) -> list:
    """Each economy-year's rows from one source, the one reporting the most rows, so a year is never counted twice;
    the source is kept as its type (LFS, ADM-IR, ...), ILOSTAT's source name before its dash."""
    count = {}
    for r in rows:
        key = (r["iso3"], r["year"])
        count.setdefault(key, {})
        count[key][r["source"]] = count[key].get(r["source"], 0) + 1
    best = {k: min(c, key=lambda s: (-c[s], s)) for k, c in count.items()}
    return [dict(r, source=r["source"].split(" - ")[0]) for r in rows if r["source"] == best[(r["iso3"], r["year"])]]


# ILOSTAT --------------------------------------------------------------------------------------------------------

def ilo_rows(cache: Path, flow: str, key: str, iso3: set) -> list:
    path = cache / f"{flow}_{re.sub(r'[^A-Za-z0-9]', '_', key)}.csv"
    if not path.exists():
        log(f"downloading {flow}")
        path.write_bytes(get(ILO.format(flow=flow, key=key, first=FIRST, last=LAST), timeout=1800, accept=ILO_CSV))
    with path.open(encoding="utf-8-sig") as f:
        return [dict(r, iso3=r["REF_AREA"], year=int(r["TIME_PERIOD"][:4]), source=r["SOURCE"])
                for r in csv.DictReader(f)
                if r["FREQ"] == "A" and r["REF_AREA"] in iso3 and r["OBS_VALUE"] and "Modelled" not in r["SOURCE"]]


def ilo(cache: Path, iso3: set) -> tuple:
    series = {}
    ilo_dir = RAW / "ilo"

    def put(name, title, header, rows, steps):
        series[f"ilo/{name}"] = {"title": title, "rows": table(ilo_dir / f"{name}.csv", header, rows), "for": steps}

    rows = ilo_rows(cache, "DF_EAR_EMTA_SEX_OCU_CUR_NB", ".A..SEX_T..CUR_TYPE_LCU", iso3)
    put("earnings_by_occupation",
        "Mean monthly earnings of employees by ISCO-08 major group (0-9, X not elsewhere classified, TOTAL), local "
        "currency, both sexes, 2015-2025 (DF_EAR_EMTA_SEX_OCU_CUR_NB), ILOSTAT",
        ["iso3", "year", "occupation", "value", "source"],
        [(r["iso3"], r["year"], r["OCU"][len("OCU_ISCO08_"):], r["OBS_VALUE"], r["source"])
         for r in rows if r["OCU"].startswith("OCU_ISCO08_")], ["S1.22", "S6.02"])

    rows = ilo_rows(cache, "DF_EAR_EMTA_SEX_EDU_NB", ".A..SEX_T.", iso3)
    put("earnings_by_education",
        "Mean monthly earnings of employees by education (aggregate levels: less than basic, basic, intermediate, "
        "advanced, not stated, total), local currency, both sexes, 2015-2025 (DF_EAR_EMTA_SEX_EDU_NB), ILOSTAT",
        ["iso3", "year", "education", "value", "source"],
        [(r["iso3"], r["year"], r["EDU"][len("EDU_AGGREGATE_"):], r["OBS_VALUE"], r["source"])
         for r in rows if r["EDU"].startswith("EDU_AGGREGATE_")], ["S6.114"])

    union = []
    for flow, measure in (("DF_ILR_CBCT_NOC_RT", "coverage"), ("DF_ILR_TUMT_NOC_RT", "density")):
        union += [(r["iso3"], r["year"], measure, r["OBS_VALUE"], r["source"])
                  for r in ilo_rows(cache, flow, "all", iso3)]
    put("unions", "Collective bargaining coverage rate and trade union density rate, % of employees, 2015-2025 "
        "(DF_ILR_CBCT_NOC_RT, DF_ILR_TUMT_NOC_RT), ILOSTAT", ["iso3", "year", "measure", "value", "source"], union,
        ["S2.214"])

    rows = ilo_rows(cache, "DF_EMP_TEMP_SEX_INS_NB", ".A..SEX_T.", iso3)
    put("employment_by_sector",
        "Employment by institutional sector (PUB public, PRI private, TOTAL), thousands, both sexes, 2015-2025 "
        "(DF_EMP_TEMP_SEX_INS_NB), ILOSTAT", ["iso3", "year", "sector", "employed", "source"],
        [(r["iso3"], r["year"], r["INS"][len("INS_SECTOR_"):], r["OBS_VALUE"], r["source"]) for r in rows],
        ["S5.122"])

    rows = snapshot([r for r in ilo_rows(cache, "DF_EMP_TEMP_SEX_ECO_INS_NB", ".A..SEX_T..", iso3)
                     if r["ECO"].startswith("ECO_ISIC4_")])
    put("employment_by_activity_sector",
        "Employment by ISIC Rev. 4 section and institutional sector (PUB, PRI, TOTAL), thousands, both sexes, one "
        "survey per economy nearest 2019 (DF_EMP_TEMP_SEX_ECO_INS_NB), ILOSTAT",
        ["iso3", "year", "activity", "sector", "employed", "source"],
        [(r["iso3"], r["year"], r["ECO"][len("ECO_ISIC4_"):], r["INS"][len("INS_SECTOR_"):], r["OBS_VALUE"],
          r["source"]) for r in rows], ["S5.122"])

    rows = snapshot([r for r in ilo_rows(cache, "DF_EMP_TEMP_SEX_EC2_NB", ".A..SEX_T.", iso3)
                     if r["EC2"].startswith("EC2_ISIC4_")])
    put("employment_by_division",
        "Employment by ISIC Rev. 4 division (two digits; transport 49-53 among them), thousands, both sexes, one "
        "survey per economy nearest 2019 (DF_EMP_TEMP_SEX_EC2_NB), ILOSTAT",
        ["iso3", "year", "division", "employed", "source"],
        [(r["iso3"], r["year"], r["EC2"][len("EC2_ISIC4_"):], r["OBS_VALUE"], r["source"]) for r in rows],
        ["S1.22", "S2.05"])

    injuries = []
    for flow, measure in (("DF_INJ_FATL_ECO_RT", "fatal"), ("DF_INJ_NFTL_ECO_RT", "nonfatal")):
        injuries += [(r["iso3"], r["year"], measure, r["ECO"][len("ECO_ISIC4_"):], r["OBS_VALUE"], r["source"])
                     for r in one_source([r for r in ilo_rows(cache, flow, "all", iso3)
                                          if r["ECO"].startswith("ECO_ISIC4_")])]
    put("injuries", "Occupational injuries per 100 000 workers in the reference group, fatal and non-fatal, by ISIC "
        "Rev. 4 section (TOTAL all activities), 2015-2025, one source per economy-year (DF_INJ_FATL_ECO_RT, "
        "DF_INJ_NFTL_ECO_RT), ILOSTAT",
        ["iso3", "year", "measure", "activity", "rate", "source"], injuries, ["S4.134"])

    cases = {(r["iso3"], r["year"], r["source"]): r["OBS_VALUE"]
             for r in ilo_rows(cache, "DF_INJ_NFTL_ECO_NB", "all", iso3) if r["ECO"] == "ECO_ISIC4_TOTAL"}
    days = [(r["iso3"], r["year"], r["cases"], r["OBS_VALUE"], r["source"])
            for r in one_source([dict(r, cases=cases[(r["iso3"], r["year"], r["source"])])
                                 for r in ilo_rows(cache, "DF_INJ_DAYS_ECO_NB", "all", iso3)
                                 if r["ECO"] == "ECO_ISIC4_TOTAL" and (r["iso3"], r["year"], r["source"]) in cases])]
    put("injury_days", "Cases of non-fatal occupational injury and days lost to those with temporary incapacity, all "
        "activities, 2015-2025 (DF_INJ_NFTL_ECO_NB, DF_INJ_DAYS_ECO_NB), ILOSTAT",
        ["iso3", "year", "cases", "days_lost", "source"], days, ["S4.134"])

    rows = one_source(ilo_rows(cache, "DF_UNE_TUNE_SEX_AGE_DUR_NB", ".A..SEX_T.AGE_YTHADULT_YGE15.", iso3))
    put("unemployment_by_duration",
        "Unemployed persons aged 15 and over by duration of unemployment (DETAILS_*: bands in months, MLT1 under "
        "1, MGE1LT3 1 to 3, ... MGE24 24 or more; AGGREGATE_*: under 6, 6-12, 12 or more; X not stated; TOTAL), "
        "thousands, both sexes, 2015-2025, one source per economy-year "
        "(DF_UNE_TUNE_SEX_AGE_DUR_NB), ILOSTAT", ["iso3", "year", "duration", "unemployed", "source"],
        [(r["iso3"], r["year"], r["DUR"][len("DUR_"):], r["OBS_VALUE"], r["source"]) for r in rows],
        ["S1.16", "S7.01"])

    sources = {"people_ilo": {"title": "ILOSTAT SDMX API", "url": ILO.format(flow="<dataflow>", key="<key>",
                                                                           first=FIRST, last=LAST), "fetched": TODAY}}
    return sources, series


# OECD -----------------------------------------------------------------------------------------------------------

def oecd_rows(cache: Path, name: str, flow: str, key: str, first: int = FIRST) -> list:
    path = cache / f"oecd_{name}.csv"
    if not path.exists():
        log(f"downloading OECD {flow}")
        path.write_bytes(get(OECD.format(flow=flow, key=key, first=first), timeout=900))
    with path.open(encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


def oecd(cache: Path, iso3: set) -> tuple:
    series = {}
    rows = [(r["REF_AREA"], int(r["TIME_PERIOD"]), r["OBS_VALUE"])
            for r in oecd_rows(cache, "gov_emp", "OECD.GOV.GIP,DSD_GOV@DF_GOV_EMPPS_REP_2025,", "all", 2000)
            if r["MEASURE"] == "EMPG" and r["UNIT_MEASURE"] == "PT_EMP" and r["REF_AREA"] in iso3 and r["OBS_VALUE"]]
    series["oecd/government_employment"] = {
        "title": "Employment in general government, % of total employment, all years published, Government at a "
                 "Glance 2025 (DSD_GOV@DF_GOV_EMPPS_REP_2025, EMPG), OECD",
        "rows": table(RAW / "oecd" / "government_employment.csv", ["iso3", "year", "share"], rows),
        "for": ["S5.122"]}
    union = []
    for flow in ("DF_TUD", "DF_CBC"):
        union += [(r["REF_AREA"], int(r["TIME_PERIOD"]), r["MEASURE"], r["OBS_VALUE"])
                  for r in oecd_rows(cache, flow, f"OECD.ELS.SAE,DSD_TUD_CBC@{flow},", "all")
                  if r["REF_AREA"] in iso3 and r["OBS_VALUE"]]
    series["oecd/unions"] = {
        "title": "Trade union density (TUD) and collective bargaining coverage (measures of DF_CBC), % of employees, "
                 "2015-2025, OECD/AIAS ICTWSS (DSD_TUD_CBC), OECD",
        "rows": table(RAW / "oecd" / "unions.csv", ["iso3", "year", "measure", "value"], union), "for": ["S1.15"]}
    rows = [(r["REF_AREA"], int(r["TIME_PERIOD"]), r["OBS_VALUE"])
            for r in oecd_rows(cache, "sick_leave", "OECD.ELS.HD,DSD_HEALTH_STAT@DF_HEALTH_STATUS,",
                               ".A.WKAB..........")
            if r["MEASURE"] == "WKAB" and r["REF_AREA"] in iso3 and r["OBS_VALUE"]]
    series["oecd/sick_leave"] = {
        "title": "Sick leave of full-time dependent employees, weeks a year per person, 2015-2025 "
                 "(DSD_HEALTH_STAT@DF_HEALTH_STATUS, WKAB), OECD",
        "rows": table(RAW / "oecd" / "sick_leave.csv", ["iso3", "year", "weeks"], rows), "for": ["S6.02"]}
    sources = {"people_oecd": {"title": "OECD Data Explorer SDMX API: Government at a Glance, ICTWSS, health status",
                               "url": OECD.format(flow="<dataflow>", key="<key>", first=FIRST), "fetched": TODAY}}
    return sources, series


# World Bank -----------------------------------------------------------------------------------------------------

WDI = {
    "SE.PRE.ENRR": (["S6.127"], "School enrolment, pre-primary, % gross"),
    "SE.PRM.ENRR": (["S6.02"], "School enrolment, primary, % gross"),
    "SE.SEC.ENRR": (["S6.02"], "School enrolment, secondary, % gross"),
    "SE.TER.ENRR": (["S6.02"], "School enrolment, tertiary, % gross"),
    "SE.COM.DURS": (["S6.02"], "Compulsory education, duration, years"),
    "SE.PRE.DURS": (["S6.114"], "Pre-primary education, duration, years"),
    "SE.PRM.DURS": (["S6.02"], "Primary education, duration, years"),
    "SE.SEC.DURS": (["S6.02"], "Secondary education, duration, years"),
    "SE.PRM.AGES": (["S6.02"], "Primary school starting age, years"),
    "SE.SEC.AGES": (["S6.02"], "Lower secondary school starting age, years"),
    "SE.PRE.ENRL.TC.ZS": (["S5.122"], "Pupil-teacher ratio, pre-primary"),
    "SE.PRM.ENRL.TC.ZS": (["S5.122"], "Pupil-teacher ratio, primary"),
    "SE.SEC.ENRL.TC.ZS": (["S5.122"], "Pupil-teacher ratio, secondary"),
    "SE.TER.ENRL.TC.ZS": (["S5.122"], "Pupil-teacher ratio, tertiary"),
    "HD.HCI.HLOS": (["S6.114"], "Harmonized test scores (Human Capital Index), 300 minimal to 625 advanced"),
    "BX.TRF.PWKR.DT.GD.ZS": (["S5.100"], "Personal remittances, received, % of GDP"),
    "BX.TRF.PWKR.CD.DT": (["S5.100"], "Personal remittances, received, current US$"),
    "BM.TRF.PWKR.CD.DT": (["S5.100"], "Personal remittances, paid, current US$"),
    "SM.POP.TOTL.ZS": (["S5.100"], "International migrant stock, % of population"),
    "SM.POP.NETM": (["S5.100", "S6.100"], "Net migration, persons"),
}


def wdi(cache: Path, iso3: set) -> tuple:
    series = {}
    for code, (steps, title) in WDI.items():
        page = json.loads(get(WB_API.format(code=code, first=FIRST, last=LAST)))
        rows = [(r["countryiso3code"], int(r["date"]), r["value"]) for r in page[1] or []
                if r["value"] is not None and r["countryiso3code"] in iso3]
        series[f"wdi/{code}"] = {"title": f"{title}, 2015-2025 ({code}), World Bank WDI",
                                 "rows": table(RAW / "wdi" / f"{code}.csv", ["iso3", "year", "value"], rows),
                                 "for": steps}
    sources = {"people_wdi": {"title": "World Bank API, World Development Indicators",
                              "url": WB_API.format(code="<series>", first=FIRST, last=LAST), "fetched": TODAY}}
    return sources, series


# Preferences ----------------------------------------------------------------------------------------------------

GPS_TRAITS = ["patience", "risk_taking", "positive_reciprocity", "negative_reciprocity", "altruism", "trust"]


def gps(cache: Path, iso3: set) -> tuple:
    """The rankings service answers for a list of the survey's own country ids; asking for more ids than exist
    returns every country."""
    ids = "&".join(f"gallop_id%5B%5D={i}" for i in range(1, 200))
    data = json.loads(cached(cache, "gps_values.json", f"{GPS}values?{ids}").read_text())
    of = iso3_of_iso2()
    rows = [(of[r["iso_3166-2"]], *(r[t] for t in GPS_TRAITS)) for r in data if of.get(r["iso_3166-2"]) in iso3]
    missing = sorted(r["country_short_name"] for r in data if of.get(r["iso_3166-2"]) not in iso3)
    if missing:
        log(f"GPS countries outside the economies kept: {missing}")
    series = {"people/gps_country": {
        "title": "Global Preferences Survey country means of patience, risk taking, positive and negative reciprocity, "
                 "altruism and trust, in standard deviations of the world's individual-level distribution (mean 0), "
                 "surveyed 2012, Falk et al. (2018, QJE 133(4)), from the survey site's rankings service",
        "rows": table(RAW / "people" / "gps_country.csv", ["iso3", *GPS_TRAITS], rows),
        "for": ["S1.12", "S2.05", "VAL.22"]}}
    sources = {"people_gps": {"title": "Global Preferences Survey, University of Bonn (country rankings service)",
                              "url": f"{GPS}rankings", "fetched": TODAY}}
    return sources, series


# Expectations ---------------------------------------------------------------------------------------------------

SCA_GROUPS = ["all", "a1834", "a3544", "a4554", "a5564", "a6597", "y15", "y25", "y35", "y45", "y55"]
SCA_FORM = {"freq": "year", "start_year": FIRST, "end_year": LAST, "alldemos": "on", "all": "on", "age_18": "on",
            "age_35": "on", "age_45": "on", "age_55": "on", "age_65": "on", "inc_1": "on", "inc_2": "on",
            "inc_3": "on", "inc_4": "on", "inc_5": "on", "px1": "on", "px5": "on", "umex": "on", "inex": "on",
            "pjob": "on", "scores": "all"}

# The chart-data sheets kept and, for a sheet by demographic group, what each of its column blocks measures (the
# block titles are partly missing in the file; the blocks follow the order of the all-households sheet).
SCE_SHEETS = {
    "Inflation expectations": ("infl", None),
    "Inflation expectations Demo": ("infl_demo", ["median one-year-ahead expected inflation, %",
                                                  "median three-year-ahead expected inflation, %"]),
    "Inflation uncertainty": ("unc", None),
    "Inflation uncertainty Demo": ("unc_demo", ["median one-year-ahead inflation uncertainty (interquartile range of "
                                                "the density), percentage points",
                                                "median three-year-ahead inflation uncertainty, percentage points"]),
    "Earnings growth": ("earn", None),
    "Earnings growth Demo": ("earn_demo", ["median expected earnings growth one year ahead, %"]),
    "HH Income Change": ("inc", None),
    "HH Income Change Demo": ("inc_demo", ["median expected household income growth one year ahead, %"]),
    "Unemployment Expectations": ("unemp", None),
    "Unemployment Expectations Demo": ("unemp_demo", ["mean probability that US unemployment will be higher in a "
                                                      "year, %"]),
    "Job separation expectation": ("sep", None),
    "Job separation expectation Demo": ("sep_demo", ["mean probability of losing one's job in 12 months, %",
                                                     "mean probability of leaving one's job voluntarily in 12 "
                                                     "months, %"]),
    "Job finding expectations": ("find", None),
    "Job finding expectations Demo": ("find_demo", ["mean probability of finding a job within 3 months if one lost "
                                                    "it today, %"]),
}
SCE_MICRO = ["frbny-sce-public-microdata-complete-13-16.xlsx", "frbny-sce-public-microdata-complete-17-19.xlsx"]
# Point forecasts of inflation one and three years ahead, the density's median and interquartile range one year
# ahead, the probabilities of higher US unemployment and of losing one's job, and the groups.
SCE_COLUMNS = {"Q8v2part2": "infl_1y", "Q9bv2part2": "infl_3y", "Q9_cent50": "infl_1y_density_median",
               "Q9_iqr": "infl_1y_density_iqr", "Q4new": "p_unemployment_higher", "Q13new": "p_lose_job",
               "_AGE_CAT": "age", "_HH_INC_CAT": "income", "_EDU_CAT": "education"}


def sca(cache: Path) -> list:
    path = cache / "sca_subset.csv"
    if not path.exists():
        log("downloading Surveys of Consumers subset")
        body = urllib.parse.urlencode(SCA_FORM).encode()
        request = urllib.request.Request(SCA, data=body, headers={"User-Agent": "phoenix-data-fetch/1"})
        with urllib.request.urlopen(request, timeout=600) as r:
            path.write_bytes(r.read())
    out = []
    for r in csv.DictReader(path.open(encoding="utf-8-sig")):
        for column, value in r.items():
            if column == "yyyy" or not value.strip():
                continue
            variable, rest = column.split("_", 1)
            stat, group = rest.rsplit("_", 1)
            out.append((int(r["yyyy"]), variable, stat, group, value.strip()))
    return out


def sce_chart(cache: Path) -> tuple:
    """The kept sheets' monthly values by measure code and group, and the codes' legend."""
    import openpyxl
    book = openpyxl.load_workbook(cached(cache, "frbny-sce-data.xlsx", SCE + "frbny-sce-data.xlsx"), read_only=True)
    out, legend = [], {}
    for sheet, (code, blocks) in SCE_SHEETS.items():
        rows = list(book[sheet].iter_rows(values_only=True))
        names, block, gap, column = {}, 0, False, 0
        for i, h in enumerate(rows[3][1:], start=1):
            if not isinstance(h, str) or not h.strip():
                gap = bool(names)
                continue
            if gap:
                block, gap = block + 1, False
            if blocks is None:
                column += 1
                measure, group, meaning = f"{code}.{column}", "all", h.strip()
            else:
                measure, group, meaning = f"{code}.{block + 1}", h.strip(), blocks[block]
            names[i] = (measure, group)
            legend[measure] = (sheet, meaning)
        for r in rows[4:]:
            if not isinstance(r[0], (int, float)):
                continue
            month = int(r[0])
            if not FIRST <= month // 100 <= LAST:
                continue
            for i, (measure, group) in names.items():
                if i < len(r) and isinstance(r[i], (int, float)):
                    out.append((f"{month // 100}-{month % 100:02d}", measure, group, f"{r[i]:.3f}"))
    return out, [(m, s, d) for m, (s, d) in legend.items()]


def sce_micro(cache: Path) -> list:
    import openpyxl
    out = []
    for name in SCE_MICRO:
        book = openpyxl.load_workbook(cached(cache, name, SCE + name), read_only=True)
        rows = book.worksheets[0].iter_rows(values_only=True)
        next(rows)
        header = list(next(rows))
        at = {c: header.index(c) for c in ["date", "userid", "tenure", "weight", *SCE_COLUMNS]}
        for r in rows:
            if r[at["date"]] is None or not FIRST <= int(r[at["date"]]) // 100 <= SNAPSHOT:
                continue
            value = lambda c: "" if r[at[c]] is None else (f"{r[at[c]]:.4g}" if isinstance(r[at[c]], float)
                                                            else str(r[at[c]]))
            out.append((int(r[at["date"]]), int(r[at["userid"]]), value("tenure"), value("weight"),
                        *(value(c) for c in SCE_COLUMNS)))
    return out


def ces(cache: Path) -> list:
    text = cached(cache, "ces.csv", CES).read_text(encoding="utf-8-sig")
    of = {**iso3_of_iso2(), "GR": "GRC"}
    return [(of.get(r["REF_AREA"], r["REF_AREA"]), r["TIME_PERIOD"], r["CES_VARIABLE"], r["CES_DENOM"], r["CES_BREAKDOWN"], r["OBS_VALUE"])
            for r in csv.DictReader(io.StringIO(text)) if r["OBS_VALUE"] and r["TIME_PERIOD"][:4] <= str(LAST)]


def expectations(cache: Path, iso3: set) -> tuple:
    series = {}
    people = RAW / "people"
    sca_rows = sca(cache)
    series["people/sca_expectations"] = {
        "title": "University of Michigan Surveys of Consumers, yearly 2015-2025, by group (all households; age 18-34, "
                 "35-44, 45-54, 55-64, 65+; income quintiles y15..y55): expected price change next year (px1) and in "
                 "5 years (px5), % — shares by band, mean, median, p25, p75, std, variance; expected unemployment "
                 "change (umex); expected income change next year (inex); probability of losing a job in 5 years "
                 "(pjob)",
        "rows": table(people / "sca_expectations.csv", ["year", "variable", "stat", "group", "value"], sca_rows),
        "for": ["S1.15"]}
    chart, legend = sce_chart(cache)
    series["people/sce_expectations"] = {
        "title": "New York Fed Survey of Consumer Expectations chart data, monthly 2015-2025: inflation expectations "
                 "and uncertainty one and three years ahead, earnings and household income growth, probabilities of "
                 "higher unemployment, losing, leaving and finding a job, all households (medians and percentiles) "
                 "and by age, education, income, numeracy and region; measure codes in sce_measures.csv",
        "rows": table(people / "sce_expectations.csv", ["month", "measure", "group", "value"], chart),
        "for": ["S1.15"]}
    series["people/sce_measures"] = {
        "title": "Legend of sce_expectations' measure codes: the chart-data sheet and what the code measures (a "
                 "sheet by demographic group has one code per column block, its blocks in the order of the "
                 "all-households sheet where the file leaves a block untitled)",
        "rows": table(people / "sce_measures.csv", ["measure", "sheet", "meaning"], legend), "for": ["S1.15"]}
    micro = sce_micro(cache)
    path = people / "sce_microdata.csv.gz"
    with gzip.open(path, "wt", newline="") as f:
        w = csv.writer(f)
        w.writerow(["month", "userid", "tenure", "weight", *SCE_COLUMNS.values()])
        w.writerows(sorted(micro))
    series["people/sce_microdata"] = {
        "title": "New York Fed Survey of Consumer Expectations public microdata, 2015-2019, one row per respondent "
                 "and month (a rotating panel of up to 12 months, userid): point forecasts of inflation one and three "
                 "years ahead, the one-year density's median and interquartile range, probabilities of higher US "
                 "unemployment and of losing one's job (%), survey weight, age, household income and education "
                 "groups (gzip CSV)",
        "rows": len(micro), "for": ["S1.455"]}
    series["people/ces_expectations"] = {
        "title": "ECB Consumer Expectations Survey, monthly from 2020-04: inflation expectations 12 months (C1120) and "
                 "3 years ahead (C1220), probabilistic 12-month mean and uncertainty (C1150_EXP, C1150_UNCERT), "
                 "household income growth expected (C3220), expected unemployment rate (C4031); weighted mean (WA), "
                 "median (WM), p25, p75; euro area (Z17, Z18: 17 and 18 members) and 11 countries (iso3), all and by age band and income "
                 "quintile",
        "rows": table(people / "ces_expectations.csv", ["area", "month", "variable", "stat", "group", "value"],
                      ces(cache)), "for": ["S1.15"]}

    # The US pattern by group, shifted to each level's inflation.
    level = levels()
    cpi = {}
    for r in csv.DictReader((RAW / "wdi" / "FP.CPI.TOTL.ZG.csv").open()):
        if FIRST <= int(r["year"]) <= SNAPSHOT:
            cpi.setdefault(r["iso3"], []).append(float(r["value"]))
    mean_cpi = {iso: statistics.mean(v) for iso, v in cpi.items() if len(v) == SNAPSHOT - FIRST + 1}
    us = mean_cpi["USA"]
    shift = {lv: statistics.median(v for iso, v in mean_cpi.items() if level.get(iso) == lv) - us
             for lv in ("developed", "emerging", "developing")}
    pattern = {}
    for year, variable, stat, group, value in sca_rows:
        if variable == "px1" and stat in ("mean", "med", "p25", "p75", "std") and year <= SNAPSHOT:
            pattern.setdefault((group, stat), []).append(float(value))
    derived = []
    for (group, stat), values in sorted(pattern.items()):
        us_value = statistics.mean(values)
        for lv, s in shift.items():
            derived.append((lv, group, stat, f"{us_value + (0 if stat == 'std' else s):.3f}", f"{us_value:.3f}",
                            f"{s:.3f}"))
    series["people/expectations_by_level"] = {
        "title": "Households' one-year-ahead expected inflation by age group and income quintile, per development "
                 "level (mean, median, p25, p75 and standard deviation, % a year), derived",
        "note": "Derived: the US pattern shifted to each level's inflation. For each Surveys of Consumers group "
                "(sca_expectations: px1, age 18-34..65+, income quintiles, all) the 2015-2019 mean of each yearly "
                "statistic is the US value; the level's shift is the median over its economies (World Bank income "
                "groups: High = developed, Upper middle = emerging, Lower middle and Low = developing) of each "
                "economy's 2015-2019 mean CPI inflation (wdi/FP.CPI.TOTL.ZG, economies with all five years) less the "
                "US's; mean, median, p25 and p75 are shifted by it, the standard deviation is kept",
        "rows": table(people / "expectations_by_level.csv", ["level", "group", "stat", "value", "us_value", "shift"],
                      derived), "for": ["S1.456"]}
    sources = {
        "people_sca": {"title": "University of Michigan, Surveys of Consumers, demographic subgroups subset", "url": SCA,
                       "fetched": TODAY},
        "people_sce": {"title": "Federal Reserve Bank of New York, Survey of Consumer Expectations: chart data and "
                                "public microdata", "url": SCE + "<file>", "fetched": TODAY},
        "people_ces": {"title": "ECB Data Portal API, Consumer Expectations Survey (CES)", "url": CES, "fetched": TODAY},
    }
    return sources, series


# Our World in Data's catalogue ---------------------------------------------------------------------------------

OWID_TABLES = {
    "migrant_stock": "garden/un/2025-03-12/migrant_stock/migrant_stock_dest_origin",
    "remittances": "garden/wb/2024-12-17/bilateral_remittance/bilateral_remittance",
    "marriage": "garden/oecd/2025-10-07/family_database/marriage_divorce_rates",
    "regions": "garden/regions/2023-01-01/regions/regions",
}
# Names the catalogue gives economies the World Bank lists under other codes.
OWID_ALIASES = {"Channel Islands": "CHI", "Kosovo": "XKX", "Saint Martin (French part)": "MAF", "World": "WLD"}


def owid(cache: Path, name: str):
    import pandas
    return pandas.read_parquet(cached(cache, f"owid_{name}.parquet", OWID + OWID_TABLES[name] + ".parquet"))


def present(v) -> bool:
    import pandas
    return not pandas.isna(v)


def owid_iso3(cache: Path) -> dict:
    regions = owid(cache, "regions").reset_index()
    return {**{n: c for n, c in zip(regions["name"], regions["code"]) if isinstance(c, str) and len(c) == 3},
            **OWID_ALIASES}


def migration(cache: Path, iso3: set) -> tuple:
    of = owid_iso3(cache)
    keep = iso3 | {"WLD"}
    stock = owid(cache, "migrant_stock")
    rows = [(of[d], of[o], int(y), int(n)) for d, o, y, n in zip(stock["country_destination"], stock["country_origin"],
                                                                 stock["year"], stock["migrants_all_sexes"])
            if of.get(d) in keep and of.get(o) in keep and int(y) in (2015, 2020, 2024) and present(n) and n > 0]
    remit = owid(cache, "remittances")
    rrows = [(of[s], of[r], int(y), int(v)) for s, r, y, v in zip(remit["country_origin"], remit["country_receiving"],
                                                                  remit["year"], remit["remittance_flows"])
             if of.get(s) in keep and of.get(r) in keep and present(v) and v > 0]
    series = {
        "people/migrant_stock_bilateral": {
            "title": "International migrant stock by country of destination and origin (WLD = all), persons, mid-2015, "
                     "2020 and 2024, UN DESA International Migrant Stock 2024, via Our World in Data's catalogue",
            "rows": table(RAW / "people" / "migrant_stock_bilateral.csv", ["destination", "origin", "year", "migrants"],
                          rows), "for": ["S5.100", "S5.174"]},
        "people/remittances_bilateral": {
            "title": "Bilateral remittance flows, sending economy to receiving economy (WLD = all), current US$, 2021, "
                     "World Bank/KNOMAD bilateral remittance matrix (estimated from migrant stocks and incomes), via "
                     "Our World in Data's catalogue",
            "rows": table(RAW / "people" / "remittances_bilateral.csv", ["sender", "receiver", "year", "usd"], rrows),
            "for": ["S5.100", "S5.174"]},
    }
    sources = {"people_owid": {"title": "Our World in Data ETL catalogue (garden tables: UN DESA migrant stock 2024, "
                                        "KNOMAD bilateral remittances, OECD Family Database, regions)",
                               "url": OWID + "<path>.parquet", "fetched": TODAY}}
    return sources, series


# Demography ----------------------------------------------------------------------------------------------------

DHS_INDICATORS = {
    # Household roster aggregates.
    "HC_MEMB_H_MNM": "mean household members", "HC_MEMB_H_1MM": "households of 1 member, %",
    "HC_MEMB_H_2MM": "households of 2, %", "HC_MEMB_H_3MM": "households of 3, %", "HC_MEMB_H_4MM": "households of 4, %",
    "HC_MEMB_H_5MM": "households of 5, %", "HC_MEMB_H_6MM": "households of 6, %", "HC_MEMB_H_7MM": "households of 7, %",
    "HC_MEMB_H_8MM": "households of 8, %", "HC_MEMB_H_9MM": "households of 9+, %",
    "HC_OLDR_H_W65": "households with members 65+, %", "HC_OLDR_H_O65": "households of members 65+ only, %",
    "HC_OLDR_H_3GN": "households of 3 generations, %", "HC_OLDR_P_NNC": "population in non-nuclear families, %",
    "HC_HHHD_H_FEM": "female-headed households, %", "HC_LVAR_C_BTH": "children under 18 living with both parents, %",
    "HC_LVAR_C_NBP": "children under 18 living with no biological parent, %",
    "HC_ORPH_H_FOS": "households with children not living with a biological parent, %",
    # Marital status of women and men 15-49.
    **{f"MA_MSTA_{s}_{c}": f"{'women' if s == 'W' else 'men'} 15-49 {label}, %" for s in "WM"
       for c, label in (("NMA", "never married"), ("MAR", "married"), ("LTG", "living together"),
                        ("DIV", "divorced"), ("SEP", "separated"), ("WID", "widowed"))},
    # Ideal number of children.
    "PR_IDLC_W_MNA": "mean ideal number of children, all women", "PR_IDLC_M_MNA": "mean ideal number of children, "
                                                                                  "all men",
    **{f"PR_IDLC_W_ID{k}": f"women whose ideal number of children is {k if k != '6' else '6+'}, %"
       for k in "0123456"},
    "PR_IDLC_W_IDN": "women giving a non-numeric ideal number of children, %",
}


def demography(cache: Path, iso3: set) -> tuple:
    series = {}
    of = owid_iso3(cache)
    fd = owid(cache, "marriage")
    rows = [(of[c], int(y), i, f"{v:.3f}") for c, y, g, i, v in
            zip(fd["country"], fd["year"], fd["gender"], fd["indicator"], fd["value"])
            if of.get(c) in iso3 and g == "Both" and present(v) and FIRST <= int(y) <= LAST]
    series["people/marriage_divorce"] = {
        "title": "Crude marriage and divorce rates, per 1 000 people, 2015-2025, OECD Family Database (SF3.1), via Our "
                 "World in Data's catalogue",
        "rows": table(RAW / "people" / "marriage_divorce.csv", ["iso3", "year", "indicator", "per_1000"], rows),
        "for": ["S6.114"]}

    dhs_iso = {c["DHS_CountryCode"]: c["ISO3_CountryCode"] for c in json.loads(cached(
        cache, "dhs_countries.json", DHS + "countries?f=json").read_text())["Data"]}
    ids = ",".join(DHS_INDICATORS)
    url = f"{DHS}data?indicatorIds={ids}&surveyYearStart=2005&breakdown=national&perpage=5000&f=json&page="
    first = json.loads(cached(cache, "dhs_people_1.json", url + "1").read_text())
    data = first["Data"]
    for page in range(2, first["TotalPages"] + 1):
        data += json.loads(cached(cache, f"dhs_people_{page}.json", url + str(page)).read_text())["Data"]
    rows = [(dhs_iso[r["DHS_CountryCode"]], int(r["SurveyYear"]), r["SurveyId"], r["IndicatorId"], r["Value"])
            for r in data if r["IsTotal"] == 1 and dhs_iso.get(r["DHS_CountryCode"]) in iso3 and r["Value"] is not None]
    series["people/dhs_indicators"] = {
        "title": "DHS survey aggregates, national, surveys from 2005: household size distribution and composition "
                 "(members 65+, three generations, non-nuclear families, female head, children's living "
                 "arrangements), current marital status of women and men 15-49, ideal number of children; "
                 "indicators: " + "; ".join(f"{k} {v}" for k, v in DHS_INDICATORS.items()) + "; DHS Program API",
        "rows": table(RAW / "people" / "dhs_indicators.csv", ["iso3", "year", "survey", "indicator", "value"], rows),
        "for": ["S1.23", "S1.403"]}

    piaac = []
    of2 = {**iso3_of_iso2(), "BE-VLG": "BEL", "GB-ENG": "GBR"}
    for t in ("PIAAC01", "PIAAC02"):
        for r in csv.reader(io.StringIO(cached(cache, f"cso_{t}.csv", CSO.format(table=t)).read_text(
                encoding="utf-8-sig"))):
            if r[0] == "STATISTIC" or of2.get(r[4]) not in iso3 or not r[7]:
                continue
            piaac.append((of2[r[4]], r[5], " ".join(r[1].split()), r[6], r[7]))
    series["people/piaac_2023"] = {
        "title": "Survey of Adult Skills (PIAAC Cycle 2, 2023), adults 16-65: mean literacy, numeracy and adaptive "
                 "problem-solving scores (0-500) and shares at proficiency levels, by participating economy (Belgium is "
                 "Flanders, the United Kingdom England), OECD results republished by the Irish CSO (PxStat PIAAC01, "
                 "PIAAC02)",
        "rows": table(RAW / "people" / "piaac_2023.csv", ["iso3", "area", "statistic", "unit", "value"], piaac),
        "for": ["S6.114"]}

    tenure = []
    for r in csv.DictReader(io.StringIO(cached(cache, "dls_tenure.tab", DLS).read_text()), delimiter="\t"):
        code = {"WBG": "PSE"}.get(r["ccode"], r["ccode"])
        if code in iso3:
            tenure.append((code, int(float(r["year"])), int(r["tenure4"]), r["t1"], r["yshare"], r["ytrate"]))
    series["people/tenure_dls"] = {
        "title": "Wage workers by job tenure band (1 under 6 months, 2 6-12 months, 3 1-5 years, 4 over 5 years): the "
                 "band's share of wage workers (share) and the quarterly rate of moving to state to_state (W wage work "
                 "in the same job, B and M the replication's other states), by economy and survey year, harmonised "
                 "rotating-panel labour force surveys, Donovan, Lu and Schoellman (2023, QJE), replication data "
                 "doi:10.7910/DVN/RXWKTV (Tenure_Transitions)",
        "rows": table(RAW / "people" / "tenure_dls.csv", ["iso3", "year", "tenure_band", "to_state", "share", "rate"],
                      tenure), "for": ["S1.411"]}
    sources = {
        "people_dhs": {"title": "DHS Program API", "url": DHS + "data?indicatorIds=<ids>", "fetched": TODAY},
        "people_piaac": {"title": "CSO Ireland PxStat, PIAAC 2023 international comparison tables",
                         "url": CSO.format(table="<table>"), "fetched": TODAY},
        "people_dls": {"title": "Harvard Dataverse, Donovan, Lu and Schoellman (2023) replication package",
                       "url": DLS, "fetched": TODAY},
    }
    return sources, series


# Tables typed from papers --------------------------------------------------------------------------------------

def returns_rows(path: Path, names: dict) -> list:
    """Annex Table 1 of the paper, read by column position: a cell left blank in the paper stays missing."""
    import pymupdf
    doc = pymupdf.open(path)
    out, unmatched = [], set()
    for page in doc:
        words = page.get_text("words")
        head = [w for w in words if len(w[4]) == 1 and w[4] in "ABCDEFGHIJK" and w[1] < 130]
        if len(head) != 11:
            continue
        cols = {w[4]: (w[0] + w[2]) / 2 for w in head}
        top = head[0][1]
        lines = {}
        for w in words:
            if w[1] <= top + 5:
                continue
            y = round((w[1] + w[3]) / 2)
            key = next((k for k in lines if abs(k - y) <= 3), y)
            lines.setdefault(key, []).append(w)
        for y in sorted(lines):
            ws = sorted(lines[y], key=lambda w: w[0])
            year = next((w for w in ws if re.fullmatch(r"(19|20)\d\d", w[4])), None)
            if year is None:
                continue
            name = " ".join(w[4] for w in ws if w[2] <= year[0] + 1)
            cells = {}
            for w in ws:
                if w[0] > year[2]:
                    cells[min(cols, key=lambda c: abs(cols[c] - (w[0] + w[2]) / 2))] = w[4]
            iso = names.get(name)
            if iso is None:
                unmatched.add(name)
                continue
            out.append((iso, int(year[4]), *(cells.get(c, "") for c in "ACDE")))
    if unmatched:
        log(f"returns to schooling, economies not kept: {sorted(unmatched)}")
    return out


INTERNAL_REGIONS = {"Africa", "Asia", "Europe", "Latin America and the Caribbean", "North America", "North",
                    "America"}


def internal_rows(path: Path, names: dict) -> list:
    """Tables A4.1 and A4.2 (intensities by education) with Table A1's census year."""
    import pymupdf
    doc = pymupdf.open(path)
    text = {i: doc[i].get_text() for i in range(doc.page_count)}
    a1 = next(t for t in text.values() if "Table A1 Migration data" in t)
    census = {}
    lines = [l.strip() for l in a1.splitlines() if l.strip()]
    for i, l in enumerate(lines):
        if re.fullmatch(r"(19|20)\d\d", l):
            name = lines[i - 1]
            if lines[i - 2] in ("Trinidad", "and") or name == "Tobago":
                name = "Trinidad and Tobago"
            census[" ".join(name.split())] = int(l)
    out, unmatched = [], set()
    for t in text.values():
        if not t.lstrip().split("\n", 2)[-1].lstrip().startswith("Table A4."):
            continue
        body = [l.strip() for l in t.split("University", 1)[1].splitlines() if l.strip()]
        name, pending = None, []
        i = 1
        while i < len(body):
            l = body[i]
            if l in ("major", "minor"):
                values = body[i + 1:i + 6]
                country = " ".join(pending) if pending else name
                name, pending = country, []
                iso = names.get(country)
                if iso is None:
                    unmatched.add(country)
                else:
                    out.append((iso, census.get(country, ""), l, *values))
                i += 6
                continue
            if l.startswith("Note") or re.fullmatch(r"\d+", l):
                i += 1
                continue
            if l not in INTERNAL_REGIONS:
                pending.append(l)
            i += 1
    if unmatched:
        log(f"internal migration, economies not kept: {sorted(unmatched)}")
    return out


def typed(cache: Path, iso3: set) -> tuple:
    names = {**{r["name"]: r["iso3"] for r in csv.DictReader((RAW / "wb" / "countries.csv").open())},
             **{n: c for n, c in owid_iso3(cache).items() if c in iso3},
             "Macedonia, FYR": "MKD", "Swaziland": "SWZ", "Czech Republic": "CZE", "Kyrgyz Republic": "KGZ",
             "São Tomé and Principe": "STP", "Côte d'Ivoire": "CIV", "Venezuela, RB": "VEN", "Turkey": "TUR",
             "Congo, Dem. Rep.": "COD", "Gambia, The": "GMB", "Yemen, Rep.": "YEM", "Korea, Rep.": "KOR",
             "Lao PDR": "LAO", "Russian Federation": "RUS", "Syrian Arab Republic": "SYR", "Egypt": "EGY",
             "Iran": "IRN", "Vietnam": "VNM", "Kyrgyzstan": "KGZ", "United States": "USA"}
    returns = [r for r in returns_rows(cached(cache, "WPS7020.pdf", RETURNS), names) if r[0] in iso3]
    internal = [r for r in internal_rows(cached(cache, "bernard_bell_cooper_2018.pdf", INTERNAL), names)
                if r[0] in iso3]
    series = {
        "people/typed_returns_to_schooling": {
            "title": "Private returns to schooling by economy and survey year, % per year: another year of schooling "
                     "(A, educyT) and a year of primary, secondary and tertiary schooling (C, D, E: educyL_P_T, "
                     "educyL_S_T, educyL_T_T), both sexes, Mincerian estimates on harmonised household surveys "
                     "1970-2013",
            "note": "Typed from Montenegro, C.E. and H.A. Patrinos (2014), 'Comparable Estimates of Returns to "
                    "Schooling Around the World', World Bank Policy Research Working Paper 7020, Annex Table 1, "
                    "pp. 20-36 (columns A, C, D, E); read from the downloaded PDF by column position, a blank cell "
                    "left missing",
            "rows": table(RAW / "people" / "typed_returns_to_schooling.csv",
                          ["iso3", "year", "return_per_year", "primary", "secondary", "tertiary"], returns),
            "for": ["S6.114"]},
        "people/typed_internal_migration": {
            "title": "Five-year crude internal migration intensity, % of the population aged 15 and over who changed "
                     "major (or minor) administrative region over five years, in total and by education (less than "
                     "primary, primary, secondary, university completed), census year",
            "note": "Typed from Bernard, A., M. Bell and J. Cooper (2018), 'Internal migration and education: a "
                    "cross-national comparison', background paper for the 2019 Global Education Monitoring Report "
                    "(UNESCO), arXiv:1812.08913, Tables A4.1 and A4.2 (pp. 53-54) with the census years of Table A1 "
                    "(p. 49); the IMAGE repository and Bell et al. (2015) are unreachable",
            "rows": table(RAW / "people" / "typed_internal_migration.csv",
                          ["iso3", "census_year", "regions", "total", "less_than_primary", "primary", "secondary",
                           "university"], internal),
            "for": ["S6.100"]},
    }
    sources = {"people_typed": {"title": "World Bank Documents (WPS7020); arXiv (1812.08913)",
                                "url": f"{RETURNS}; {INTERNAL}", "fetched": TODAY}}
    return sources, series


SOURCES = {"ilo": ilo, "oecd": oecd, "wdi": wdi, "gps": gps, "expectations": expectations, "migration": migration,
           "demography": demography, "typed": typed}


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--cache", type=Path, default=None)
    parser.add_argument("--only", nargs="*", choices=sorted(SOURCES))
    args = parser.parse_args()
    cache = args.cache or Path(tempfile.mkdtemp())
    cache.mkdir(parents=True, exist_ok=True)
    (RAW / "people").mkdir(exist_ok=True)
    iso3 = countries()
    for name in args.only or list(SOURCES):
        sources, series = SOURCES[name](cache, iso3)
        merge_manifest(sources, series)
        for key, s in series.items():
            log(f"{key}: {s['rows']} rows")


if __name__ == "__main__":
    main()
