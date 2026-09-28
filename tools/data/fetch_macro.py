#!/usr/bin/env python3
"""Fetches the macro series the realism reads judge the world against, per development level, into
data/sources/raw/macro/, and derives from them each measurable fact's range per level (level_ranges.csv).

- Penn World Table 11.0 (DataverseNL), every economy, 1950-2023: real GDP (national accounts and expenditure side),
  population, employment, hours, human capital, the labour share and the expenditure shares.
- Quarterly national accounts: the OECD's (chain-linked volumes: GDP, household consumption, fixed investment,
  household consumption by durability, employment) and the IMF's QNEA (volumes: GDP, private consumption, fixed
  investment); one source per economy and series, seasonally adjusted where published.
- Quarterly labour series: the ILO's survey unemployment rate and employment (national estimates, not modelled),
  the OECD's seasonally adjusted unemployment rate, job vacancies and hourly earnings.
- Quarterly prices and rates: the IMF's consumer price index and interest rates (policy, lending, Treasury bill),
  the OECD's three-month and long-term rates.
- Annual series: the World Bank's WDI (real GDP, household consumption and fixed investment, inflation, the
  financial and current accounts in dollars, GDP in dollars), the ILO's survey unemployment
  rate, the OECD's average annual wages and hours.
- The Jordà-Schularick-Taylor Macrohistory Database, release 6 (developed economies, 1870-2020).
- BIS US-dollar exchange rates (WS_XRU), monthly, end of period, 2000 on.
- Typed paper tables: Aguiar and Gopinath (NBER WP 10734) Tables 2A-2C, Neumeyer and Perri (NBER WP 10387) Tables
  1A-1B, Calderón and Fuentes (World Bank PRWP 5343) Table 1.

    python3 tools/data/fetch_macro.py [--cache DIR] [--only NAME ...]

`ranges` recomputes level_ranges.csv from the files on disk and needs no network.
"""
import argparse
import bisect
import csv
import datetime
import io
import json
import math
import statistics
import tempfile
import zipfile
from pathlib import Path

import openpyxl

from fetch import RAW, get, log, cached, merge_manifest
from fetch_pop import countries, table

OUT = RAW / "macro"
TODAY = datetime.date.today().isoformat()
UA_CSV = "application/vnd.sdmx.data+csv;version=1.0.0"
# Statistics end before the pandemic, as the opening's reference year does.
LAST_YEAR = 2019

PWT_URL = "https://dataverse.nl/api/access/datafile/554105"
PWT_COLUMNS = ["rgdpe", "rgdpna", "pop", "emp", "avh", "hc", "labsh", "csh_c", "csh_i", "csh_g"]
JST_URL = "https://www.macrohistory.net/app/download/9834512569/JSTdatasetR6.xlsx"
JST_COLUMNS = ["pop", "rgdpmad", "rgdpbarro", "rconsbarro", "gdp", "iy", "cpi", "ca", "money", "stir", "ltrate",
               "unemp", "wage", "xrusd", "tloans", "tmort", "thh", "tbus", "debtgdp", "hpnom", "crisisJST", "peg"]
XRU_URL = "https://data.bis.org/static/bulk/WS_XRU_csv_flat.zip"
XRU_FIRST = 2000
RATES_FIRST = 1970
IMF = "https://api.imf.org/external/sdmx/2.1/data/IMF.STA,{flow}/{key}?detail=dataonly"
OECD = "https://sdmx.oecd.org/public/rest/data/{flow}/{key}?format=csvfile"
ILO = "https://sdmx.ilo.org/rest/data/ILO,{flow},1.0/{key}?detail=dataonly"
WB = "https://api.worldbank.org/v2/country/all/indicator/{code}?format=json&date=1960:2024&per_page=20000&page={page}"

# The IMF's QNEA indicator codes and the names kept.
QNEA = {"B1GQ": "gdp", "P3_PS": "consumption_private", "P51G": "investment"}
# The OECD's quarterly accounts: (sector, transaction) and the names kept.
QNA = {("S1", "B1GQ"): "gdp", ("S1M", "P3"): "consumption_private", ("S1", "P51G"): "investment",
       ("S14", "P311"): "consumption_durable", ("S14", "P312"): "consumption_semidurable",
       ("S14", "P313"): "consumption_nondurable", ("S14", "P314"): "consumption_services", ("S1", "EMP"): "employment"}
QNA_FLOWS = {
    "OECD.SDD.NAD,DSD_NAMAIN1@DF_QNA_EXPENDITURE_NATIO_CURR": "Q.Y+N..S1+S1M..B1GQ+P3+P51G.....L..",
    "OECD.SDD.NAD,DSD_NAMAIN1@DF_QNA_EXPENDITURE_DURABILITY": "Q.Y+N..S14..P311+P312+P313+P314.....L..",
    "OECD.SDD.NAD,DSD_NAMAIN1@DF_QNA_POP_EMPNC": "Q.Y+N..S1..EMP.......",
}
MFS_IR = {"MFS166_RT_PT_A_PT": "policy_rate", "MFS162_RT_PT_A_PT": "lending_rate",
          "GSTBILY_RT_PT_A_PT": "tbill_rate"}
FINMARK = {"IR3TIB": "short_rate_3m", "IRLT": "long_rate"}
WB_SERIES = {
    "NY.GDP.MKTP.KD": "gdp_real", "NE.CON.PRVT.KD": "consumption_household_real", "NE.GDI.FTOT.KD": "investment_real",
    "FP.CPI.TOTL.ZG": "cpi_inflation", "BN.FIN.TOTL.CD": "financial_account_usd",
    "BN.CAB.XOKA.CD": "current_account_usd",
    "NY.GDP.MKTP.CD": "gdp_usd", "FS.AST.PRVT.GD.ZS": "private_credit_pct_gdp",
}
# Codes the sources use for economies whose ISO code differs.
ALIAS = {"KOS": "XKX"}


def levels() -> dict:
    """Each economy's development level: developed, emerging or developing, from its World Bank income group."""
    group = {"High income": "developed", "Upper middle income": "emerging", "Lower middle income": "developing",
             "Low income": "developing"}
    return {r["iso3"]: group[r["income_group"]] for r in csv.DictReader((RAW / "wb" / "countries.csv").open())}


def sdmx_rows(text: str):
    yield from csv.DictReader(io.StringIO(text.lstrip("﻿")))


def g6(v) -> str:
    return f"{float(v):.6g}"


# ---------------------------------------------------------------------------------------------------- fetchers


def pwt(cache: Path, manifest: dict, iso3: set) -> None:
    book = openpyxl.load_workbook(cached(cache, "pwt110.xlsx", PWT_URL), read_only=True)
    rows = book["Data"].iter_rows(values_only=True)
    header = list(next(rows))
    at = [header.index(c) for c in PWT_COLUMNS]
    out = []
    for r in rows:
        if r[0] in iso3 and r[3] is not None:
            values = ["" if r[i] is None else g6(r[i]) for i in at]
            if any(values):
                out.append((r[0], int(r[3]), *values))
    manifest["series"]["macro/pwt"] = {
        "title": "Penn World Table 11.0, every economy, 1950-2023: real GDP (rgdpe at chained PPPs; rgdpna at "
                 "constant national prices, for growth), population and persons engaged (millions), average annual "
                 "hours (avh), human capital (hc), labour share (labsh), household, investment and government shares "
                 "of GDP at current PPPs (csh_c, csh_i, csh_g)",
        "rows": table(OUT / "pwt.csv", ["iso3", "year", *PWT_COLUMNS], out),
        "for": ["S1.16", "S7.01"],
    }
    manifest["sources"]["macro/pwt"] = {"title": "Penn World Table 11.0 (Feenstra, Inklaar and Timmer), DataverseNL, "
                                        "doi:10.34894/FABVLR, pwt110.xlsx", "url": PWT_URL, "fetched": TODAY,
                                        "release": "11.0"}


def qna(cache: Path, manifest: dict, iso3: set) -> None:
    """Quarterly volumes: the OECD's chain-linked series and the IMF's QNEA, seasonally adjusted where published."""
    found = {}
    for flow, key in QNA_FLOWS.items():
        name = flow.split("@")[1]
        text = cached(cache, f"oecd_{name}.csv", OECD.format(flow=flow, key=key)).read_text(encoding="utf-8-sig")
        for r in sdmx_rows(text):
            series = QNA.get((r["SECTOR"], r["TRANSACTION"]))
            if series and r["REF_AREA"] in iso3 and r["OBS_VALUE"] and r["FREQ"] == "Q":
                if r["TRANSACTION"] == "EMP" and r["UNIT_MEASURE"] != "PS":
                    continue
                adj = "SA" if r["ADJUSTMENT"] == "Y" else "NSA"
                found.setdefault(("OECD", r["REF_AREA"], series, adj), {})[r["TIME_PERIOD"]] = r["OBS_VALUE"]
    url = IMF.format(flow="QNEA", key=".B1GQ+P3_PS+P51G.Q.SA+NSA.XDC.Q")
    text = _imf(cache, "imf_qnea.csv", url)
    for r in sdmx_rows(text):
        iso = ALIAS.get(r["COUNTRY"], r["COUNTRY"])
        if iso in iso3 and r["OBS_VALUE"]:
            key = ("IMF", iso, QNEA[r["INDICATOR"]], r["S_ADJUSTMENT"])
            found.setdefault(key, {})[r["TIME_PERIOD"]] = r["OBS_VALUE"]
    # One source per economy and series: adjusted before unadjusted, then the longer run to 2019, the OECD on ties.
    best = {}
    for (source, iso, series, adj), obs in found.items():
        span = runs({qindex(q): 0 for q in obs}, last=LAST_YEAR * 4 + 3)
        rank = (adj == "SA", len(span), source == "OECD")
        if (iso, series) not in best or rank > best[(iso, series)][0]:
            best[(iso, series)] = (rank, source, adj, obs)
    out = [(iso, source, series, adj, q, g6(v)) for (iso, series), (_, source, adj, obs) in best.items()
           for q, v in obs.items()]
    manifest["series"]["macro/qna"] = {
        "title": "Quarterly national accounts in volumes (national currency; chain-linked volumes for the OECD, "
                 "constant prices for the IMF): GDP, private (households and NPISH) final consumption, "
                 "gross fixed capital formation, household consumption of durable, semi-durable and non-durable goods "
                 "and services (OECD), and total employment in persons (OECD, national concept); OECD Quarterly "
                 "National Accounts and IMF Quarterly National Accounts (QNEA)",
        "note": "One source per economy and series: seasonally adjusted (SA) before unadjusted (NSA), then the longer "
                "run of consecutive quarters to 2019Q4, the OECD's on ties.",
        "rows": table(OUT / "qna.csv", ["iso3", "source", "series", "adjustment", "quarter", "value"], out),
        "for": ["S1.16", "S7.01"],
    }
    manifest["sources"]["macro/oecd_qna"] = {"title": "OECD Data Explorer, SDMX API, Quarterly National Accounts",
                                             "url": OECD.format(flow="OECD.SDD.NAD,DSD_NAMAIN1@<table>", key="<key>"),
                                             "fetched": TODAY}
    manifest["sources"]["macro/imf_qnea"] = {"title": "IMF SDMX 2.1 API, Quarterly National Accounts (QNEA)",
                                             "url": url, "fetched": TODAY}


def _imf(cache: Path, name: str, url: str) -> str:
    path = cache / name
    if not path.exists():
        log(f"downloading {url}")
        path.write_bytes(get(url, timeout=900, accept=UA_CSV))
    return path.read_text(encoding="utf-8-sig")


def _ilo(cache: Path, name: str, flow: str, key: str) -> str:
    path = cache / name
    if not path.exists():
        url = ILO.format(flow=flow, key=key)
        log(f"downloading {url}")
        path.write_bytes(get(url + "&startPeriod=1950", timeout=900, accept=UA_CSV))
    return path.read_text(encoding="utf-8-sig")


def labour(cache: Path, manifest: dict, iso3: set) -> None:
    out = []
    for name, flow in (("unemployment_rate", "DF_UNE_DEAP_SEX_AGE_RT"), ("employment", "DF_EMP_TEMP_SEX_AGE_NB")):
        text = _ilo(cache, f"ilo_{name}_q.csv", flow, ".Q..SEX_T.AGE_YTHADULT_YGE15")
        out += [(r["REF_AREA"], "ILO", name, "NSA", r["TIME_PERIOD"], g6(r["OBS_VALUE"])) for r in sdmx_rows(text)
                if r["REF_AREA"] in iso3 and r["OBS_VALUE"]]
    oecd = {
        "une": ("OECD.SDD.TPS,DSD_LFS@DF_IALFS_UNE_M", ".UNE_LF_M.PT_LF_SUB._Z.Y._T.Y_GE15._Z.Q",
                lambda r: "unemployment_rate"),
        "vac": ("OECD.SDD.TPS,DSD_OLAB@DF_OIALAB_INDIC", ".VAC_U.PS._Z.Y+N..Q",
                lambda r: "vacancies" if r["MEASURE"] == "VAC_U" else None),
        "ear": ("OECD.SDD.TPS,DSD_EAR@DF_HOU_EAR", ".EAR.IX.Y+N...Q",
                lambda r: "hourly_earnings_manufacturing" if r["ACTIVITY"] == "C" else "hourly_earnings"),
    }
    found = {}
    for name, (flow, key, series) in oecd.items():
        text = cached(cache, f"oecd_{name}.csv", OECD.format(flow=flow, key=key)).read_text(encoding="utf-8-sig")
        for r in sdmx_rows(text):
            if r["REF_AREA"] in iso3 and r["OBS_VALUE"] and r["FREQ"] == "Q" and series(r):
                adj = "SA" if r["ADJUSTMENT"] == "Y" else "NSA"
                found.setdefault((r["REF_AREA"], series(r), adj), {})[r["TIME_PERIOD"]] = r["OBS_VALUE"]
    for (iso, series, adj), obs in found.items():
        # Manufacturing earnings stand in only where the private sector's are not published.
        if (adj == "NSA" and (iso, series, "SA") in found) or (
                series == "hourly_earnings_manufacturing" and any((iso, "hourly_earnings", a) in found for a in
                                                                   ("SA", "NSA"))):
            continue
        out += [(iso, "OECD", series.replace("_manufacturing", ""), adj, q, g6(v)) for q, v in obs.items()]
    manifest["series"]["macro/labour_q"] = {
        "title": "Quarterly labour series: unemployment rate (per cent of the labour force, ages 15+) and employment "
                 "(thousands, ages 15+) from national labour force surveys (ILO, national estimates, not modelled, "
                 "unadjusted); the OECD's harmonised unemployment rate (per cent, ages 15+), unfilled job vacancies "
                 "(persons) and an hourly earnings index (private sector, or manufacturing where the private "
                 "sector's is not published), seasonally adjusted where published",
        "rows": table(OUT / "labour_q.csv", ["iso3", "source", "series", "adjustment", "quarter", "value"], out),
        "for": ["S1.16", "S7.01"],
    }
    manifest["sources"]["macro/ilo_sdmx"] = {"title": "ILO SDMX API (ILOSTAT), DF_UNE_DEAP_SEX_AGE_RT and "
                                             "DF_EMP_TEMP_SEX_AGE_NB", "url": ILO.format(flow="<flow>", key="<key>"),
                                             "fetched": TODAY}
    manifest["sources"]["macro/oecd_labour"] = {"title": "OECD Data Explorer, SDMX API, infra-annual labour "
                                                "statistics, job vacancies, hourly earnings",
                                                "url": OECD.format(flow="<flow>", key="<key>"), "fetched": TODAY}


def prices(cache: Path, manifest: dict, iso3: set) -> None:
    out = []
    text = _imf(cache, "imf_cpi.csv", IMF.format(flow="CPI", key=".CPI._T.IX.Q"))
    out += [(ALIAS.get(r["COUNTRY"], r["COUNTRY"]), "cpi", r["TIME_PERIOD"], g6(r["OBS_VALUE"]))
            for r in sdmx_rows(text) if ALIAS.get(r["COUNTRY"], r["COUNTRY"]) in iso3 and r["OBS_VALUE"]]
    text = _imf(cache, "imf_mfs_ir.csv", IMF.format(flow="MFS_IR", key="." + "+".join(MFS_IR) + ".Q"))
    out += [(ALIAS.get(r["COUNTRY"], r["COUNTRY"]), MFS_IR[r["INDICATOR"]], r["TIME_PERIOD"],
             g6(r["OBS_VALUE"])) for r in sdmx_rows(text)
            if ALIAS.get(r["COUNTRY"], r["COUNTRY"]) in iso3 and r["OBS_VALUE"]]
    url = OECD.format(flow="OECD.SDD.STES,DSD_STES@DF_FINMARK", key=".Q." + "+".join(FINMARK) + ".......")
    text = cached(cache, "oecd_finmark.csv", url).read_text(encoding="utf-8-sig")
    out += [(r["REF_AREA"], FINMARK[r["MEASURE"]], r["TIME_PERIOD"], g6(r["OBS_VALUE"]))
            for r in sdmx_rows(text) if r["REF_AREA"] in iso3 and r["OBS_VALUE"]]
    manifest["series"]["macro/prices_rates_q"] = {
        "title": f"Quarterly, {RATES_FIRST} on: consumer price index, all items (IMF CPI, index), and interest "
                 "rates in per cent a year, "
                 "quarterly averages: the monetary-policy rate, banks' lending rate and the Treasury-bill yield "
                 "(IMF Monetary and Financial Statistics, MFS_IR), and the three-month and "
                 "long-term (ten-year government) rates (OECD Financial market statistics)",
        "rows": table(OUT / "prices_rates_q.csv", ["iso3", "series", "quarter", "value"],
                      [r for r in out if r[2][:4] >= str(RATES_FIRST)]),
        "for": ["S1.16", "S7.01"],
    }
    manifest["sources"]["macro/imf_cpi_ir"] = {"title": "IMF SDMX 2.1 API, CPI and MFS_IR dataflows",
                                               "url": IMF.format(flow="<CPI|MFS_IR>", key="<key>"), "fetched": TODAY}
    manifest["sources"]["macro/oecd_finmark"] = {"title": "OECD Data Explorer, SDMX API, financial market statistics "
                                                 "(DF_FINMARK)", "url": url, "fetched": TODAY}


def annual(cache: Path, manifest: dict, iso3: set) -> None:
    by = {}
    for code, name in WB_SERIES.items():
        path = cache / f"wb_{code}.json"
        if not path.exists():
            rows, page = [], 1
            while True:
                data = json.loads(get(WB.format(code=code, page=page), timeout=300))
                rows += data[1] or []
                if page >= data[0]["pages"]:
                    break
                page += 1
            path.write_text(json.dumps(rows))
        for r in json.loads(path.read_text()):
            if r["countryiso3code"] in iso3 and r["value"] is not None:
                by.setdefault((r["countryiso3code"], int(r["date"])), {})[name] = g6(r["value"])
    text = _ilo(cache, "ilo_unemployment_a.csv", "DF_UNE_DEAP_SEX_AGE_RT", ".A..SEX_T.AGE_YTHADULT_YGE15")
    for r in sdmx_rows(text):
        if r["REF_AREA"] in iso3 and r["OBS_VALUE"]:
            by.setdefault((r["REF_AREA"], int(r["TIME_PERIOD"])), {})["unemployment_survey"] = g6(r["OBS_VALUE"])
    url = OECD.format(flow="OECD.ELS.SAE,DSD_EARNINGS@AV_AN_WAGE", key="all")
    for r in sdmx_rows(cached(cache, "oecd_wages.csv", url).read_text(encoding="utf-8-sig")):
        if r["REF_AREA"] in iso3 and r["OBS_VALUE"] and r["UNIT_MEASURE"] == "USD_PPP" and r["PRICE_BASE"] == "Q":
            by.setdefault((r["REF_AREA"], int(r["TIME_PERIOD"])), {})["average_wage_usd_ppp"] = g6(r["OBS_VALUE"])
    url_h = OECD.format(flow="OECD.ELS.SAE,DSD_HW@DF_AVG_ANN_HRS_WKD", key="all")
    for r in sdmx_rows(cached(cache, "oecd_hours.csv", url_h).read_text(encoding="utf-8-sig")):
        if r["REF_AREA"] in iso3 and r["OBS_VALUE"] and r["WORKER_STATUS"] == "_T":
            by.setdefault((r["REF_AREA"], int(r["TIME_PERIOD"])), {})["average_hours"] = g6(r["OBS_VALUE"])
    columns = [*WB_SERIES.values(), "unemployment_survey", "average_wage_usd_ppp", "average_hours"]
    rows = [(iso, year, *(v.get(c, "") for c in columns)) for (iso, year), v in by.items()]
    manifest["series"]["macro/annual"] = {
        "title": "Annual series, 1960-2024: GDP, household consumption and gross fixed capital formation in constant "
                 "2015 dollars, CPI inflation (per cent), net financial account and current account (BoP, "
                 "current dollars), GDP in current dollars, domestic credit to the private sector (per cent of GDP) (World Bank WDI); the "
                 "unemployment rate from national labour force surveys (ILO, national estimates, ages 15+); average "
                 "annual wages (constant 2025 US dollars at PPP) and average annual hours actually worked per worker "
                 "(OECD)",
        "rows": table(OUT / "annual.csv", ["iso3", "year", *columns], rows),
        "for": ["S1.16", "S7.01"],
    }
    manifest["sources"]["macro/wb_api"] = {"title": "World Bank API, WDI series, full history", "url": WB,
                                           "fetched": TODAY}
    manifest["sources"]["macro/oecd_wages_hours"] = {"title": "OECD Data Explorer, average annual wages "
                                                     "(AV_AN_WAGE) and average annual hours worked", "url": url_h,
                                                     "fetched": TODAY}


def jst(cache: Path, manifest: dict, iso3: set) -> None:
    book = openpyxl.load_workbook(cached(cache, "JSTdatasetR6.xlsx", JST_URL), read_only=True)
    rows = book[book.sheetnames[0]].iter_rows(values_only=True)
    header = list(next(rows))
    at = [header.index(c) for c in JST_COLUMNS]
    out = [(r[header.index("iso")], int(r[0]), *("" if r[i] is None else g6(r[i]) for i in at))
           for r in rows if r[0] is not None]
    manifest["series"]["macro/jst"] = {
        "title": "Jordà-Schularick-Taylor Macrohistory Database R6, 18 developed economies, 1870-2020: population, "
                 "real GDP per capita (Maddison; Barro index), real consumption per capita (index), nominal GDP, "
                 "investment to GDP, CPI, current account, money, short and long rates, unemployment, wages, dollar "
                 "exchange rate, total loans to the non-financial private sector and their mortgage, household and "
                 "business parts (national currency), public debt to GDP, nominal house prices, systemic crisis "
                 "dummy (crisisJST), peg dummy",
        "rows": table(OUT / "jst.csv", ["iso3", "year", *JST_COLUMNS], out),
        "for": ["S7.01"],
    }
    manifest["sources"]["macro/jst"] = {"title": "Jordà, Schularick and Taylor, Macrohistory Database, release 6",
                                        "url": JST_URL, "fetched": TODAY, "release": "R6"}


def xru(cache: Path, manifest: dict, iso3: set) -> None:
    iso3_of = {r["iso2"]: r["iso3"] for r in csv.DictReader((RAW / "wb" / "countries.csv").open()) if r["iso2"]}
    z = zipfile.ZipFile(cached(cache, "bis_xru.zip", XRU_URL))
    code = lambda text: text.split(":", 1)[0].strip()
    out = []
    reader = csv.reader(io.TextIOWrapper(z.open(z.namelist()[0]), encoding="utf-8-sig"))
    header = [code(h) for h in next(reader)]
    at = {h: header.index(h) for h in ("FREQ", "REF_AREA", "COLLECTION", "TIME_PERIOD", "OBS_VALUE")}
    for r in reader:
        f = {h: code(r[i]) for h, i in at.items()}
        iso = iso3_of.get(f["REF_AREA"])
        if f["FREQ"] == "M" and f["COLLECTION"] == "E" and iso in iso3 and f["OBS_VALUE"] \
                and int(f["TIME_PERIOD"][:4]) >= XRU_FIRST and f["OBS_VALUE"] != "NaN":
            out.append((iso, f["TIME_PERIOD"], g6(f["OBS_VALUE"])))
    release = max(i.date_time for i in z.infolist())
    manifest["series"]["macro/xru_monthly"] = {
        "title": f"US-dollar exchange rates, national currency per dollar, monthly, end of period, {XRU_FIRST} on "
                 "(BIS WS_XRU); the economy's currency as the BIS publishes it (euro-area members in euro)",
        "rows": table(OUT / "xru_monthly.csv", ["iso3", "month", "value"], out),
        "for": ["S5.04", "S7.01"],
    }
    manifest["sources"]["macro/bis_xru"] = {"title": "BIS, US dollar exchange rates (WS_XRU), bulk CSV", "url": XRU_URL,
                                            "fetched": TODAY, "release": datetime.date(*release[:3]).isoformat()}


# The papers' tables, typed from the PDFs downloaded for this fetch: country, iso3, the paper's group, then values.
AG_REF = ("Aguiar and Gopinath, 'Emerging Market Business Cycles: The Cycle is the Trend', NBER Working Paper 10734 "
          "(August 2004; published JPE 115(1), 2007), Tables 2A-2C, pp. 32-34 of the PDF: quarterly, deseasonalised, "
          "log, HP 1600; sd in per cent")
AG = """Argentina,ARG,emerging,3.68,2.28,0.85,0.61,1.38,2.53,2.56,0.90,0.96,-0.70
Brazil,BRA,emerging,1.98,1.69,0.65,0.35,2.01,3.08,2.61,0.41,0.62,0.01
Ecuador,ECU,emerging,2.44,1.52,0.82,0.15,2.39,5.56,5.68,0.73,0.89,-0.79
Israel,ISR,emerging,1.95,1.99,0.50,0.27,1.60,3.42,2.12,0.45,0.49,0.12
Korea,KOR,emerging,2.51,1.71,0.78,0.17,1.23,2.50,2.32,0.85,0.78,-0.61
Malaysia,MYS,emerging,3.10,1.84,0.85,0.56,1.70,4.82,5.30,0.76,0.86,-0.74
Mexico,MEX,emerging,2.48,1.53,0.82,0.27,1.24,4.05,2.19,0.92,0.91,-0.74
Peru,PER,emerging,3.68,2.97,0.64,0.12,0.92,2.37,1.25,0.78,0.85,-0.24
Philippines,PHL,emerging,3.00,1.66,0.87,0.17,0.62,4.66,3.21,0.59,0.76,-0.41
Slovak Republic,SVK,emerging,1.24,1.06,0.66,0.20,2.04,7.77,4.29,0.42,0.46,-0.44
South Africa,ZAF,emerging,1.62,0.82,0.89,0.58,1.61,3.94,2.57,0.72,0.75,-0.54
Thailand,THA,emerging,4.35,2.25,0.89,0.42,1.09,3.49,4.58,0.92,0.91,-0.83
Turkey,TUR,emerging,3.57,2.92,0.67,0.05,1.09,2.71,3.23,0.89,0.83,-0.69
Australia,AUS,developed,1.39,0.84,0.84,0.36,0.69,3.69,1.08,0.48,0.80,-0.43
Austria,AUT,developed,0.89,0.47,0.90,0.52,0.87,2.75,0.65,0.74,0.75,0.10
Belgium,BEL,developed,1.02,0.71,0.79,0.18,0.81,3.72,0.91,0.67,0.62,-0.04
Canada,CAN,developed,1.64,0.81,0.91,0.55,0.77,2.63,0.91,0.88,0.77,-0.20
Denmark,DNK,developed,1.02,1.04,0.49,0.15,1.19,3.90,0.88,0.36,0.51,-0.08
Finland,FIN,developed,2.18,1.32,0.85,0.01,0.94,3.26,1.11,0.84,0.88,-0.45
Netherlands,NLD,developed,1.20,0.88,0.77,0.03,1.07,2.92,0.71,0.72,0.70,-0.19
New Zealand,NZL,developed,1.56,1.13,0.77,0.02,0.90,4.38,1.37,0.76,0.82,-0.26
Norway,NOR,developed,1.40,1.46,0.48,0.46,1.32,4.33,1.73,0.63,0.00,0.11
Portugal,PRT,developed,1.34,1.03,0.72,0.28,1.02,2.88,1.16,0.75,0.70,-0.11
Spain,ESP,developed,1.11,0.75,0.82,0.08,1.11,3.70,0.86,0.83,0.83,-0.60
Sweden,SWE,developed,1.52,1.45,0.53,0.35,0.97,3.66,0.94,0.35,0.68,0.01
Switzerland,CHE,developed,1.11,0.50,0.92,0.81,0.51,2.56,0.96,0.58,0.69,-0.03"""
AG_COLUMNS = ["sd_y", "sd_dy", "rho_y", "rho_dy", "sd_c_over_sd_y", "sd_i_over_sd_y", "sd_nx", "corr_c_y",
              "corr_i_y", "corr_nx_y"]
NP_REF = ("Neumeyer and Perri, 'Business Cycles in Emerging Economies: the Role of Interest Rates', NBER Working Paper "
          "10387 (March 2004; published JME 52(2), 2005), Tables 1A and 1B, pp. 43-44 of the PDF: quarterly, log, HP; "
          "EMP and HRS on semiannual data; relative sd are per cent of GDP's sd over 100; NA left empty")
NP = """Argentina,ARG,emerging,4.22,1.08,1.17,2.95,0.39,0.57,0.94,0.97,0.94,0.36,0.52
Brazil,BRA,emerging,1.76,1.93,1.24,3.05,0.89,1.95,0.48,0.58,0.80,0.62,0.75
Korea,KOR,emerging,3.54,1.34,2.05,2.20,0.59,0.71,0.96,0.92,0.94,0.91,0.96
Mexico,MEX,emerging,2.98,1.21,1.29,3.83,0.43,0.33,0.93,0.96,0.96,0.56,0.37
Philippines,PHL,emerging,1.44,0.93,2.78,4.44,1.34,,0.69,0.51,0.76,0.26,
Australia,AUS,developed,1.19,0.84,1.20,4.13,1.13,1.40,0.63,0.79,0.87,0.77,0.76
Canada,CAN,developed,1.39,0.74,0.84,2.91,0.75,0.82,0.83,0.86,0.73,0.93,0.93
Netherlands,NLD,developed,0.93,1.17,1.44,2.66,1.27,,0.64,0.77,0.58,0.81,
New Zealand,NZL,developed,1.99,0.82,0.86,3.32,1.15,1.28,0.72,0.59,0.66,0.73,0.73
Sweden,SWE,developed,1.35,1.01,1.67,4.18,1.24,2.94,0.55,0.38,0.81,0.81,0.93"""
NP_COLUMNS = ["sd_y", "sd_pc_over_sd_y", "sd_tc_over_sd_y", "sd_i_over_sd_y", "sd_emp_over_sd_y",
              "sd_hrs_over_sd_y", "corr_pc_y", "corr_tc_y", "corr_i_y", "corr_emp_y", "corr_hrs_y"]
CF_REF = ("Calderón and Fuentes, 'Characterizing the Business Cycles of Emerging Economies', World Bank Policy "
          "Research Working Paper 5343 (June 2010), Table 1, p. 11: classical cycles by Harding and Pagan's quarterly "
          "algorithm, GDP 1980Q1-2006Q2; mean durations in quarters, amplitudes in per cent; '.' left empty")
CF = """Argentina,ARG,LAC,4.5,7.1,-9.4,12.1,8
Bolivia,BOL,LAC,3.0,,-2.0,,1
Brazil,BRA,LAC,2.8,6.9,-4.3,9.6,10
Chile,CHL,LAC,3.3,30.0,-10.1,55.6,3
Colombia,COL,LAC,3.0,62.0,-4.1,59.9,2
Costa Rica,CRI,LAC,2.5,9.0,-0.6,13.1,2
Ecuador,ECU,LAC,3.2,10.0,-5.2,11.6,6
Mexico,MEX,LAC,3.7,12.6,-4.4,13.9,6
Paraguay,PRY,LAC,3.5,5.3,-6.1,7.2,4
Peru,PER,LAC,3.7,8.2,-12.6,18.0,7
Uruguay,URY,LAC,5.5,8.5,-9.8,12.4,4
Venezuela,VEN,LAC,4.6,6.3,-9.0,8.5,8
Hong Kong,HKG,Asia,3.0,13.6,-4.2,24.7,6
Indonesia,IDN,Asia,3.3,29.0,-7.5,62.6,3
Korea,KOR,Asia,3.0,,-9.3,,1
Malaysia,MYS,Asia,4.5,44.0,-8.0,91.3,2
Philippines,PHL,Asia,5.0,17.0,-6.3,20.3,4
Singapore,SGP,Asia,3.8,19.3,-4.6,43.1,4
Taiwan,TWN,Asia,3.0,5.0,-3.2,7.7,2
Thailand,THA,Asia,8.0,,-16.1,,1
India,IND,other emerging,3.0,30.0,-2.0,57.1,2
South Africa,ZAF,other emerging,8.3,9.5,-4.6,8.7,3
Turkey,TUR,other emerging,3.0,11.8,-7.8,20.9,5"""
CF_COLUMNS = ["contraction_quarters", "expansion_quarters", "contraction_amplitude", "expansion_amplitude",
              "contractions"]
PAPER_NAMES = {"aguiar_gopinath": "Aguiar and Gopinath", "neumeyer_perri": "Neumeyer and Perri",
               "calderon_fuentes": "Calderón and Fuentes"}
PAPERS = {
    "aguiar_gopinath": (AG, AG_COLUMNS, AG_REF, "https://www.nber.org/system/files/working_papers/w10734/w10734.pdf"),
    "neumeyer_perri": (NP, NP_COLUMNS, NP_REF, "https://www.nber.org/system/files/working_papers/w10387/w10387.pdf"),
    "calderon_fuentes": (CF, CF_COLUMNS, CF_REF,
                         "https://documents.worldbank.org/curated/en/524811468325165721/pdf/WPS5343.pdf"),
}


def typed(cache: Path, manifest: dict, iso3: set) -> None:
    for name, (text, columns, ref, url) in PAPERS.items():
        rows = [r.split(",") for r in text.splitlines()]
        manifest["series"][f"macro/typed_{name}"] = {
            "title": f"Business-cycle moments per economy, typed: {ref}",
            "rows": table(OUT / f"typed_{name}.csv", ["country", "iso3", "group", *columns], rows),
            "for": ["S1.16", "S7.01"],
            "note": f"Typed by hand from the paper's table in the PDF downloaded on {TODAY} from {url}; the group is "
                    "the paper's own classification.",
        }
        manifest["sources"][f"macro/{name}"] = {"title": ref, "url": url, "fetched": TODAY}


# ---------------------------------------------------------------------------------------------------- statistics


def quantile(xs: list, p: float) -> float:
    xs = sorted(xs)
    h = (len(xs) - 1) * p
    lo = math.floor(h)
    return xs[lo] + (h - lo) * (xs[min(lo + 1, len(xs) - 1)] - xs[lo])


def hp(y: list, lam: float) -> list:
    """The Hodrick-Prescott cycle, y less the trend solving (I + lam K'K) trend = y, by banded Cholesky."""
    n = len(y)
    diag = [1 + 6 * lam] * n
    diag[0] = diag[-1] = 1 + lam
    diag[1] = diag[-2] = 1 + 5 * lam
    off1 = [-4 * lam] * (n - 1)
    off1[0] = off1[-1] = -2 * lam
    band = lambda i, j: diag[i] if i == j else off1[j] if i - j == 1 else lam if i - j == 2 else 0.0
    low = [[0.0, 0.0, 0.0] for _ in range(n)]  # low[i][k] is L[i][i-2+k]
    for i in range(n):
        for j in range(max(0, i - 2), i + 1):
            s = band(i, j) - sum(low[i][k - i + 2] * low[j][k - j + 2] for k in range(max(0, i - 2), j))
            low[i][j - i + 2] = math.sqrt(s) if i == j else s / low[j][2]
    z = [0.0] * n
    for i in range(n):
        z[i] = (y[i] - sum(low[i][k - i + 2] * z[k] for k in range(max(0, i - 2), i))) / low[i][2]
    trend = [0.0] * n
    for i in reversed(range(n)):
        trend[i] = (z[i] - sum(low[k][i - k + 2] * trend[k] for k in range(i + 1, min(n, i + 3)))) / low[i][2]
    return [a - b for a, b in zip(y, trend)]


def deseason(y: list, first_quarter: int) -> list:
    """Classical additive deseasonalising of a quarterly log series: each quarter's mean deviation from the centred
    2x4 moving average, centred on zero, is removed."""
    n = len(y)
    dev = {q: [] for q in range(4)}
    for t in range(2, n - 2):
        ma = (y[t - 2] / 2 + y[t - 1] + y[t] + y[t + 1] + y[t + 2] / 2) / 4
        dev[(first_quarter + t) % 4].append(y[t] - ma)
    factor = {q: statistics.fmean(v) for q, v in dev.items()}
    centre = statistics.fmean(factor.values())
    return [y[t] - factor[(first_quarter + t) % 4] + centre for t in range(n)]


def corr(a: list, b: list) -> float:
    return statistics.correlation(a, b)


def ols(x: list, y: list) -> tuple:
    """Slope and intercept of y on x."""
    mx, my = statistics.fmean(x), statistics.fmean(y)
    sxx = sum((v - mx) ** 2 for v in x)
    slope = sum((u - mx) * (v - my) for u, v in zip(x, y)) / sxx
    return slope, my - slope * mx


def solve(a: list, b: list) -> list:
    """Gaussian elimination with partial pivoting for the small normal equations of logit and probit."""
    n = len(b)
    m = [row[:] + [b[i]] for i, row in enumerate(a)]
    for c in range(n):
        p = max(range(c, n), key=lambda r: abs(m[r][c]))
        m[c], m[p] = m[p], m[c]
        for r in range(c + 1, n):
            f = m[r][c] / m[c][c]
            for k in range(c, n + 1):
                m[r][k] -= f * m[c][k]
    x = [0.0] * n
    for r in reversed(range(n)):
        x[r] = (m[r][n] - sum(m[r][k] * x[k] for k in range(r + 1, n))) / m[r][r]
    return x


def binary_fit(xs: list, ys: list, link: str) -> list:
    """Maximum-likelihood logit or probit coefficients (constant first) by Newton's method."""
    k = len(xs[0]) + 1
    beta = [0.0] * k
    phi = lambda z: math.exp(-z * z / 2) / math.sqrt(2 * math.pi)
    for _ in range(100):
        grad, hess = [0.0] * k, [[0.0] * k for _ in range(k)]
        for x, y in zip(xs, ys):
            row = [1.0, *x]
            z = sum(b * v for b, v in zip(beta, row))
            if link == "logit":
                p = 1 / (1 + math.exp(-z))
                g, w = y - p, p * (1 - p)
            else:
                p = min(max(0.5 * (1 + math.erf(z / math.sqrt(2))), 1e-12), 1 - 1e-12)
                d = phi(z)
                g = d * (y - p) / (p * (1 - p))
                w = d * d / (p * (1 - p))
            for i in range(k):
                grad[i] += g * row[i]
                for j in range(k):
                    hess[i][j] += w * row[i] * row[j]
        step = solve(hess, grad)
        beta = [b + s for b, s in zip(beta, step)]
        if max(abs(s) for s in step) < 1e-10:
            break
    return beta


def auc(scores: list, ys: list) -> float:
    pos = [s for s, y in zip(scores, ys) if y]
    neg = sorted(s for s, y in zip(scores, ys) if not y)
    wins = sum(bisect.bisect_left(neg, s) + (bisect.bisect_right(neg, s) - bisect.bisect_left(neg, s)) / 2
               for s in pos)
    return wins / (len(pos) * len(neg))


def subbotin(xs: list) -> float:
    """Exponential-power shape b of standardised data by profile maximum likelihood (1 Laplace, 2 normal)."""
    n = len(xs)

    def loglik(b):
        a = (sum(abs(x) ** b for x in xs) / n) ** (1 / b)
        return -n * (math.log(2 * a) + math.log(b) / b + math.lgamma(1 + 1 / b)) - n / b

    grid = [0.2 + 0.05 * i for i in range(97)]
    b = max(grid, key=loglik)
    lo, hi = max(0.15, b - 0.05), b + 0.05
    for _ in range(60):
        m1, m2 = lo + (hi - lo) / 3, hi - (hi - lo) / 3
        lo, hi = (lo, m2) if loglik(m1) > loglik(m2) else (m1, hi)
    return (lo + hi) / 2


def standardise(xs: list) -> list:
    m, s = statistics.fmean(xs), statistics.stdev(xs)
    return [(x - m) / s for x in xs]


def turning_points(y: list, window: int, min_phase: int, min_cycle: int) -> list:
    """Harding and Pagan's dating: local extremes over +-window, alternated, the ends and short phases and cycles
    censored. Returns (index, 'P' or 'T') in time order."""
    n = len(y)
    pts = []
    for t in range(window, n - window):
        around = y[t - window:t] + y[t + 1:t + window + 1]
        if all(y[t] > v for v in around):
            pts.append((t, "P"))
        elif all(y[t] < v for v in around):
            pts.append((t, "T"))
    better = lambda a, b: (y[a[0]] >= y[b[0]]) if a[1] == "P" else (y[a[0]] <= y[b[0]])
    changed = True
    while changed:
        changed = False
        out = []
        for p in pts:
            if out and out[-1][1] == p[1]:
                out[-1] = out[-1] if better(out[-1], p) else p
                changed = True
            else:
                out.append(p)
        pts = out
        # The first and last points must be more extreme than the series' ends.
        if pts and ((pts[0][1] == "P" and y[pts[0][0]] < y[0]) or (pts[0][1] == "T" and y[pts[0][0]] > y[0])):
            pts, changed = pts[1:], True
        if pts and ((pts[-1][1] == "P" and y[pts[-1][0]] < y[-1]) or (pts[-1][1] == "T" and y[pts[-1][0]] > y[-1])):
            pts, changed = pts[:-1], True
        for i in range(len(pts) - 1):
            if pts[i + 1][0] - pts[i][0] < min_phase:
                pts = pts[:i] + pts[i + 2:]
                changed = True
                break
        else:
            for i in range(len(pts) - 2):
                if pts[i + 2][0] - pts[i][0] < min_cycle:
                    drop = i + 2 if better(pts[i], pts[i + 2]) else i
                    pts = [p for j, p in enumerate(pts) if j not in (drop, i + 1)]
                    changed = True
                    break
    return pts


def phases(pts: list) -> tuple:
    recessions = [b[0] - a[0] for a, b in zip(pts, pts[1:]) if a[1] == "P"]
    expansions = [b[0] - a[0] for a, b in zip(pts, pts[1:]) if a[1] == "T"]
    return recessions, expansions


# ---------------------------------------------------------------------------------------------------- the ranges


def qindex(q: str) -> int:
    return int(q[:4]) * 4 + int(q[-1]) - 1


def runs(obs: dict, first=None, last=None) -> list:
    """The longest run of consecutive quarters (or years) of a {period index: value} map, within the bounds."""
    keys = sorted(k for k in obs if (first is None or k >= first) and (last is None or k <= last))
    best, cur = [], []
    for k in keys:
        cur = cur + [k] if cur and k == cur[-1] + 1 else [k]
        if len(cur) > len(best):
            best = cur
    return best


class Data:
    """The fetched files read back, by economy."""

    def __init__(self):
        self.level = levels()
        self.q = {}  # (iso, series) -> ({quarter index: value}, adjustment), from the chosen source
        self.qsrc = {}
        qna_rows = list(csv.DictReader((OUT / "qna.csv").open()))
        candidates = {}
        for r in qna_rows:
            candidates.setdefault((r["iso3"], r["series"], r["source"], r["adjustment"]), {})[qindex(r["quarter"])] = \
                float(r["value"])
        lab = list(csv.DictReader((OUT / "labour_q.csv").open()))
        for r in lab:
            candidates.setdefault((r["iso3"], r["series"], r["source"], r["adjustment"]), {})[qindex(r["quarter"])] = \
                float(r["value"])
        # A series is taken from one source: adjusted before unadjusted, then the longer run to 2019, OECD on ties.
        best = {}
        for (iso, series, source, adj), obs in candidates.items():
            key = (adj == "SA", len(runs(obs, last=LAST_YEAR * 4 + 3)), source == "OECD")
            if (iso, series) not in best or key > best[(iso, series)][0]:
                best[(iso, series)] = (key, source, adj, obs)
        for (iso, series), (_, source, adj, obs) in best.items():
            self.q[(iso, series)] = (obs, adj)
            self.qsrc[(iso, series)] = f"{source} {adj}"
        self.pr = {}
        for r in csv.DictReader((OUT / "prices_rates_q.csv").open()):
            self.pr.setdefault((r["iso3"], r["series"]), {})[qindex(r["quarter"])] = float(r["value"])
        self.pwt = {}
        for r in csv.DictReader((OUT / "pwt.csv").open()):
            self.pwt.setdefault(r["iso3"], {})[int(r["year"])] = r
        self.ann = {}
        for r in csv.DictReader((OUT / "annual.csv").open()):
            self.ann.setdefault(r["iso3"], {})[int(r["year"])] = r
        self.jst = {}
        for r in csv.DictReader((OUT / "jst.csv").open()):
            self.jst.setdefault(r["iso3"], {})[int(r["year"])] = r

    def quarterly(self, iso: str, series: str, log_: bool = True) -> dict:
        """A quarterly series to 2019, logged if asked, deseasonalised if the source did not adjust it."""
        if (iso, series) not in self.q:
            return {}
        obs, adj = self.q[(iso, series)]
        span = runs({k: v for k, v in obs.items() if v > 0 or not log_}, last=LAST_YEAR * 4 + 3)
        if len(span) < 12:
            return {}
        y = [math.log(obs[k]) if log_ else obs[k] for k in span]
        if adj == "NSA":
            y = deseason(y, span[0] % 4)
        return dict(zip(span, y))


def cycles(series: dict, lam: float) -> dict:
    keys = sorted(series)
    return dict(zip(keys, hp([series[k] for k in keys], lam)))


def common(*maps) -> list:
    keys = set(maps[0])
    for m in maps[1:]:
        keys &= set(m)
    return sorted(keys)


class Ranges:
    def __init__(self, data: Data):
        self.d = data
        self.rows = []

    def add(self, fact, statistic, values: dict, frequency, sample, method, source, pooled: dict = None):
        """One row per level: the median and interquartile range of the per-economy values."""
        for level in ("developed", "emerging", "developing"):
            xs = [v for iso, v in values.items() if self.d.level.get(iso) == level and v is not None
                  and math.isfinite(v)]
            if not xs and not (pooled and level in pooled):
                continue
            row = [fact, statistic, level, len(xs),
                   f"{statistics.median(xs):.4g}" if xs else "", f"{quantile(xs, 0.25):.4g}" if xs else "",
                   f"{quantile(xs, 0.75):.4g}" if xs else "",
                   f"{pooled[level]:.4g}" if pooled and level in pooled else "", frequency, sample, method, source,
                   " ".join(sorted(iso for iso, v in values.items() if self.d.level.get(iso) == level
                                   and v is not None and math.isfinite(v)))]
            self.rows.append(row)

    # F01 --------------------------------------------------------------------------------------------------------
    def f01(self):
        growth = {}
        for iso, years in self.d.pwt.items():
            g = [100 * math.log(float(years[y]["rgdpna"]) / float(years[y - 1]["rgdpna"]))
                 for y in years if y - 1 in years and y <= LAST_YEAR and years[y]["rgdpna"] and years[y - 1]["rgdpna"]]
            if len(g) >= 20:
                growth[iso] = statistics.fmean(g)
        self.add("F01", "mean annual growth of real output, per cent", growth, "annual", "1950-2019, economies with "
                 "20 or more years", "per economy the mean of 100 x annual log change of rgdpna", "macro/pwt")
        ratio, rec, cv = {}, {}, {}
        for iso in self.d.level:
            y = self.d.quarterly(iso, "gdp")
            if len(y) < 80:
                continue
            keys = sorted(y)
            r, e = phases(turning_points([y[k] for k in keys], 2, 2, 5))
            if r:
                rec[iso] = statistics.fmean(r)
            if r and e:
                ratio[iso] = statistics.fmean(e) / statistics.fmean(r)
            if len(e) >= 2:
                cv[iso] = statistics.stdev(e) / statistics.fmean(e)
        src, sample = "macro/qna (OECD QNA, IMF QNEA)", "quarterly real GDP to 2019Q4, economies with 80 or more " \
                                                         "consecutive quarters"
        hpm = "Harding-Pagan dating of log real GDP (window 2 quarters, phase 2, cycle 5, ends censored); " \
              "unadjusted series deseasonalised by quarter means around a centred 2x4 moving average"
        self.add("F01", "mean expansion length over mean recession length", ratio, "quarterly", sample,
                 hpm + "; economies with at least one complete expansion and recession", src)
        self.add("F01", "mean recession length, quarters", rec, "quarterly", sample, hpm, src)
        self.add("F01", "coefficient of variation of expansion lengths", cv, "quarterly", sample,
                 hpm + "; sample sd over mean, economies with two or more complete expansions", src)
        # Annual dating reaches the economies with no long quarterly series.
        ratio_a, rec_a = {}, {}
        for iso, years in self.d.pwt.items():
            span = runs({y: float(v["rgdpna"]) for y, v in years.items() if v["rgdpna"]}, last=LAST_YEAR)
            if len(span) < 20:
                continue
            r, e = phases(turning_points([math.log(float(years[k]["rgdpna"])) for k in span], 1, 1, 2))
            if r:
                rec_a[iso] = 4 * statistics.fmean(r)
            if r and e:
                ratio_a[iso] = statistics.fmean(e) / statistics.fmean(r)
        amethod = "Harding-Pagan dating of annual log rgdpna (window 1 year, phase 1, cycle 2, ends censored); " \
                  "annual dating misses short recessions, so lengths run longer than quarterly"
        self.add("F01", "mean expansion length over mean recession length (annual dating)", ratio_a, "annual",
                 "1950-2019, economies with 20 or more years", amethod, "macro/pwt")
        self.add("F01", "mean recession length, quarters (annual dating, years x 4)", rec_a, "annual",
                 "1950-2019, economies with 20 or more years", amethod, "macro/pwt")
        cf = self.typed("calderon_fuentes")
        tmethod = "per-economy values typed from Calderón and Fuentes (2010) Table 1, grouped by current World Bank " \
                  "level"
        self.add("F01", "mean recession length, quarters (Calderón and Fuentes)",
                 {i: v["contraction_quarters"] for i, v in cf.items()},
                 "quarterly", "1980Q1-2006Q2", tmethod, "macro/typed_calderon_fuentes")
        self.add("F01", "mean expansion length over mean recession length (Calderón and Fuentes)",
                 {i: v["expansion_quarters"] / v["contraction_quarters"] for i, v in cf.items()
                  if v["expansion_quarters"] is not None}, "quarterly", "1980Q1-2006Q2", tmethod,
                 "macro/typed_calderon_fuentes")

    def typed(self, name: str) -> dict:
        out = {}
        for r in csv.DictReader((OUT / f"typed_{name}.csv").open()):
            out[r["iso3"]] = {k: (float(v) if v not in ("", None) else None) for k, v in r.items()
                              if k not in ("country", "iso3", "group")}
        return out

    # F02, F03 ---------------------------------------------------------------------------------------------------
    def f02_f03(self):
        rel, co = {}, {}
        for iso in self.d.level:
            y = self.d.quarterly(iso, "gdp")
            if len(y) < 40:
                continue
            cy = cycles(y, 1600)
            parts = {s: self.d.quarterly(iso, s) for s in ("investment", "consumption_private", "employment")}
            nds = [self.d.quarterly(iso, s) for s in ("consumption_semidurable", "consumption_nondurable",
                                                       "consumption_services")]
            if nds[1] and nds[2]:
                keys = common(*[m for m in nds if m])
                parts["consumption_nds"] = {k: math.log(sum(math.exp(m[k]) for m in nds if m)) for k in keys}
            u = self.d.quarterly(iso, "unemployment_rate", log_=False)
            if u:
                parts["unemployment_rate"] = u
            for s, m in parts.items():
                keys = common(cy, m)
                if len(keys) < 40:
                    continue
                cs = cycles({k: m[k] for k in keys}, 1600)
                a, b = [cy[k] for k in keys], [cs[k] for k in keys]
                rel.setdefault(s, {})[iso] = statistics.pstdev(b) / statistics.pstdev(a)
                co.setdefault(s, {})[iso] = corr(a, b)
        sample = "quarterly to 2019Q4, economies with 40 or more common consecutive quarters"
        m = "HP 1600 on logs (the unemployment rate in levels) over each pair's common quarters; unadjusted series " \
            "deseasonalised by quarter means around a centred 2x4 moving average"
        src = "macro/qna, macro/labour_q"
        names = {"investment": "fixed investment", "consumption_private": "private consumption (households and NPISH, "
                 "durables included)", "consumption_nds": "household consumption of semi-durables, non-durables and "
                 "services (their volumes summed)", "employment": "employment", "unemployment_rate":
                 "unemployment rate"}
        for s in ("investment", "consumption_nds", "consumption_private"):
            self.add("F02", f"sd of cyclical {names[s]} over sd of cyclical output", rel.get(s, {}), "quarterly",
                     sample, m, src)
        for s in ("consumption_nds", "consumption_private", "investment", "employment", "unemployment_rate"):
            self.add("F03", f"correlation of cyclical {names[s]} with cyclical output", co.get(s, {}), "quarterly",
                     sample, m, src)
        # Annual: every economy with national accounts, HP 100.
        rel_a, co_a = {}, {}
        for iso, years in self.d.ann.items():
            series = {c: {y: math.log(float(v[c])) for y, v in years.items() if v[c] and float(v[c]) > 0
                          and y <= LAST_YEAR} for c in ("gdp_real", "investment_real", "consumption_household_real")}
            span = runs(series["gdp_real"])
            if len(span) < 20:
                continue
            cy = cycles({k: series["gdp_real"][k] for k in span}, 100)
            for c in ("investment_real", "consumption_household_real"):
                keys = runs({k: 0 for k in common(cy, series[c])})
                if len(keys) < 20:
                    continue
                cs = cycles({k: series[c][k] for k in keys}, 100)
                cyk = cycles({k: series["gdp_real"][k] for k in keys}, 100)
                a, b = [cyk[k] for k in keys], [cs[k] for k in keys]
                rel_a.setdefault(c, {})[iso] = statistics.pstdev(b) / statistics.pstdev(a)
                co_a.setdefault(c, {})[iso] = corr(a, b)
        am = "HP 100 on annual logs over each pair's common years"
        asample = "1960-2019, economies with 20 or more common years"
        self.add("F02", "sd of cyclical fixed investment over sd of cyclical output (annual)",
                 rel_a.get("investment_real", {}), "annual", asample, am, "macro/annual (WDI)")
        self.add("F02", "sd of cyclical household consumption over sd of cyclical output (annual)",
                 rel_a.get("consumption_household_real", {}), "annual", asample, am, "macro/annual (WDI)")
        self.add("F03", "correlation of cyclical household consumption with cyclical output (annual)",
                 co_a.get("consumption_household_real", {}), "annual", asample, am, "macro/annual (WDI)")
        self.add("F03", "correlation of cyclical fixed investment with cyclical output (annual)",
                 co_a.get("investment_real", {}), "annual", asample, am, "macro/annual (WDI)")
        for name, cols in (("aguiar_gopinath", {"F02": [("sd_i_over_sd_y", "fixed investment"),
                                                        ("sd_c_over_sd_y", "private consumption")],
                                                "F03": [("corr_c_y", "private consumption"),
                                                        ("corr_i_y", "fixed investment")]}),
                           ("neumeyer_perri", {"F02": [("sd_i_over_sd_y", "fixed investment"),
                                                       ("sd_pc_over_sd_y", "private consumption")],
                                               "F03": [("corr_pc_y", "private consumption"),
                                                       ("corr_i_y", "fixed investment"),
                                                       ("corr_emp_y", "employment")]})):
            t = self.typed(name)
            for fact, pairs in cols.items():
                for col, label in pairs:
                    stat = f"sd of cyclical {label} over sd of cyclical output ({PAPER_NAMES[name]})" \
                        if fact == "F02" else f"correlation of cyclical {label} with cyclical output " \
                                              f"({PAPER_NAMES[name]})"
                    self.add(fact, stat, {i: v[col] for i, v in t.items() if v[col] is not None}, "quarterly",
                             "the paper's samples (1980s-2003)",
                             f"per-economy values typed from {PAPER_NAMES[name]} (HP 1600), "
                             "grouped by current World Bank level", f"macro/typed_{name}")

    # F04 --------------------------------------------------------------------------------------------------------
    def f04(self):
        slope = {}
        for iso, years in self.d.ann.items():
            pw = self.d.pwt.get(iso, {})
            pts = []
            for y in years:
                if y > LAST_YEAR or y - 1 not in years:
                    continue
                u1, u0 = years[y]["unemployment_survey"], years[y - 1]["unemployment_survey"]
                g1, g0 = years[y]["gdp_real"], years[y - 1]["gdp_real"]
                if not (g1 and g0) and y in pw and y - 1 in pw:
                    g1, g0 = pw[y]["rgdpna"], pw[y - 1]["rgdpna"]
                if u1 and u0 and g1 and g0:
                    pts.append((100 * math.log(float(g1) / float(g0)), float(u1) - float(u0)))
            if len(pts) >= 15:
                slope[iso] = ols([p[0] for p in pts], [p[1] for p in pts])[0]
        self.add("F04", "OLS slope of the change in the unemployment rate on output growth", slope, "annual",
                 "to 2019, economies with 15 or more years", "per economy OLS of the yearly change in the survey "
                 "unemployment rate (points) on 100 x the log change of real GDP (WDI, else PWT rgdpna)",
                 "macro/annual (ILO national survey estimates, WDI), macro/pwt")

    # F05 --------------------------------------------------------------------------------------------------------
    def f05(self):
        r = {}
        for iso in self.d.level:
            v = self.d.quarterly(iso, "vacancies")
            u = self.d.quarterly(iso, "unemployment_rate")
            keys = common(v, u)
            span = runs({k: 0 for k in keys})
            if len(span) < 40:
                continue
            cv = cycles({k: v[k] for k in span}, 100000)
            cu = cycles({k: u[k] for k in span}, 100000)
            r[iso] = corr([cv[k] for k in span], [cu[k] for k in span])
        self.add("F05", "correlation of cyclical log vacancies with cyclical log unemployment", r, "quarterly",
                 "to 2019Q4, economies with 40 or more common quarters", "HP 1e5 on log unfilled vacancies (OECD) and "
                 "the log unemployment rate over their common quarters", "macro/labour_q")

    # F06 --------------------------------------------------------------------------------------------------------
    def f06(self):
        med, share = {}, {}
        for iso in self.d.level:
            w = self.d.quarterly(iso, "hourly_earnings")
            u = self.d.quarterly(iso, "unemployment_rate", log_=False)
            if not w or not u:
                continue
            inflation = {k: 400 * (w[k] - w[k - 1]) for k in w if k - 1 in w}
            span = runs({k: 0 for k in common(inflation, u)})
            if len(span) < 80:
                continue
            gap = cycles({k: u[k] for k in span}, 1600)
            slopes = []
            for start in range(0, len(span) - 40 + 1, 4):
                ks = span[start:start + 40]
                slopes.append(ols([gap[k] for k in ks], [inflation[k] for k in ks])[0])
            m = statistics.median(slopes)
            med[iso] = m
            share[iso] = sum(1 for s in slopes if s * m < 0) / len(slopes)
        sample = "to 2019Q4, economies with 80 or more common quarters"
        method = "wage inflation = 400 x quarterly log change of the hourly earnings index (private sector, else " \
                 "manufacturing); gap = unemployment rate less its HP 1600 trend; OLS over 40-quarter windows " \
                 "stepped by 4"
        self.add("F06", "median slope of wage inflation on the unemployment gap across windows", med, "quarterly",
                 sample, method, "macro/labour_q")
        self.add("F06", "share of windows whose slope has the opposite sign to the median's", share, "quarterly",
                 sample, method, "macro/labour_q")

    # F07 --------------------------------------------------------------------------------------------------------
    def f07(self):
        per, pool = {}, {}
        for iso in self.d.level:
            y = self.d.quarterly(iso, "gdp")
            if len(y) < 80:
                continue
            keys = sorted(y)
            z = standardise([y[k] - y[k - 1] for k in keys[1:]])
            per[iso] = subbotin(z)
            pool.setdefault(self.d.level[iso], []).extend(z)
        self.add("F07", "exponential-power shape of standardised output growth (1 Laplace, 2 normal)", per,
                 "quarterly", "to 2019Q4, economies with 80 or more quarters", "profile maximum likelihood per "
                 "economy on quarterly log growth standardised by its own mean and sd; 'pooled' is the fit to the "
                 "level's pooled standardised growths", "macro/qna",
                 pooled={lvl: subbotin(z) for lvl, z in pool.items()})
        per_a, pool_a = {}, {}
        for iso, years in self.d.pwt.items():
            span = runs({y: 0 for y, v in years.items() if v["rgdpna"]}, last=LAST_YEAR)
            if len(span) < 30:
                continue
            z = standardise([math.log(float(years[k]["rgdpna"]) / float(years[k - 1]["rgdpna"])) for k in span[1:]])
            per_a[iso] = subbotin(z)
            pool_a.setdefault(self.d.level.get(iso), []).extend(z)
        self.add("F07", "exponential-power shape of standardised output growth (annual)", per_a, "annual",
                 "1950-2019, economies with 30 or more years", "as the quarterly row, on annual log growth of rgdpna",
                 "macro/pwt", pooled={lvl: subbotin(z) for lvl, z in pool_a.items() if lvl})

    # F13 --------------------------------------------------------------------------------------------------------
    def f13(self):
        shares = {}
        for r in csv.DictReader((RAW / "wid" / "shares.csv").open()):
            if int(r["year"]) <= LAST_YEAR:
                shares.setdefault((r["iso3"], r["what"]), {}).setdefault(int(r["year"]), {})[r["group"]] = \
                    float(r["share"])
        alpha = {"income": {}, "wealth": {}}
        for (iso, what), years in shares.items():
            ys = [y for y, g in years.items() if g.get("p99p100") and g.get("p99.9p100")]
            if ys:
                g = years[max(ys)]
                alpha[what][iso] = 1 / (1 + math.log10(g["p99.9p100"] / g["p99p100"]))
        method = "alpha = 1 / (1 + log10(S0.1 / S1)), S the top 0.1% and top 1% shares (a Pareto tail's identity), " \
                 "latest year to 2019"
        src = "wid/shares (World Inequality Database)"
        self.add("F13", "Pareto exponent of top incomes", alpha["income"], "annual", "latest year to 2019", method, src)
        self.add("F13", "Pareto exponent of top wealth", alpha["wealth"], "annual", "latest year to 2019", method, src)
        diff = {i: alpha["wealth"][i] - alpha["income"][i] for i in alpha["wealth"] if i in alpha["income"]}
        self.add("F13", "wealth's exponent less income's", diff, "annual", "latest year to 2019", method, src)

    # F16, F18: JST (developed) --------------------------------------------------------------------------------
    def f16_f18(self):
        conc = {}
        xs, ys = [], []
        for iso, years in self.d.jst.items():
            real = {y: math.log(float(v["tloans"]) / float(v["cpi"])) for y, v in years.items()
                    if v["tloans"] and v["cpi"] and float(v["tloans"]) > 0}
            gdp = {y: math.log(float(v["rgdpbarro"])) for y, v in years.items() if v["rgdpbarro"]}
            span = runs({y: 0 for y in common(real, gdp)}, first=1950, last=LAST_YEAR)
            if len(span) >= 30:
                phase = []
                for s in (real, gdp):
                    pts = turning_points([s[k] for k in span], 1, 1, 2)
                    state = {}
                    for (a, ta), (b, _) in zip(pts, pts[1:]):
                        for t in range(a + 1, b + 1):
                            state[t] = ta == "T"
                    phase.append(state)
                both = common(*phase)
                if both:
                    conc[iso] = sum(1 for t in both if phase[0][t] == phase[1][t]) / len(both)
            for y, v in years.items():
                lags = [y - k for k in range(1, 6)]
                if y > LAST_YEAR or y < 1870 + 6 or not v["crisisJST"] or any(
                        (l not in real or l - 1 not in real) for l in lags):
                    continue
                if years.get(y - 1, {}).get("crisisJST") == "1":
                    continue
                xs.append([real[l] - real[l - 1] for l in lags])
                ys.append(1 if v["crisisJST"] == "1" else 0)
        self.add("F16", "concordance of credit and output cycles: the share of years both are in the same phase",
                 conc, "annual", "1950-2019, JST economies",
                 "Harding-Pagan dating (annual: window 1, phase 1, cycle 2) "
                 "of log real private loans (tloans / cpi) and log real GDP per capita (rgdpbarro); share of years in "
                 "the same phase", "macro/jst")
        beta = binary_fit(xs, ys, "logit")
        scores = [sum(b * v for b, v in zip(beta, [1.0, *x])) for x in xs]
        for stat, value in (("sum of the coefficients on lagged real credit growth", sum(beta[1:])),
                            ("area under the ROC curve", auc(scores, ys))):
            self.rows.append(["F18", stat, "developed", len(self.d.jst), "", "", "", f"{value:.4g}", "annual",
                              f"1876-2019 pooled, {len(ys)} country-years, {sum(ys)} onsets",
                              "pooled logit of a systemic-crisis onset (crisisJST, years after a crisis year dropped) "
                              "on five lags of the log change of real private loans, no country effects",
                              "macro/jst", " ".join(sorted(self.d.jst))])

    # F16, F18 by level: WDI private credit and Laeven-Valencia banking crises ------------------------------------
    def f16_f18_levels(self):
        onsets = {}
        for r in csv.DictReader((RAW / "bank" / "crisis_years.csv").open()):
            if r["kind"] == "banking":
                onsets.setdefault(r["iso3"], set()).add(int(r["year"]))
        conc, xs, ys, used = {}, {}, {}, {}
        for iso, years in self.d.ann.items():
            level = self.d.level.get(iso)
            # Credit in constant prices: its share of GDP times real GDP, so credit is deflated as output is.
            real = {y: math.log(float(v["private_credit_pct_gdp"]) * float(v["gdp_real"])) for y, v in years.items()
                    if y <= LAST_YEAR and v["private_credit_pct_gdp"] and v["gdp_real"]
                    and float(v["private_credit_pct_gdp"]) > 0}
            gdp = {y: math.log(float(v["gdp_real"])) for y, v in years.items() if y <= LAST_YEAR and v["gdp_real"]}
            span = runs({y: 0 for y in common(real, gdp)}, last=LAST_YEAR)
            if len(span) >= 30:
                phase = []
                for series in (real, gdp):
                    pts = turning_points([series[k] for k in span], 1, 1, 2)
                    state = {}
                    for (a, ta), (b, _) in zip(pts, pts[1:]):
                        for t in range(a + 1, b + 1):
                            state[t] = ta == "T"
                    phase.append(state)
                both = common(*phase)
                if both:
                    conc[iso] = sum(1 for t in both if phase[0][t] == phase[1][t]) / len(both)
            if not level:
                continue
            crises = onsets.get(iso, set())
            for y in span:
                lags = [y - k for k in range(1, 6)]
                if any(l not in real or l - 1 not in real for l in lags) or y - 1 in crises:
                    continue
                xs.setdefault(level, []).append([real[l] - real[l - 1] for l in lags])
                ys.setdefault(level, []).append(1 if y in crises else 0)
                used.setdefault(level, set()).add(iso)
        self.add("F16", "concordance of credit and output cycles: the share of years both are in the same phase",
                 conc, "annual", "1960-2019, 30 or more years",
                 "Harding-Pagan dating (annual: window 1, phase 1, cycle 2) of log real private credit (WDI "
                 "FS.AST.PRVT.GD.ZS times real GDP) and log real GDP; share of years in the same phase",
                 "macro/annual")
        for level in ("developed", "emerging", "developing"):
            if level not in ys or sum(ys[level]) < 5:
                continue
            beta = binary_fit(xs[level], ys[level], "logit")
            scores = [sum(b * v for b, v in zip(beta, [1.0, *x])) for x in xs[level]]
            for stat, value in (("sum of the coefficients on lagged real credit growth", sum(beta[1:])),
                                ("area under the ROC curve", auc(scores, ys[level]))):
                self.rows.append(["F18", stat, level, len(used[level]), "", "", "", f"{value:.4g}", "annual",
                                  f"1960-2019 pooled, {len(ys[level])} country-years, {sum(ys[level])} onsets",
                                  "pooled logit of a systemic banking crisis onset (Laeven-Valencia, years after an "
                                  "onset dropped) on five lags of the log change of real private credit (WDI "
                                  "FS.AST.PRVT.GD.ZS times real GDP), no country effects",
                                  "macro/annual, bank/crisis_years", " ".join(sorted(used[level]))])

    # F19 --------------------------------------------------------------------------------------------------------
    def f19(self):
        impact, long_run = {}, {}
        for iso in self.d.level:
            lend, pol = self.d.pr.get((iso, "lending_rate")), self.d.pr.get((iso, "policy_rate"))
            if not lend or not pol:
                continue
            span = runs({k: 0 for k in common(lend, pol)}, last=LAST_YEAR * 4 + 3)
            if len(span) < 40 or statistics.pstdev([pol[k] for k in span]) == 0:
                continue
            theta, c = ols([pol[k] for k in span], [lend[k] for k in span])
            rows = [(pol[k] - pol[k - 1], lend[k - 1] - c - theta * pol[k - 1], lend[k] - lend[k - 1])
                    for k in span[1:]]
            xs = [[r[0], r[1]] for r in rows]
            yv = [r[2] for r in rows]
            # OLS of the change on the policy change and the lagged gap, by the normal equations.
            xtx = [[sum(a * b for a, b in zip(col_i, col_j)) for col_j in zip(*[[1.0, *x] for x in xs])]
                   for col_i in zip(*[[1.0, *x] for x in xs])]
            xty = [sum(a * b for a, b in zip(col, yv)) for col in zip(*[[1.0, *x] for x in xs])]
            impact[iso] = solve(xtx, xty)[1]
            long_run[iso] = theta
        method = "Engle-Granger error correction on quarterly averages: long run = OLS of the lending rate on the " \
                 "policy rate in levels; impact = coefficient on the policy rate's change in the change of the " \
                 "lending rate with the lagged long-run gap (quarterly, not monthly as the fact reads)"
        sample = "to 2019Q4, economies with 40 or more consecutive common quarters"
        self.add("F19", "impact pass-through to loan rates", impact, "quarterly", sample, method,
                 "macro/prices_rates_q")
        self.add("F19", "long-run pass-through to loan rates", long_run, "quarterly", sample, method,
                 "macro/prices_rates_q")

    # F23 --------------------------------------------------------------------------------------------------------
    def f23(self):
        spread, coef = {}, {}
        for iso in self.d.level:
            long_ = self.d.pr.get((iso, "long_rate"))
            short = self.d.pr.get((iso, "tbill_rate")) or self.d.pr.get((iso, "short_rate_3m"))
            if not long_ or not short:
                continue
            keys = [k for k in common(long_, short) if k <= LAST_YEAR * 4 + 3]
            if len(keys) < 40:
                continue
            s = {k: long_[k] - short[k] for k in keys}
            spread[iso] = statistics.fmean(s.values())
            y = self.d.quarterly(iso, "gdp")
            if len(y) < 80:
                continue
            ky = sorted(y)
            pts = turning_points([y[k] for k in ky], 2, 2, 5)
            recession = set()
            for (a, ta), (b, _) in zip(pts, pts[1:]):
                if ta == "P":
                    recession.update(ky[t] for t in range(a + 1, b + 1))
            obs = [(s[k], 1 if any(k + h in recession for h in range(1, 5)) else 0) for k in keys
                   if k + 4 <= ky[-1] and k >= ky[0] and k not in recession]
            if len(obs) >= 80 and 3 <= sum(o[1] for o in obs) < len(obs) - 3:
                coef[iso] = binary_fit([[o[0]] for o in obs], [o[1] for o in obs], "probit")[1]
        self.add("F23", "mean ten-year less bill spread, percentage points", spread, "quarterly",
                 "to 2019Q4, economies with 40 or more common quarters", "long-term government bond yield (OECD) "
                 "less the Treasury-bill yield (IMF, else OECD three-month rate), quarterly averages",
                 "macro/prices_rates_q")
        self.add("F23", "probit coefficient on the spread, four quarters ahead", coef, "quarterly",
                 "to 2019Q4, economies with 80 or more quarters and three or more recession-ahead quarters",
                 "probit of a recession (by the F01 dating) beginning within the next four quarters on the spread, "
                 "quarters already in recession dropped", "macro/prices_rates_q, macro/qna")

    # F27 --------------------------------------------------------------------------------------------------------
    def f27(self):
        freq, pool = {}, {}
        for iso, years in self.d.ann.items():
            inflow = {y: -float(v["financial_account_usd"]) / float(v["gdp_usd"]) * 100 for y, v in years.items()
                      if v["financial_account_usd"] and v["gdp_usd"] and y <= LAST_YEAR}
            change = {y: inflow[y] - inflow[y - 1] for y in inflow if y - 1 in inflow}
            if len(change) < 20:
                continue
            m, s = statistics.fmean(change.values()), statistics.stdev(change.values())
            growth = {y: float(v["gdp_real"]) for y, v in years.items() if v["gdp_real"]}
            stops = [y for y, c in change.items() if c < m - 2 * s and any(
                y + h in growth and y + h - 1 in growth and growth[y + h] < growth[y + h - 1] for h in (0, 1))]
            freq[iso] = len(stops) / len(change)
            lvl = self.d.level.get(iso)
            if lvl:
                counts = pool.setdefault(lvl, [0, 0])
                counts[0] += len(stops)
                counts[1] += len(change)
        self.add("F27", "sudden stops per country-year", freq, "annual", "to 2019, economies with 20 or more years",
                 "annual version of Calvo, Izquierdo and Mejía: net capital inflow = minus the net financial account "
                 "(BPM6) over GDP; a stop is a year its change falls two sd below the economy's mean change, with "
                 "real GDP falling that year or the next; 'pooled' is the level's stops over its country-years",
                 "macro/annual (WDI)", pooled={lvl: n / years for lvl, (n, years) in pool.items()})

    # S1.16's macro reads ---------------------------------------------------------------------------------------
    def reads(self):
        def pwt_mean(col):
            out = {}
            for iso, years in self.d.pwt.items():
                xs = [float(years[y][col]) for y in range(2010, LAST_YEAR + 1) if y in years and years[y][col]]
                if len(xs) >= 5:
                    out[iso] = statistics.fmean(xs)
            return out
        src = "macro/pwt"
        self.add("S1.16", "labour share of GDP", pwt_mean("labsh"), "annual", "mean of 2010-2019", "PWT labsh", src)
        self.add("S1.16", "investment share of GDP (current PPPs)", pwt_mean("csh_i"), "annual", "mean of 2010-2019",
                 "PWT csh_i", src)
        self.add("S1.16", "household consumption share of GDP (current PPPs)", pwt_mean("csh_c"), "annual",
                 "mean of 2010-2019", "PWT csh_c", src)
        self.add("S1.16", "average annual hours per person engaged", pwt_mean("avh"), "annual", "mean of 2010-2019",
                 "PWT avh", src)
        u = {}
        for iso, years in self.d.ann.items():
            xs = [float(years[y]["unemployment_survey"]) for y in range(2010, LAST_YEAR + 1)
                  if y in years and years[y]["unemployment_survey"]]
            if len(xs) >= 5:
                u[iso] = statistics.fmean(xs)
        self.add("S1.16", "unemployment rate, per cent", u, "annual", "mean of 2010-2019",
                 "ILO national survey estimates", "macro/annual")
        mean_inf, persistence = {}, {}
        for iso in self.d.level:
            p = self.d.pr.get((iso, "cpi"))
            if not p:
                continue
            span = runs({k: 0 for k, v in p.items() if v > 0}, first=2000 * 4, last=LAST_YEAR * 4 + 3)
            if len(span) < 40:
                continue
            lp = deseason([math.log(p[k]) for k in span], span[0] % 4)
            inf = [400 * (b - a) for a, b in zip(lp, lp[1:])]
            mean_inf[iso] = statistics.fmean(inf)
            persistence[iso] = ols(inf[:-1], inf[1:])[0]
        self.add("S1.16", "CPI inflation, per cent a year", mean_inf, "quarterly", "2000Q1-2019Q4",
                 "mean of 400 x quarterly log change of the CPI, deseasonalised", "macro/prices_rates_q")
        self.add("S1.16", "persistence of CPI inflation (AR(1) coefficient)", persistence, "quarterly",
                 "2000Q1-2019Q4", "OLS of quarterly annualised inflation on its lag", "macro/prices_rates_q")
        sd, ac = {}, {}
        for iso in self.d.level:
            y = self.d.quarterly(iso, "gdp")
            if len(y) < 40:
                continue
            keys = sorted(y)
            cy = cycles(y, 1600)
            sd[iso] = 100 * statistics.pstdev(cy.values())
            g = [y[k] - y[k - 1] for k in keys[1:]]
            ac[iso] = corr(g[:-1], g[1:])
        self.add("S1.16", "sd of cyclical output, per cent", sd, "quarterly", "to 2019Q4, 40 or more quarters",
                 "HP 1600 on log real GDP", "macro/qna")
        self.add("S1.16", "first-order autocorrelation of quarterly output growth", ac, "quarterly",
                 "to 2019Q4, 40 or more quarters", "quarterly log growth of real GDP", "macro/qna")


def ranges(cache: Path, manifest: dict, iso3: set) -> None:
    r = Ranges(Data())
    for step in (r.f01, r.f02_f03, r.f04, r.f05, r.f06, r.f07, r.f13, r.f16_f18, r.f16_f18_levels, r.f19, r.f23, r.f27, r.reads):
        step()
        log(f"ranges: {step.__name__}")
    header = ["fact", "statistic", "level", "economies", "median", "p25", "p75", "pooled", "frequency", "sample",
              "method", "source", "iso3"]
    with (OUT / "level_ranges.csv").open("w", newline="") as f:
        w = csv.writer(f)
        w.writerow(header)
        w.writerows(r.rows)
    manifest["series"]["macro/level_ranges"] = {
        "title": "Each measurable N3 fact's statistic (and S1.16's macro reads), per development level (World Bank "
                 "income groups: developed High; emerging Upper-middle; developing Lower-middle and Low): median and "
                 "interquartile range across the level's economies, with the method, sample and member economies",
        "rows": len(r.rows),
        "for": ["S1.16", "S7.01"],
        "note": "Derived by fetch_macro.py (ranges) from the macro/ files and wid/shares; each row's method column "
                "states the computation. Quarterly statistics end at 2019Q4. Rows naming a paper group the per-economy "
                "values typed from it by current World Bank level.",
    }


SOURCES = {"pwt": pwt, "qna": qna, "labour": labour, "prices": prices, "annual": annual, "jst": jst, "xru": xru,
           "typed": typed, "ranges": ranges}


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--cache", type=Path, default=None)
    parser.add_argument("--only", nargs="*", choices=sorted(SOURCES))
    args = parser.parse_args()
    cache = args.cache or Path(tempfile.mkdtemp())
    cache.mkdir(parents=True, exist_ok=True)
    manifest = {"sources": {}, "series": {}}
    iso3 = countries()
    for name in args.only or list(SOURCES):
        SOURCES[name](cache, manifest, iso3)
        log(f"{name} done")
    merge_manifest(manifest["sources"], manifest["series"])


if __name__ == "__main__":
    main()
